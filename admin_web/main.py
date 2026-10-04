from fastapi import FastAPI, Request, Form, Depends, HTTPException, status, BackgroundTasks, File, UploadFile, Query
from fastapi.responses import HTMLResponse, RedirectResponse, JSONResponse, Response, FileResponse
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles
from fastapi.security import HTTPBasic, HTTPBasicCredentials
from starlette.middleware.sessions import SessionMiddleware
import secrets
import bcrypt
import aiosqlite
import os
import asyncio
import httpx
import json
import shutil
import mimetypes
from aiogram import Bot
from bot.config import load_settings
from bot.strings import get_string
from datetime import datetime, timedelta
import re
from contextlib import asynccontextmanager


BASE_DIR = "/app"
DB_PATH = "/app/data/bot.db"

os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)

UPLOAD_DIR = "/app/data/uploads/models"
STATIC_DIR = "/app/admin_web/static"


os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(STATIC_DIR, exist_ok=True)

async def run_migrations(db: aiosqlite.Connection):

    await db.execute("""
        CREATE TABLE IF NOT EXISTS gcp_usage (
            project_id TEXT PRIMARY KEY,
            generation_count INTEGER NOT NULL DEFAULT 0
        );
        """)
    await db.commit()
    await db.execute(
        """
        CREATE TABLE IF NOT EXISTS zveno_usage_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            key_id INTEGER,
            key_preview TEXT,
            model TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'ok',
            prompt_tokens INTEGER NOT NULL DEFAULT 0,
            completion_tokens INTEGER NOT NULL DEFAULT 0,
            total_tokens INTEGER NOT NULL DEFAULT 0,
            input_rub REAL NOT NULL DEFAULT 0,
            output_rub REAL NOT NULL DEFAULT 0,
            total_rub REAL NOT NULL DEFAULT 0,
            latency_ms INTEGER NOT NULL DEFAULT 0,
            error_text TEXT
        );
        """
    )
    await db.commit()

    
    async with db.execute("PRAGMA table_info(subscriptions)") as cur:
        cols = [row[1] for row in await cur.fetchall()]
    if cols and "plan_id" not in cols:
        try:
            await db.execute("ALTER TABLE subscriptions ADD COLUMN plan_id INTEGER")
            await db.commit()
        except Exception as e:
            print(f"Migration error (subscriptions.plan_id): {e}")

    if cols and "individual_api_key" not in cols:
        try:
            await db.execute("ALTER TABLE subscriptions ADD COLUMN individual_api_key TEXT")
            await db.commit()
        except Exception as e:
            print(f"Migration error (subscriptions.individual_api_key): {e}")

    
    await db.execute("""
    CREATE TABLE IF NOT EXISTS proxies (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT DEFAULT '',
        url TEXT NOT NULL,
        is_active INTEGER NOT NULL DEFAULT 1,
        status TEXT DEFAULT 'unknown',
        last_check TIMESTAMP,
        error_message TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """)
    await db.commit()

    
    try:
        async with db.execute("PRAGMA table_info(proxies)") as cur:
            cols = [row[1] for row in await cur.fetchall()]
        if cols and "name" not in cols:
            await db.execute("ALTER TABLE proxies ADD COLUMN name TEXT DEFAULT ''")
            await db.commit()
    except Exception:
        pass

    
    await db.execute("""
    CREATE TABLE IF NOT EXISTS api_key_proxies (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        key_id INTEGER NOT NULL,
        proxy_id INTEGER NOT NULL,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        UNIQUE(key_id, proxy_id),
        FOREIGN KEY(key_id) REFERENCES api_keys(id) ON DELETE CASCADE,
        FOREIGN KEY(proxy_id) REFERENCES proxies(id) ON DELETE CASCADE
    );
    """)
    await db.commit()

    
    try:
        import sys
        if BASE_DIR not in sys.path:
            sys.path.append(BASE_DIR)
        from scripts.migrate_proxies import migrate_proxies
        await migrate_proxies()
        
        
        DEFAULT_PROXY = "http://BKM30EMK:Z7DGJIDF@141.11.162.29:47066"
        async with db.execute("SELECT COUNT(*) FROM proxies WHERE url LIKE ?", ("%141.11.162.29%",)) as cur:
            has_default = (await cur.fetchone())[0] > 0
        if not has_default:
            await db.execute("INSERT INTO proxies (url, is_active, status) VALUES (?, 1, 'unknown')", (DEFAULT_PROXY,))
            await db.commit()
            print("Added default proxy (141.11.162.29)")
        
        
        async with db.execute("SELECT id, url FROM proxies") as cur:
            rows = await cur.fetchall()
        for p_id, url in rows:
            url = url.strip()
            parts = url.split(':')
            if len(parts) == 4 and "://" not in url:
                new_url = f"http://{parts[2]}:{parts[3]}@{parts[0]}:{parts[1]}"
                await db.execute("UPDATE proxies SET url = ? WHERE id = ?", (new_url, p_id))
        await db.commit()
        
    except Exception as e:
        print(f"Proxy migration warning (non-critical): {e}")

    async with db.execute("PRAGMA table_info(users)") as cur:
        cols = [row[1] for row in await cur.fetchall()]
    if cols:
        if "trial_used" not in cols:
            try:
                await db.execute("ALTER TABLE users ADD COLUMN trial_used INTEGER NOT NULL DEFAULT 0")
                await db.commit()
            except Exception as e:
                print(f"Migration error (users.trial_used): {e}")
        if "balance" not in cols:
            try:
                await db.execute("ALTER TABLE users ADD COLUMN balance INTEGER NOT NULL DEFAULT 50")
                await db.commit()
            except Exception as e:
                print(f"Migration error (users.balance): {e}")
        if "generation_price" not in cols:
            try:
                await db.execute("ALTER TABLE users ADD COLUMN generation_price INTEGER NOT NULL DEFAULT 30")
                await db.commit()
            except Exception as e:
                print(f"Migration error (users.generation_price): {e}")
        
        
        try:
            await db.execute("UPDATE users SET generation_price = 30 WHERE generation_price IS NULL")
            await db.commit()
        except Exception:
            pass
        async with db.execute("PRAGMA table_info(users)") as cur:
            cols_u2 = [row[1] for row in await cur.fetchall()]
        if "custom_1k_price_rub" not in cols_u2:
            try:
                await db.execute("ALTER TABLE users ADD COLUMN custom_1k_price_rub INTEGER NULL")
                await db.execute(
                    "UPDATE users SET custom_1k_price_rub = generation_price "
                    "WHERE generation_price IS NOT NULL AND generation_price NOT IN (30, 25)"
                )
                await db.execute(
                    "UPDATE users SET generation_price = COALESCE(custom_1k_price_rub, 30)"
                )
                await db.commit()
            except Exception as e:
                print(f"Migration error (users.custom_1k_price_rub): {e}")
        if "created_at" not in cols:
            try:
                await db.execute("ALTER TABLE users ADD COLUMN created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP")
                await db.commit()
            except Exception as e:
                print(f"Migration error (users.created_at): {e}")
        if "source" not in cols:
            try:
                await db.execute("ALTER TABLE users ADD COLUMN source TEXT NOT NULL DEFAULT 'telegram'")
                await db.commit()
            except Exception as e:
                print(f"Migration error (users.source): {e}")

    
    await db.execute("""
    CREATE TABLE IF NOT EXISTS balance_history (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_id INTEGER NOT NULL,
        amount INTEGER NOT NULL,
        new_balance INTEGER NOT NULL,
        reason TEXT NOT NULL,
        admin_id TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(user_id) REFERENCES users(id)
    );
    """)
    await db.commit()

    
    await db.execute("""
    CREATE TABLE IF NOT EXISTS site_users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        email TEXT UNIQUE NOT NULL,
        password_hash TEXT NOT NULL,
        balance INTEGER NOT NULL DEFAULT 50,
        language TEXT NOT NULL DEFAULT 'ru',
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """)
    await db.commit()

    
    async with db.execute("PRAGMA table_info(step_options)") as cur:
        cols = [row[1] for row in await cur.fetchall()]
    if cols and "custom_prompt" not in cols:
        try:
            await db.execute("ALTER TABLE step_options ADD COLUMN custom_prompt TEXT")
            await db.commit()
            print("Migration: Added custom_prompt to step_options")
        except Exception as e:
            print(f"Migration error (step_options.custom_prompt): {e}")

    
    try:
        async with db.execute("PRAGMA table_info(steps)") as cur:
            s_cols = [row[1] for row in await cur.fetchall()]
        if s_cols and "question_text_en" not in s_cols:
            await db.execute("ALTER TABLE steps ADD COLUMN question_text_en TEXT")
            await db.commit()
        if s_cols and "question_text_vi" not in s_cols:
            await db.execute("ALTER TABLE steps ADD COLUMN question_text_vi TEXT")
            await db.commit()
    except Exception as e:
        print(f"Migration error (steps translations): {e}")

    
    try:
        async with db.execute("PRAGMA table_info(library_options)") as cur:
            l_cols = [row[1] for row in await cur.fetchall()]
        if l_cols and "option_text_en" not in l_cols:
            await db.execute("ALTER TABLE library_options ADD COLUMN option_text_en TEXT")
            await db.commit()
        if l_cols and "option_text_vi" not in l_cols:
            await db.execute("ALTER TABLE library_options ADD COLUMN option_text_vi TEXT")
            await db.commit()
    except Exception as e:
        print(f"Migration error (library_options translations): {e}")

    
    try:
        async with db.execute("PRAGMA table_info(step_options)") as cur:
            so_cols = [row[1] for row in await cur.fetchall()]
        if so_cols and "option_text_en" not in so_cols:
            await db.execute("ALTER TABLE step_options ADD COLUMN option_text_en TEXT")
            await db.commit()
        if so_cols and "option_text_vi" not in so_cols:
            await db.execute("ALTER TABLE step_options ADD COLUMN option_text_vi TEXT")
            await db.commit()
    except Exception as e:
        print(f"Migration error (step_options translations): {e}")

    
    try:
        await db.execute(
            """
            CREATE TABLE IF NOT EXISTS library_step_options (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                step_id INTEGER NOT NULL,
                option_text TEXT NOT NULL,
                option_value TEXT NOT NULL,
                order_index INTEGER NOT NULL DEFAULT 0,
                custom_prompt TEXT,
                FOREIGN KEY(step_id) REFERENCES library_steps(id)
            );
            """
        )
        await db.commit()
    except Exception as e:
        print(f"Migration error (library_step_options): {e}")

    
    try:
        
        async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Системные",)) as cur:
            row = await cur.fetchone()
        if row:
            sys_cat_id = row[0]
        else:
            await db.execute("INSERT INTO button_categories (name) VALUES (?)", ("Системные",))
            await db.commit()
            async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Системные",)) as cur:
                sys_cat_id = (await cur.fetchone())[0]

        system_buttons = [
            ("Свой вариант", "custom", "Введите ваш вариант текста:"),
            ("Пропустить", "skip", None),
            ("Назад", "back", None),
        ]
        for text, value, prompt in system_buttons:
            async with db.execute(
                "SELECT id FROM library_options WHERE category_id=? AND option_value=?",
                (sys_cat_id, value)
            ) as cur:
                if not await cur.fetchone():
                    await db.execute(
                        "INSERT INTO library_options (category_id, option_text, option_value, custom_prompt) VALUES (?, ?, ?, ?)",
                        (sys_cat_id, text, value, prompt)
                    )
        await db.commit()
    except Exception as e:
        print(f"Migration error (library system buttons): {e}")

    try:
        
        async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Праздники",)) as cur:
            row = await cur.fetchone()
        if row:
            holiday_cat_id = row[0]
        else:
            await db.execute("INSERT INTO button_categories (name) VALUES (?)", ("Праздники",))
            await db.commit()
            async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Праздники",)) as cur:
                holiday_cat_id = (await cur.fetchone())[0]

        holiday_buttons = [
            ("Новый год", "new_year", None),
            ("Рождество", "christmas", None),
            ("День рождения", "birthday", None),
            ("8 марта", "mar8", None),
            ("14 февраля", "feb14", None),
            ("23 февраля", "feb23", None),
            ("Свадьба", "wedding", None),
            ("Выпускной", "graduation", None),
            ("Хэллоуин", "halloween", None),
        ]
        for text, value, prompt in holiday_buttons:
            async with db.execute(
                "SELECT id FROM library_options WHERE category_id=? AND option_value=?",
                (holiday_cat_id, value)
            ) as cur:
                if not await cur.fetchone():
                    await db.execute(
                        "INSERT INTO library_options (category_id, option_text, option_value, custom_prompt) VALUES (?, ?, ?, ?)",
                        (holiday_cat_id, text, value, prompt)
                    )
        await db.commit()
    except Exception as e:
        print(f"Migration error (library holidays): {e}")

    try:
        
        async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Размеры",)) as cur:
            row = await cur.fetchone()
        if row:
            dim_cat_id = row[0]
        else:
            await db.execute("INSERT INTO button_categories (name) VALUES (?)", ("Размеры",))
            await db.commit()
            async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Размеры",)) as cur:
                dim_cat_id = (await cur.fetchone())[0]

        dim_buttons = [
            ("Ширина", "width", None),
            ("Высота", "height", None),
            ("Длина", "length", None),
        ]
        for text, value, prompt in dim_buttons:
            async with db.execute(
                "SELECT id FROM library_options WHERE category_id=? AND option_value=?",
                (dim_cat_id, value)
            ) as cur:
                if not await cur.fetchone():
                    await db.execute(
                        "INSERT INTO library_options (category_id, option_text, option_value, custom_prompt) VALUES (?, ?, ?, ?)",
                        (dim_cat_id, text, value, prompt)
                    )
        await db.commit()
    except Exception as e:
        print(f"Migration error (library dimensions): {e}")

    try:
        
        async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Преимущества",)) as cur:
            row = await cur.fetchone()
        if row:
            adv_cat_id = row[0]
        else:
            await db.execute("INSERT INTO button_categories (name) VALUES (?)", ("Преимущества",))
            await db.commit()
            async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Преимущества",)) as cur:
                adv_cat_id = (await cur.fetchone())[0]

        adv_buttons = [
            ("Топ 1 преимущества товара", "adv_1", None),
            ("Топ 2 преимущества товара", "adv_2", None),
            ("Топ 3 преимущества товара", "adv_3", None),
        ]
        for text, value, prompt in adv_buttons:
            async with db.execute(
                "SELECT id FROM library_options WHERE category_id=? AND option_value=?",
                (adv_cat_id, value)
            ) as cur:
                if not await cur.fetchone():
                    await db.execute(
                        "INSERT INTO library_options (category_id, option_text, option_value, custom_prompt) VALUES (?, ?, ?, ?)",
                        (adv_cat_id, text, value, prompt)
                    )
        await db.commit()
    except Exception as e:
        print(f"Migration error (library advantages): {e}")

    try:
        
        async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Дополнительно",)) as cur:
            row = await cur.fetchone()
        if row:
            extra_cat_id = row[0]
        else:
            await db.execute("INSERT INTO button_categories (name) VALUES (?)", ("Дополнительно",))
            await db.commit()
            async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Дополнительно",)) as cur:
                extra_cat_id = (await cur.fetchone())[0]

        extra_buttons = [
            ("Дополнительная информация о продукте", "extra_info", None),
        ]
        for text, value, prompt in extra_buttons:
            async with db.execute(
                "SELECT id FROM library_options WHERE category_id=? AND option_value=?",
                (extra_cat_id, value)
            ) as cur:
                if not await cur.fetchone():
                    await db.execute(
                        "INSERT INTO library_options (category_id, option_text, option_value, custom_prompt) VALUES (?, ?, ?, ?)",
                        (extra_cat_id, text, value, prompt)
                    )
        await db.commit()
    except Exception as e:
        print(f"Migration error (library extra): {e}")

    try:
        
        async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Человек",)) as cur:
            row = await cur.fetchone()
        if row:
            person_cat_id = row[0]
        else:
            await db.execute("INSERT INTO button_categories (name) VALUES (?)", ("Человек",))
            await db.commit()
            async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Человек",)) as cur:
                person_cat_id = (await cur.fetchone())[0]

        person_buttons = [
            ("Человек присутствует", "person_yes", None),
            ("Без человека", "person_no", None),
        ]
        for text, value, prompt in person_buttons:
            async with db.execute(
                "SELECT id FROM library_options WHERE category_id=? AND option_value=?",
                (person_cat_id, value)
            ) as cur:
                if not await cur.fetchone():
                    await db.execute(
                        "INSERT INTO library_options (category_id, option_text, option_value, custom_prompt) VALUES (?, ?, ?, ?)",
                        (person_cat_id, text, value, prompt)
                    )
        await db.commit()
    except Exception as e:
        print(f"Migration error (library person): {e}")

    try:
        
        async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Продукт",)) as cur:
            row = await cur.fetchone()
        if row:
            product_cat_id = row[0]
        else:
            await db.execute("INSERT INTO button_categories (name) VALUES (?)", ("Продукт",))
            await db.commit()
            async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Продукт",)) as cur:
                product_cat_id = (await cur.fetchone())[0]

        product_buttons = [
            ("Название товара/бренда", "product_name", None),
        ]
        for text, value, prompt in product_buttons:
            async with db.execute(
                "SELECT id FROM library_options WHERE category_id=? AND option_value=?",
                (product_cat_id, value)
            ) as cur:
                if not await cur.fetchone():
                    await db.execute(
                        "INSERT INTO library_options (category_id, option_text, option_value, custom_prompt) VALUES (?, ?, ?, ?)",
                        (product_cat_id, text, value, prompt)
                    )
        await db.commit()
    except Exception as e:
        print(f"Migration error (library product): {e}")

    try:
        
        async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Стиль",)) as cur:
            row = await cur.fetchone()
        if row:
            style_cat_id = row[0]
        else:
            await db.execute("INSERT INTO button_categories (name) VALUES (?)", ("Стиль",))
            await db.commit()
            async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Стиль",)) as cur:
                style_cat_id = (await cur.fetchone())[0]

        style_buttons = [
            ("Стиль", "style", None),
        ]
        for text, value, prompt in style_buttons:
            async with db.execute(
                "SELECT id FROM library_options WHERE category_id=? AND option_value=?",
                (style_cat_id, value)
            ) as cur:
                if not await cur.fetchone():
                    await db.execute(
                        "INSERT INTO library_options (category_id, option_text, option_value, custom_prompt) VALUES (?, ?, ?, ?)",
                        (style_cat_id, text, value, prompt)
                    )
        await db.commit()
    except Exception as e:
        print(f"Migration error (library style): {e}")

    try:
        
        async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Язык",)) as cur:
            row = await cur.fetchone()
        if row:
            lang_cat_id = row[0]
        else:
            await db.execute("INSERT INTO button_categories (name) VALUES (?)", ("Язык",))
            await db.commit()
            async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Язык",)) as cur:
                lang_cat_id = (await cur.fetchone())[0]

        lang_buttons = [
            ("Русский", "lang_ru", None),
            ("English", "lang_en", None),
            ("Tiếng Việt", "lang_vi", None),
        ]
        for text, value, prompt in lang_buttons:
            async with db.execute(
                "SELECT id FROM library_options WHERE category_id=? AND option_value=?",
                (lang_cat_id, value)
            ) as cur:
                if not await cur.fetchone():
                    await db.execute(
                        "INSERT INTO library_options (category_id, option_text, option_value, custom_prompt) VALUES (?, ?, ?, ?)",
                        (lang_cat_id, text, value, prompt)
                    )
        await db.commit()
    except Exception as e:
        print(f"Migration error (library language): {e}")

    try:
        
        async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Формат",)) as cur:
            row = await cur.fetchone()
        if row:
            format_cat_id = row[0]
        else:
            await db.execute("INSERT INTO button_categories (name) VALUES (?)", ("Формат",))
            await db.commit()
            async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Формат",)) as cur:
                format_cat_id = (await cur.fetchone())[0]

        format_buttons = [
            ("1:1", "1:1", None),
            ("9:16", "9:16", None),
            ("16:9", "16:9", None),
            ("3:4", "3:4", None),
            ("4:3", "4:3", None),
            ("3:2", "3:2", None),
            ("2:3", "2:3", None),
            ("5:4", "5:4", None),
            ("4:5", "4:5", None),
            ("21:9", "21:9", None),
        ]
        for text, value, prompt in format_buttons:
            async with db.execute(
                "SELECT id FROM library_options WHERE category_id=? AND option_value=?",
                (format_cat_id, value)
            ) as cur:
                if not await cur.fetchone():
                    await db.execute(
                        "INSERT INTO library_options (category_id, option_text, option_value, custom_prompt) VALUES (?, ?, ?, ?)",
                        (format_cat_id, text, value, prompt)
                    )
        await db.commit()
    except Exception as e:
        print(f"Migration error (library format): {e}")

    try:
        
        async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Локация (Улица)",)) as cur:
            row = await cur.fetchone()
        if row:
            out_cat_id = row[0]
        else:
            await db.execute("INSERT INTO button_categories (name) VALUES (?)", ("Локация (Улица)",))
            await db.commit()
            async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Локация (Улица)",)) as cur:
                out_cat_id = (await cur.fetchone())[0]

        out_buttons = [
            ("У машины", "car", None),
            ("У кофейни", "cafe", None),
            ("У стены", "wall", None),
            ("У здания", "building", None),
            ("Москва сити", "moscow_city", None),
            ("В лесу", "forest", None),
            ("В горах", "mountains", None),
            ("На аллее", "alley", None),
            ("В парке", "park", None),
            ("В городе", "city", None),
            ("Свой вариант", "custom", "Введите ваш вариант локации на улице (до 100 симв):"),
        ]
        for text, value, prompt in out_buttons:
            async with db.execute(
                "SELECT id FROM library_options WHERE category_id=? AND option_value=?",
                (out_cat_id, value)
            ) as cur:
                if not await cur.fetchone():
                    await db.execute(
                        "INSERT INTO library_options (category_id, option_text, option_value, custom_prompt) VALUES (?, ?, ?, ?)",
                        (out_cat_id, text, value, prompt)
                    )
        await db.commit()
    except Exception as e:
        print(f"Migration error (library street location): {e}")

    try:
        
        async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Локация (Помещение)",)) as cur:
            row = await cur.fetchone()
        if row:
            in_cat_id = row[0]
        else:
            await db.execute("INSERT INTO button_categories (name) VALUES (?)", ("Локация (Помещение)",))
            await db.commit()
            async with db.execute("SELECT id FROM button_categories WHERE name=?", ("Локация (Помещение)",)) as cur:
                in_cat_id = (await cur.fetchone())[0]

        in_buttons = [
            ("Фотостудия", "photo_studio", None),
            ("В комнате", "room", None),
            ("В ресторане", "restaurant", None),
            ("В гостинице", "hotel", None),
            ("В торговом центре", "mall", None),
            ("Свой вариант", "custom", "Введите ваш вариант помещения (до 100 симв):"),
        ]
        for text, value, prompt in in_buttons:
            async with db.execute(
                "SELECT id FROM library_options WHERE category_id=? AND option_value=?",
                (in_cat_id, value)
            ) as cur:
                if not await cur.fetchone():
                    await db.execute(
                        "INSERT INTO library_options (category_id, option_text, option_value, custom_prompt) VALUES (?, ?, ?, ?)",
                        (in_cat_id, text, value, prompt)
                    )
        await db.commit()
    except Exception as e:
        print(f"Migration error (library indoor location): {e}")

    try:
        
        async with db.execute("SELECT id FROM library_steps WHERE step_key=?", ("model_select",)) as cur:
            if not await cur.fetchone():
                await db.execute(
                    "INSERT INTO library_steps (step_key, question_text, input_type) VALUES (?, ?, ?)",
                    ("model_select", "💃 Выберите модель:", "model_select")
                )
                await db.commit()
    except Exception as e:
        print(f"Migration error (library_steps.model_select): {e}")

    try:
        
        async with db.execute("SELECT id FROM library_steps WHERE step_key=?", ("info_load",)) as cur:
            if not await cur.fetchone():
                await db.execute(
                    "INSERT INTO library_steps (step_key, question_text, input_type) VALUES (?, ?, ?)",
                    ("info_load", "📊 Укажите нагруженность инфографики (1-10):", "text")
                )
                await db.commit()
    except Exception as e:
        print(f"Migration error (library_steps.info_load): {e}")

    try:
        
        async with db.execute("SELECT id FROM library_steps WHERE step_key=?", ("aspect",)) as cur:
            if not await cur.fetchone():
                await db.execute(
                    "INSERT INTO library_steps (step_key, question_text, input_type) VALUES (?, ?, ?)",
                    ("aspect", "🖼 Выберите формат фото:", "buttons")
                )
                await db.commit()
    except Exception as e:
        print(f"Migration error (library_steps.aspect): {e}")

    try:
        
        ready_steps = [
            ("adv_1", "🏆 Укажите преимущество 1:", "text"),
            ("adv_2", "🏆 Укажите преимущество 2:", "text"),
            ("adv_3", "🏆 Укажите преимущество 3:", "text"),
            ("info_lang", "🌐 Выберите язык инфографики:", "buttons"),
            ("extra_info", "ℹ️ Дополнительная информация о продукте:", "text"),
            ("brand_name", "🏷️ Укажите название бренда/товара:", "text"),
            ("holiday", "🎉 Выберите праздник:", "buttons"),
            ("has_person", "👤 Присутствует ли человек на фото?", "buttons"),
            ("rand_location_indoor", "🏠 Выберите стиль (В помещении):", "buttons"),
            ("rand_location_outdoor", "🌳 Выберите стиль (На улице):", "buttons"),
        ]
        for key, question, i_type in ready_steps:
            async with db.execute("SELECT id FROM library_steps WHERE step_key=?", (key,)) as cur:
                if not await cur.fetchone():
                    await db.execute(
                        "INSERT INTO library_steps (step_key, question_text, input_type) VALUES (?, ?, ?)",
                        (key, question, i_type)
                    )
        await db.commit()
    except Exception as e:
        print(f"Migration error (library_steps additions): {e}")

    try:
        
        async with db.execute("SELECT id FROM library_steps WHERE step_key=?", ("info_lang",)) as cur:
            row = await cur.fetchone()
        if row:
            lang_step_id = row[0]
            async with db.execute("SELECT COUNT(*) FROM library_step_options WHERE step_id=?", (lang_step_id,)) as cur:
                if (await cur.fetchone())[0] == 0:
                    langs = [("Русский", "lang_ru"), ("English", "lang_en"), ("Tiếng Việt", "lang_vi")]
                    for idx, (t, v) in enumerate(langs, 1):
                        await db.execute(
                            "INSERT INTO library_step_options (step_id, option_text, option_value, order_index) VALUES (?, ?, ?, ?)",
                            (lang_step_id, t, v, idx)
                        )
                    await db.commit()

        
        async with db.execute("SELECT id FROM library_steps WHERE step_key=?", ("holiday",)) as cur:
            row = await cur.fetchone()
        if row:
            h_step_id = row[0]
            async with db.execute("SELECT COUNT(*) FROM library_step_options WHERE step_id=?", (h_step_id,)) as cur:
                if (await cur.fetchone())[0] == 0:
                    hols = [
                        ("Новый год", "newyear"), ("Рождество", "christmas"), 
                        ("День рождения", "birthday"), ("8 марта", "mar8"), 
                        ("Свадьба", "wedding"), ("Пропустить", "skip")
                    ]
                    for idx, (t, v) in enumerate(hols, 1):
                        await db.execute(
                            "INSERT INTO library_step_options (step_id, option_text, option_value, order_index) VALUES (?, ?, ?, ?)",
                            (h_step_id, t, v, idx)
                        )
                    await db.commit()

        
        async with db.execute("SELECT id FROM library_steps WHERE step_key=?", ("has_person",)) as cur:
            row = await cur.fetchone()
        if row:
            p_step_id = row[0]
            async with db.execute("SELECT COUNT(*) FROM library_step_options WHERE step_id=?", (p_step_id,)) as cur:
                if (await cur.fetchone())[0] == 0:
                    btns = [("Да", "person_yes"), ("Нет", "person_no")]
                    for idx, (t, v) in enumerate(btns, 1):
                        await db.execute(
                            "INSERT INTO library_step_options (step_id, option_text, option_value, order_index) VALUES (?, ?, ?, ?)",
                            (p_step_id, t, v, idx)
                        )
                    await db.commit()

        
        async with db.execute("SELECT id FROM library_steps WHERE step_key='rand_location_indoor'") as cur:
            row = await cur.fetchone()
        if row:
            step_id = row[0]
            async with db.execute("SELECT COUNT(*) FROM library_step_options WHERE step_id=?", (step_id,)) as cur:
                if (await cur.fetchone())[0] == 0:
                    btns = [
                        ("Фотостудия", "photo_studio"), 
                        ("В комнате", "room"), 
                        ("В ресторане", "restaurant"), 
                        ("В гостинице", "hotel"), 
                        ("В торговом центре", "mall"), 
                        ("Свой вариант", "custom")
                    ]
                    for idx, (t, v) in enumerate(btns, 1):
                        prompt = "Введите ваш вариант помещения (до 100 симв):" if v == "custom" else None
                        await db.execute(
                            "INSERT INTO library_step_options (step_id, option_text, option_value, order_index, custom_prompt) VALUES (?, ?, ?, ?, ?)",
                            (step_id, t, v, idx, prompt)
                        )
                    await db.commit()

        
        async with db.execute("SELECT id FROM library_steps WHERE step_key='rand_location_outdoor'") as cur:
            row = await cur.fetchone()
        if row:
            step_id = row[0]
            async with db.execute("SELECT COUNT(*) FROM library_step_options WHERE step_id=?", (step_id,)) as cur:
                if (await cur.fetchone())[0] == 0:
                    btns = [
                        ("У машины", "car"), 
                        ("У кофейни", "cafe"), 
                        ("У стены", "wall"), 
                        ("У здания", "building"), 
                        ("Москва сити", "moscow_city"), 
                        ("В лесу", "forest"), 
                        ("В горах", "mountains"), 
                        ("На аллее", "alley"), 
                        ("В парке", "park"), 
                        ("В городе", "city"), 
                        ("Свой вариант", "custom")
                    ]
                    for idx, (t, v) in enumerate(btns, 1):
                        prompt = "Введите ваш вариант локации на улице (до 100 симв):" if v == "custom" else None
                        await db.execute(
                            "INSERT INTO library_step_options (step_id, option_text, option_value, order_index, custom_prompt) VALUES (?, ?, ?, ?, ?)",
                            (step_id, t, v, idx, prompt)
                        )
                    await db.commit()
    except Exception as e:
        print(f"Migration error (library_step_options defaults): {e}")

    try:
        
        async with db.execute("SELECT id FROM library_steps WHERE step_key=?", ("aspect",)) as cur:
            row = await cur.fetchone()
        if row:
            aspect_step_id = row[0]
            async with db.execute("SELECT COUNT(*) FROM library_step_options WHERE step_id=?", (aspect_step_id,)) as cur:
                count = (await cur.fetchone())[0]
            if count == 0:
                format_buttons = [
                    ("1:1", "1:1", None),
                    ("9:16", "9:16", None),
                    ("16:9", "16:9", None),
                    ("3:4", "3:4", None),
                    ("4:3", "4:3", None),
                    ("3:2", "3:2", None),
                    ("2:3", "2:3", None),
                    ("5:4", "5:4", None),
                    ("4:5", "4:5", None),
                    ("21:9", "21:9", None),
                ]
                for idx, (text, value, prompt) in enumerate(format_buttons, start=1):
                    await db.execute(
                        "INSERT INTO library_step_options (step_id, option_text, option_value, order_index, custom_prompt) VALUES (?, ?, ?, ?, ?)",
                        (aspect_step_id, text, value, idx, prompt)
                    )
                await db.commit()
    except Exception as e:
        print(f"Migration error (library_steps.aspect buttons): {e}")

    try:
        
        async with db.execute("SELECT id FROM categories WHERE key=?", ("presets",)) as cur:
            cat_row = await cur.fetchone()
        if cat_row:
            presets_id = cat_row[0]
            async with db.execute(
                "SELECT id FROM steps WHERE category_id=? AND step_key=?",
                (presets_id, "model_select")
            ) as cur:
                exists = await cur.fetchone()
            if not exists:
                await db.execute(
                    "UPDATE steps SET order_index = order_index + 1 WHERE category_id=? AND order_index >= 13",
                    (presets_id,)
                )
                await db.execute(
                    "INSERT INTO steps (category_id, step_key, question_text, input_type, is_optional, order_index) VALUES (?, ?, ?, ?, 0, 13)",
                    (presets_id, "model_select", "💃 Выберите модель (пресет):", "model_select")
                )
                await db.commit()
    except Exception as e:
        print(f"Migration error (presets.model_select): {e}")

    async with db.execute("PRAGMA table_info(subscription_plans)") as cur:
        p_cols = [row[1] for row in await cur.fetchall()]
    
    for lang_code in ["ru", "en", "vi"]:
        col_name = f"description_{lang_code}"
        if p_cols and col_name not in p_cols:
            try:
                await db.execute(f"ALTER TABLE subscription_plans ADD COLUMN {col_name} TEXT")
                await db.commit()
                print(f"Migration: Added {col_name} to subscription_plans")
            except Exception as e:
                print(f"Migration error (subscription_plans.{col_name}): {e}")

    
    try:
        async with db.execute("PRAGMA table_info(support_messages)") as cur:
            support_cols = [row[1] for row in await cur.fetchall()]
        
        if not support_cols:
            await db.execute("""
            CREATE TABLE IF NOT EXISTS support_messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                message_text TEXT,
                file_id TEXT,
                file_type TEXT DEFAULT 'text',
                is_admin INTEGER NOT NULL DEFAULT 0,
                is_read INTEGER NOT NULL DEFAULT 0,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(user_id) REFERENCES users(id)
            )
            """)
        else:
            if "file_id" not in support_cols:
                await db.execute("ALTER TABLE support_messages ADD COLUMN file_id TEXT")
            if "file_type" not in support_cols:
                await db.execute("ALTER TABLE support_messages ADD COLUMN file_type TEXT DEFAULT 'text'")
        
        await db.commit()
    except Exception as e:
        print(f"Migration error (support_messages): {e}")

    
    try:
        
        def has_old_tags(text: str) -> bool:
            if not text: return True
            return any(x in text for x in ["(ТУТ УКАЗЫВАЕМ", "УКАЗЫВАЕМ ПОЛ", "НАГРУЖЕННОСТЬ", "ПРИМУЩЕСТВО", "УКАЗЫВАЕМ  УКАЗЫВАЕМ"])

        
        async with db.execute("SELECT value FROM app_settings WHERE key='infographic_clothing_prompt'") as cur:
            row = await cur.fetchone()
            if not row or has_old_tags(row[0]):
                text = r'''ROLE / TASK
Generate ONE modern high-fashion advertising poster in a magazine cover style,
derived from the uploaded product photo.
for a marketplace using ONLY the uploaded product photo.

Poster must be:
– stylish
– creative
– informative
– commercially strong

PRODUCT FIDELITY ONLY (visual accuracy of the clothing item must be exact)
All design, typography, layout and decorative elements are free, modern,
fashion-forward and may be highly expressive.

COLOR & SHADE LOCK
Reproduce the product color EXACTLY as in the photo:
– same hue, tone, temperature, saturation
– no recolor, no tint, no warming/cooling
– no harmonization with background or model
All small elements (belt, cuffs, quilting zones) must match the exact color.
Lighting changes must not shift product color.

NO PHOTO INSERTION
Do not insert or clip the original photo.
Render a completely new editorial-style image.

ADVANCED TYPOGRAPHY RANDOMIZATION (CRITICAL)

The poster must use a RANDOM modern, stylish, visually expressive font
that matches the product category and aesthetic.

Allowed styles:
– geometric modern sans
– elegant fashion serif
– high-contrast editorial serif
– minimalist grotesk
– soft rounded modern
– stylish condensed typefaces
– expressive display fonts (if still readable)

The font must always look:
– premium
– modern
– aesthetically strong
– visually distinctive
– not similar to default system fonts

Forbidden:
– generic system fonts (Arial, Roboto, Helvetica-like defaults)
– plain or basic typographic styles
– overly neutral, dull or standard typefaces
– repeated typographic look across outputs

The system must randomize typography selection every time, choosing
a visually appealing, fashionable, marketplace-ready typeface.

TYPOGRAPHY FREEDOM MODE

The system must select a unique font style for every output, fully independent
from previous generations.

Allowed:
– expressive display fonts
– fashion editorial serifs
– handwritten-like stylish fonts (if readable)
– geometric, bold or condensed styles
– decorative modern typography

The style may vary freely:
– weight
– contrast
– serif/sans serif
– proportions
– curves
– terminals

No limitation on baseline shape, stroke contrast or typographic geometry,
as long as the text remains readable.

LANGUAGE CONTROL (CRITICAL)
Infographic Language: {Язык инфографики}
ALL text must be written ONLY in this language.
Fix grammar automatically.
Professional marketplace copywriting.

All text must be rendered with ZERO AI artifacts.

The typography must be:
– perfectly smooth and clean
– free from distortions, glitches and warped geometry
– free from broken, duplicated, or melted characters
– free from random symbols or hallucinated glyphs
– free from uneven thickness, inconsistent stroke weight or pixel noise
– free from jagged or stair-stepped edges

Letters must always appear:
– with clean kerning and spacing

Forbidden:
– fragmented or half-rendered glyphs
– “AI handwriting” effect
– any visual signs of neural network generation

VISIBLE = REAL, INVISIBLE = NON-EXISTENT
Only features visible in the product photo may exist in the output.
No invented hardware, seams, textures, structure or lining.

TEXTURE & CONSTRUCTION RULE
Textures and patterns appear only where visible.
Construction, panels and quilting must match exactly.

NO CUT–PASTE
Render the product from scratch while preserving perfect fidelity.

IMAGE QUALITY
– high detail
– sharp
– clean
– realistic texture
– no blur or plastic effects

STYLE MODE BY CATEGORY
Product gender category: {Пол модели}

If WOMEN: modern editorial, elegant, trendy.
If MEN: строгий, геометричный, контрастный.
If KIDS: мягкие формы, дружелюбные элементы.

COLOR SYSTEM
Use 1–3 main colors plus neutrals.

TEXT CONTRAST
Text must remain readable in all cases (light text on dark, dark text on light).

INFOGRAPHIC COMPLEXITY LEVEL: {Нагруженность инфографики}

Level 1–3:
– 1 headline
– 2 advantages
– 1–2 icons
– 0 insets

Level 4–6:
– 1 headline
– 3 advantages
– 3–5 icons
– 1 inset
– 1 micro-copy

Level 7–8:
– 1 headline
– 3 advantages
– 5–7 icons
– 2 insets
– 2–3 micro-copies
– 1 extra block

Level 9–10:
– 1 headline + subheadline
– 3 advantages + 1–2 extra claims
– 7–10 icons
– 3 insets
– 3–5 micro-copies
– rich layout (side columns, badges)

INSET RULE
Inset panels must show only visible real details:
– seams
– texture
– quilting
– cuffs
– collar
– material surface
If insufficient unique zones exist, reuse a visible detail with different zoom.

INSET FIDELITY LOCK (CRITICAL)

Inset panels must reproduce the selected detail EXACTLY as it appears
in the uploaded product photo, with zero modification of:

– shape
– geometry
– fold direction
– curvature
– thickness
– stitching pattern
– quilting rhythm
– collar/hood construction
– material volume and texture
– fabric tension and compression

Inset details must look like a magnified crop of the REAL visible area from the
uploaded photo, NOT a cleaner, smoother, more symmetrical or more standard version.

INSET ANGLE LOCK (ABSOLUTE)

Inset panels must always use the EXACT SAME viewing angle and geometry
as the original product photo.

Forbidden for insets:
– rotating the detail to a different orientation
– showing the detail from a “cleaner” or “more frontal” angle
– changing perspective or projection
– straightening or centering the detail
– reconstructing hidden or obscured sides of the detail
– showing the collar, belt, cuff or quilting from an angle that is NOT visible in the photo

Insets must appear as if they are a magnified crop of the SAME camera angle,
SAME perspective, SAME lighting direction, and SAME geometry as the original image.

If the photo shows the detail partially, cut off, or at a shallow angle → the inset must
repeat that same partial shape without reconstruction or correction.

If the angle of a detail cannot be extracted without guessing → the inset MUST NOT be generated.

STRICT INSET SOURCE RULE (ABSOLUTE)

Inset panels must ALWAYS be a faithful magnified recreation ONLY of the EXACT
visible fragment from the uploaded product photo.

The inset must show:
– the same shape
– the same geometry
– the same folds and curves
– the same stitching and quilting
– the same texture
– the same lighting behavior
– the same viewing angle
– the same partial visibility if the detail is cropped

Forbidden for insets:
– inventing or adding any missing parts
– rotating the detail to another angle
– idealizing or beautifying the construction
– reconstructing hidden areas
– correcting asymmetry
– smoothing folds or straightening lines
– “guessing” how the detail should look
– using category knowledge to complete the design

If the exact detail cannot be reproduced WITHOUT INVENTION,
the inset MUST NOT be generated.

STRICTLY FORBIDDEN IN INSETS:
– improving, beautifying or idealizing the detail
– correcting asymmetry or irregularities
– smoothing or refinishing the geometry
– inventing missing parts
– completing unclear edges
– rounding or sharpening shapes
– using category-typical details when visibility is low

If the visible part of a detail is partially obscured, cropped or cut off in the
photo → the inset MUST show it with the SAME partial visibility.
No reconstruction or completion of the hidden zone.

If the model cannot identify a unique detail clearly → the inset must be skipped
or replaced with a zoom of another clearly visible area.

LAYOUT

VISUAL DESIGN RANDOMIZATION RULE (CRITICAL)

The infographic design must vary each time within a modern editorial style.

The system must randomly choose:
– icon style (outline / filled / geometric / linear / minimalistic / artistic)
– block styling (rounded cards, rectangles, shadows, strokes, layered panels)
– connector lines (straight / curved / dotted / geometric)
– accents (shapes, abstract forms, minimal patterns)
– micro-decor elements (subtle backgrounds, geometric fragments)

Allowed:
– stylish modern icons
– premium line-art
– soft geometric abstract forms
– editorial minimalism
– fashion-style micro accents

Forbidden:
– default generic icons
– basic pack icons
– repeated icon style across outputs
– outdated or overly simplistic visuals

All variations must remain coherent, premium and marketplace-ready.

ADVANCED ART DIRECTION RANDOMIZATION

The visual style of the infographic must vary every time within
the boundaries of modern fashion graphics.

Allowed variations:
– minimalistic, clean layouts
– bold fashion magazine compositions
– collage-style layering
– abstract geometric shapes
– soft pastel backgrounds
– high-contrast editorial blocks
– neon or vibrant accents
– smooth gradients
– artistic brush elements
– cutout-style frames
– modern big-typography compositions

The system must select a coherent but unique combination of:
– layout geometry
– card shapes
– panel layering
– decorative elements
– text placement logic
– accents
– visual hierarchy

Design must always remain:
– modern
– premium
– fashion-oriented
– marketplace-ready

Forbidden:
– template-like repetitive structure
– identical layout patterns across outputs
– standard “neutral” or “corporate” infographic styles

MICRO COPY
Short, relevant, and derived strictly from real visible characteristics.

PRODUCT INFO
Product / Brand Name: {Название товара}
Top 3 advantages: 
1. {Преимущество 1}
2. {Преимущество 2}
3. {Преимущество 3}

Дополнительная информация о продукте:
{Доп информация}
→ use as extra block when complexity >= 6

MODEL LOGIC (ONLY WHEN PRODUCT NEEDS A MODEL)
If product is одежда → model allowed.
If product is обувь / аксессуар / small item → close-up product priority.

MODEL USE GATE (CRITICAL)
Use a model ONLY if the garment can be reproduced on the model with 100% construction fidelity.
If there is any risk of invented elements → render the product without a model.


GARMENT COMPLETION RULE
If the uploaded product is:
– outerwear or tops → model must wear pants
– dress → model must not wear pants
– pants → model must wear a top that does not match product color and does not create a set
– footwear → model must wear neutral pants and neutral top
Non-product clothing must never match product color.

RANDOM MODEL APPEARANCE RULE

The model’s appearance must be generated randomly within the allowed range
for each output, without repeating a fixed look.

Allowed:
– European, Caucasian, Slavic, Mediterranean, Middle Eastern, Latin types
– natural variations of jawline, nose shape, lips, eyes, eyebrows
– natural variation of hairstyle, hair length, color (except unnatural colors)

Forbidden:
– Asian appearance
– African or Afro-American appearance
– unrealistic or stylized facial proportions
– cartoon-like or AI-distorted faces

The appearance must stay natural, realistic and human.
Hair, facial features and styling must vary freely, without pattern repetition.

MODEL ORIENTATION: {Угол камеры}

Если Спереди:
– фронтальный вид
– допускается лёгкий наклон корпуса или головы

Если Сзади:
– вид со спины
– допускается лёгкий поворот головы

MODEL POSE
{Поза}

Обычный:
– нейтральная коммерческая стойка
– спокойное положение тела
– руки в естественном положении

Вульгарный:
– выразительная, динамичная, провокационная поза
– смелая, эффектная, но реалистичная

Нестандартный:
– момент движения
– лёгкий шаг, поворот, наклон
– живая натуральная динамика

Любая поза не должна деформировать изделие или скрывать конструкцию.

CAMERA & FRAMING
Camera angle: {ракурс фотографии}

Дальний:
– модель в полный рост

Средний:
– от головы до середины бёдер
– акцент на пропорции товара

Близкий:
– пояс–голова или крупный план изделия
– максимальный акцент на товар

Без смешения планов и перспективных искажений.

MODEL PARAMETERS
SERVICE DATA VISIBILITY LOCK (ABSOLUTE)

The following parameters are service information ONLY and must NEVER appear
as text or labels on the final poster:

– model height
– model age
– model gender
– model body size
– garment parameters (length, sleeve type, pants cut)
– camera angle
– pose name
– orientation (front/back)
– any numeric or descriptive technical inputs

These parameters are strictly for internal AI logic.
Only the product name, brand name and product advantages may appear on the poster.
No other info may be shown.

Body size: {Размер модели}

BODY SIZE FIDELITY RULE:

Body size defines real physical body volume, mass and softness — NOT styling,
posture, fitness, angle or artistic interpretation.

Each size must show clear, realistic body differences:

42–44 → very slim
46–48 → slim
50–52 → slim-curvy with visible softness, NO athletic appearance
54–56 → curvy with clear belly volume and wide waist
58–60 → heavy with substantial mass
60+ → very large, massive body

Strictly forbidden for sizes 50+:
– flat abdomen
– tight abdomen
– athletic or fit body definition
– visible toned muscles
– narrow waist
– “plus-size but slim” interpretations

Pose, posture, camera angle, lighting and clothing must NOT hide, slim or compress
body volume. Body size ALWAYS overrides pose, styling, age and camera angle.

Height: {Рост модели}
Age: {Возраст модели}
Gender: {Пол модели}

GARMENT PARAMETERS
Product length: {Длина изделия}
Sleeve type: {Тип рукава}
Pants cut: {Тип кроя штанов}
If empty → use only what is visible.

CATEGORY DETECTION
Determine product category only from visible construction, without guessing.

NO SHAPE MODIFICATION
Garment must keep original fullness, silhouette, quilting, collar geometry, sleeve width and belt structure.
Model pose must not distort the garment.

FALLBACK RULE
If a detail cannot be confidently detected → omit it.
Never invent or assume.

FINAL QUALITY CHECK
– perfect text contrast
– correct layout per complexity
– insets only from visible zones
– modern icons and callouts
– 100% product fidelity
– exact color match
– zero invented details
– model neutral and product unobstructed'''
                await db.execute("INSERT INTO app_settings (key, value) VALUES ('infographic_clothing_prompt', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (text.strip(),))
                print("Migration: Updated infographic_clothing_prompt")

        
        async with db.execute("SELECT value FROM app_settings WHERE key='random_other_prompt'") as cur:
            row = await cur.fetchone()
            if not row or has_old_tags(row[0]):
                text = r'''You are a professional product stylist and commercial photographer.
Analyze the uploaded image and identify the main product only.
Treat the product as a finished commercial item with a final, approved design.
Product Integrity (STRICT)
The product must remain 100% unchanged.
Do NOT:
rotate, flip, mirror, or invert the product
change orientation (front/back/side)
modify, redraw, rearrange, stylize, or “improve” any graphics, text, logos, or illustrations
alter colors, proportions, or layout of the product
The product must look exactly like the original photo, as if photographed in real life.
Scale & Reality (CRITICAL)
Use the real dimensions below.
The product must appear correctly sized relative to the environment and surrounding objects.
Overscaling or unrealistic proportions are not allowed.
Color & Visual Style (IMPORTANT)
Select a color palette that complements and enhances the product.
Rules:
Base the palette on the product’s colors, materials, and mood
Avoid default neutral-only palettes (plain white, beige, gray) unless clearly justified
Use modern, tasteful, and creative color combinations
Background and props may contrast with the product, but must not overpower it
Colors should feel intentional, contemporary, and visually rich
Safe, boring color choices are discouraged
Controlled creativity is required.
Scene & Complexity Control (VERY IMPORTANT)
Build the scene around the product.
Scene complexity level (1–10) is a hard limit, not a suggestion:
1–2 → almost empty, studio-like, max 1–2 simple props
3–4 → minimal lifestyle, few elements, lots of negative space
5–6 → balanced lifestyle scene
7–8 → rich but controlled environment
9–10 → dense, detailed lifestyle scene
Do NOT exceed the specified complexity level.
Low values must stay visibly simple.
The product must remain the clear focal point.
Product Details
Product type: {Название товара}
Photo angle / shot type: {Угол камеры}
Scene complexity level (1–10): {Нагруженность}
Parameters:
Width: {Ширина}
Height: {Высота}
Length: {Длина}
Season: {Сезон}
Style: {Стиль}


Human Presence (NEW – REQUIRED)
Human presence: {Присутствует ли человек на фото}
If Human presence is Yes:
Add one person interacting naturally with the product shown in the uploaded image
The interaction must be realistic, functional, and relevant to the product
The person must not distract from the product and must support its use or context
Add the following required field:
Gender: {Пол модели}
The model must be well-dressed, stylish, and visually appropriate to:
the product type
the season
the scene complexity
the overall visual mood
Natural Face Rendering (CRITICAL)
If a human is present in the scene, the face must look natural and realistic.
Rules:
No heavy beauty retouching
No “plastic”, “oily”, or overly smooth skin
Skin texture must be visible and realistic
Light natural wrinkles, fine lines, and pores are allowed and encouraged
Natural facial asymmetry is acceptable
Realistic skin imperfections are allowed
Avoid artificial beauty filters or CGI-like skin
The face should look like a real person photographed with professional lighting, not digitally polished.
If Human presence is No:
Do not add any people to the scene
Skip the Gender field entirely
Output Requirements
Ultra-realistic
Professional commercial photography
Correct perspective and lighting
No added text or watermarks'''
                await db.execute("INSERT INTO app_settings (key, value) VALUES ('random_other_prompt', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (text.strip(),))
                print("Migration: Updated random_other_prompt")

        
        async with db.execute("SELECT value FROM app_settings WHERE key='random_prompt'") as cur:
            row = await cur.fetchone()
            if not row or has_old_tags(row[0]):
                text = r'''You are a professional commercial fashion imaging AI. Your task is to take the uploaded product photo as the ONLY source of truth and present the product on a realistic human model strictly for sale purposes.

ABSOLUTE PRODUCT FIDELITY (CRITICAL):
The uploaded product photo is the only reference. The product must be reproduced with 100% accuracy including exact shape, silhouette, proportions, cut, construction, seams, stitching, materials, fabric texture, thickness impression, exact color and shade. It is strictly forbidden to redesign, stylize, improve, simplify, reinterpret, add or remove elements or invent details not visible on the photo. If a detail, material or texture is not clearly visible on the photo, it must be treated as non-existent.

FABRIC & TEXTURE DISTRIBUTION LOCK (ABSOLUTE):
Fabric type, knit structure and texture MUST be reproduced exactly as shown on the product photo AND ONLY in the areas where they are visible. If a specific texture (e.g. ribbed, knitted, “lapsha”, structured knit, smooth knit, woven fabric) appears only in certain zones of the product (such as neckline, collar, cuffs, hem or panels), it MUST remain strictly limited to those zones. It is strictly forbidden to extend, repeat, generalize or apply a visible texture or knit pattern to other parts of the product where it is not present on the photo. Any global material unification is a critical failure.

STRUCTURAL CONSTRUCTION LOCK (ABSOLUTE):
Any structural construction elements that exist on the product (including but not limited to neckline, collar, neck opening, cuffs, hems, closures, zippers, tongues, waistlines or openings) MUST be reproduced exactly as shown on the uploaded product photo. The AI is strictly forbidden to modify, reinterpret, stylize, adjust or “improve” the shape, height, width, depth, curvature, openness, fit or construction logic of any existing structural element. If a structural element is not clearly visible on the product photo, it must be treated as non-existent and must not be inferred.

GARMENT SURFACE CONDITION RULE (ABSOLUTE):
If the product on the photo shows wrinkles, creases or folds caused by storage or hanging, the product MUST be presented neatly steamed and smooth on the model. Removing wrinkles is allowed ONLY as surface smoothing. It is strictly forbidden to change fabric behavior, stiffness, thickness, drape, elasticity, knit structure or material appearance while smoothing. No reshaping, tightening or texture alteration is allowed.

PRODUCT COLOR LOCK (ABSOLUTE):
The AI must determine the exact product color exclusively from the uploaded photo. Color must match base color, undertone, saturation and brightness exactly. No recoloring, enhancement, lighting reinterpretation or stylistic shift is allowed. Lighting must not alter perceived color.

IMAGE RESOLUTION, SHARPNESS & OPTICAL QUALITY (CRITICAL)

The generated image must simulate a true high-end commercial fashion photograph captured in studio-grade Ultra High Definition (true 4K).

MANDATORY TECHNICAL REQUIREMENTS:
– ultra-high resolution look with true 4K-equivalent pixel density
– extreme micro-detail clarity on fabric, seams, stitching, edges and construction lines
– tack-sharp focus across the entire visible product area
– no softness, haze, blur, diffusion, glow or cinematic softness
– no motion blur, depth blur, lens blur or artificial bokeh on the product
– no painterly, illustrative or AI-smoothed textures

OPTICAL & SENSOR REALISM RULES:
– simulate professional commercial camera optics and sensor behavior (high-end fashion photography setup)
– realistic depth of field, but the product must remain 100% in sharp focus at all times
– natural micro-contrast and edge acuity without digital oversharpening
– edge definition must be crisp, clean and physically precise
– fabric texture must be clearly readable at close inspection
– stitching, seams and material transitions must be visibly distinct and realistic

ANTI-AI SMOOTHING & ANTI-RENDER LOCK:
– strictly forbid any neural smoothing, beauty smoothing or texture averaging
– strictly forbid plastic, waxy, rubbery or over-processed surfaces
– forbid CGI-like, render-like or synthetic surface appearance
– preserve natural fabric micro-structure and physical irregularities without exaggeration

COMPRESSION & OUTPUT QUALITY:
– image must look uncompressed and lossless
– no JPEG artifacts, noise blobs, aliasing or texture breakup
– clean tonal transitions without banding or posterization

PRIORITY RULE:
Image sharpness, resolution, texture readability and physical realism must NEVER be reduced for style, mood, lighting effects or artistic interpretation.

STRUCTURAL CONSTRUCTION LOCK (ABSOLUTE):
Any structural construction elements that exist on the product (including but not limited to neckline, collar, neck opening, cuffs, hems, closures, zippers, tongues, collars, waistlines or openings) MUST be reproduced exactly as shown on the uploaded product photo. The AI is strictly forbidden to modify, reinterpret, stylize, adjust or “improve” the shape, height, width, depth, curvature, openness, fit or construction logic of any existing structural element. If a structural element is not clearly visible on the product photo, it must be treated as non-existent and must not be inferred. Any deviation from original structural construction is a critical failure.

Location type: {Тип локации}
Rules: If На улице → image must be outdoors. If В помещении → image must be indoors. Mixing is forbidden.

2.LOCATION DESIGN DIVERSITY & STYLE RANDOMIZATION (CRITICAL): {Стиль локации}

If the Location style / subtype is a short generic phrase (length ≤ 18 characters, e.g. у кофейни, на улице, в помещении, у машины), the AI MUST generate a visually strong, contemporary environment aligned with global fashion, architecture and visual design standards of 2025.

The design language MUST be modern, bold and editorial, not cozy or lifestyle-oriented.

FORBIDDEN STYLES & MOODS:
– warm cozy interiors
– beige / cream / soft brown dominant palettes
– “comfortable”, “homey”, “vintage”, “loft lifestyle” aesthetics
– rustic, retro, classic, Scandinavian hygge
– warm ambient lighting as dominant mood

MANDATORY VISUAL CHARACTERISTICS:
– contemporary materials (glass, metal, concrete, stone, technical surfaces)
– clear geometry and spatial contrast
– neutral-to-cool or mixed temperature lighting with directional light
– fashion-editorial or architectural composition
– sense of modernity, clarity and visual tension

The AI MUST vary visual direction across generations using modern design approaches such as contemporary minimalism, architectural modern, urban editorial, experimental space, brutal-modern, high-fashion showroom or cinematic contemporary environments.

Repetition of the same color palette, lighting mood or spatial logic across generations is strictly forbidden.

Environmental lighting and design must never cast color reflections onto the product or alter perceived product color.

If the location refers to a real existing place, the AI MUST follow its real-world architectural identity regardless of text length.

Season: {Сезон}
Applied ONLY if Location type = На улице. Зима → visible snow; Лето → clear weather and sunlight; Осень → dry autumn leaves; Весна → fresh spring atmosphere. If В помещении → ignore completely.

Holiday: {Праздник}
Apply only a light, subtle holiday hint without heavy decorations or product changes. If empty → ignore.

Model’s pose: {Поза модели}
Обычный → natural commercial pose. Вульгарный → openly provocative pose. Нестандартная → unstable or transitional pose. Pose must never distort, hide or deform the product.

Body size of the model: {Размер модели}
BODY SIZE GEOMETRY RULE (ABSOLUTE): Body size defines real physical body volume, mass and softness, not styling, posture or fitness. Each higher size must be visibly larger and heavier than the previous one.
42–44 very slim; 46–48 slim; 50–52 slim-curvy with visible softness and NO athletic appearance; 54–56 curvy with clear belly volume and wide waist; 58–60 heavy with substantial mass; 60+ very large massive body.
FORBIDDEN FOR SIZES 50 AND ABOVE: flat or tight abdomen, athletic or fit body, toned muscles, narrow waist, “plus-size but slim” interpretation.
Pose, posture, camera angle or clothing must not slim or hide body volume. Body size always overrides age, pose, styling and camera angle.

Model height: {Рост модели}
Use exactly as specified.

Model’s age: {Возраст модели}
50+ → visible age signs and natural wrinkles; 65+ → pronounced aging. Any rejuvenation, blur or beauty filters are forbidden.

Pants cut type: {Тип кроя штанов}
If empty → choose a cut that naturally fits the product.

Sleeve type: {Тип рукава}
If empty → replicate sleeves exactly as on the product photo. If footwear → ignore.

Product length: {Длина изделия}
If specified → follow precisely. If empty → determine strictly from the product photo.

Model orientation: {Угол камеры}
Спереди → front view with slight tilt allowed. Сзади → back view with head turn allowed.

Camera angle: {ракурс фотографии}
CAMERA FRAMING LOCK (ABSOLUTE):
Дальний → full body head to feet only.
Средний → head to mid-thigh ONLY; legs below mid-thigh must not be visible.
Близкий → waist to head ONLY; hips and legs must not be visible.
Product category rules must NOT expand framing beyond the specified camera angle.

Gender: {Пол модели}
Specify explicitly.

PRODUCT CATEGORY LOGIC (CRITICAL):
Upper garments require suitable bottoms that do not hide the product. Dresses must never have pants. Long coats or robes must not show pants. Pants as product require a top that does not cover waist or cut. Footwear requires suitable complementary clothing.

MODEL APPEARANCE AND RANDOMIZATION (CRITICAL):

Randomization is mandatory and should vary depending on certain categories of appearance.

FORBIDDEN BY APPEARANCE:
Asian, African, and African-American types are strictly prohibited.

Any kind of hairstyle and hair styling is allowed!

SKIN REALISM LOCK (ABSOLUTE):
Skin must be natural and physically realistic. Glossy, oily, plastic, porcelain or beauty-retouched skin is forbidden. Visible pores and natural texture are required.

MODEL REUSE PREVENTION (FAIL-SAFE):
If face, hairstyle, hair color or overall appearance is visually similar to a previous generation, the result is INVALID and must be regenerated.

FINAL GOAL:
Produce a realistic, professional, sale-ready image where commercial presentation NEVER overrides product construction, proportions, fit, structure or color accuracy.'''
                await db.execute("INSERT INTO app_settings (key, value) VALUES ('random_prompt', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (text.strip(),))
                print("Migration: Updated random_prompt")

        
        async with db.execute("SELECT value FROM app_settings WHERE key='whitebg_prompt'") as cur:
            row = await cur.fetchone()
            if not row or has_old_tags(row[0]):
                text = r'''ROLE / SYSTEM TASK

You are a professional commercial product visualization AI used for marketplaces, brand catalogs, and e-commerce vitrines.

Your task is to analyze the uploaded real product photo and generate a photorealistic 3D-render-style product image on a pure white background, visually indistinguishable from a real studio photograph.

The result must look like a high-end retail vitrine product photo, NOT like obvious CGI or synthetic 3D.

SOURCE OF TRUTH (CRITICAL)

The uploaded product photo is the ONLY and ABSOLUTE SOURCE OF TRUTH.

You must study the product carefully.

You must analyze only the product itself, NOT the background.

Color(s), material, shape, proportions, construction, seams, edges, thickness, texture, and all physical details must be derived strictly from the product visible in the photo.

🚫 DO NOT copy-paste, overlay, or reuse the original photo in any form.
🚫 DO NOT place the original photo on top of a background.
🚫 DO NOT partially reuse pixels, silhouettes, shadows, or cutouts from the source image.

The final image must be a newly generated, reconstructed product, not a manipulation of the original image.

PRODUCT RECONSTRUCTION RULES (HIGHEST PRIORITY)

Recreate the product with 100% structural fidelity:

Exact silhouette and proportions

Exact cut and geometry

Exact number and placement of seams, stitches, panels

Exact edges, folds, thickness, and volume

Exact material appearance (fabric, leather, plastic, metal, etc.)

Exact color(s) of the product itself (ignore background lighting)

🚫 FORBIDDEN ACTIONS

Adding buttons, zippers, pockets, logos, straps, patterns, prints, textures, or any elements not present in the original photo

Removing elements that exist in the original photo

“Improving”, “enhancing”, “stylizing”, or “modernizing” the design

Guessing hidden details that are not visible

Symmetry correction if the product is asymmetrical in the photo

If a detail is not clearly visible, keep it neutral and minimal, never invented.

COLOR & MATERIAL INTELLIGENCE

Analyze only the product’s real color(s), not reflections from background or lighting.

Do not change hue, saturation, brightness, or tone.

Do not “beautify” colors.

Material realism is mandatory: fabric grain, matte/gloss level, softness, rigidity must match reality.

WRINKLES & CONDITION NORMALIZATION

If the product in the photo shows:

Wrinkles

Creases

Folding marks

Signs of being un-ironed or deformed by storage

You must render the product as:

Perfectly straightened

Neatly arranged

Ironed / smoothed

Retail-ready condition

⚠️ This applies ONLY to surface condition.
⚠️ Structural shape, cut, and proportions must remain unchanged.

PHOTOREALISTIC 3D STYLE (CRITICAL)

The output must:

Look like a real studio product photograph

Be indistinguishable from a DSLR or medium-format camera shot

Avoid visible CGI traits (plastic look, fake highlights, unnatural edges)

Lighting:

Soft, neutral studio lighting

Natural shadows under the product

No dramatic, artistic, or cinematic lighting

Background:

Pure white (#FFFFFF)

No gradients

No textures

No reflections

No environment

Camera:

Neutral focal length (no distortion)

Product centered

Clean, commercial framing

ABSOLUTE PROHIBITIONS

🚫 Do NOT:

Copy the original photo

Paste the original product image onto a background

Add branding, text, labels, watermarks

Add props, models, hands, stands (unless visible in original photo)

Change product orientation unless clearly required for vitrine presentation

Add “AI enhancements” or fictional improvements

QUALITY CONTROL CHECK (MANDATORY INTERNAL STEP)

Before finalizing the image, verify internally:

Is this a newly generated product, not a reused photo?

Does every visible detail exist in the original photo?

Are there ZERO invented elements?

Does it look like a real retail product photo, not CGI?

Is the background pure white and clean?

Is the product ironed and neat, without altering structure?

If any answer is “NO” → regenerate.

FINAL OUTPUT REQUIREMENT

Produce ONE single, high-resolution, photorealistic product image suitable for:

Marketplace product card

Online store vitrine

Brand catalog

No explanations.
No captions.
No text overlays.
Image only.'''
                await db.execute("INSERT INTO app_settings (key, value) VALUES ('whitebg_prompt', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (text.strip(),))
                print("Migration: Updated whitebg_prompt")

        async with db.execute("SELECT value FROM app_settings WHERE key='collage_prompt'") as cur:
            row = await cur.fetchone()
            if not row or not (row[0] or "").strip():
                collage_text = open(os.path.join(os.path.dirname(__file__), "_collage_prompt_default.txt"), encoding="utf-8").read().strip() if os.path.isfile(os.path.join(os.path.dirname(__file__), "_collage_prompt_default.txt")) else ""
                if not collage_text:
                    collage_text = """SYSTEM / ROLE:

You are a professional AI fashion photoshoot and apparel transfer model.

Your task is to create ONE ultra realistic 4-frame collage from:
1) one reference photo of a human model,
2) product photos of a clothing or footwear item.

MAIN OBJECTIVE

Preserve the model exactly as in the original reference image: same face, same identity, same body shape, same proportions, same skin texture, same hairstyle, same age, same ethnicity, same anatomy.
The person must remain fully recognizable as the same individual.
Do NOT recreate, remember, or copy the clothing from the original model photo.

PRODUCT RECREATION

Recreate the uploaded product as a premium commercial fashion item while preserving: same category, same colors, same materials, same silhouette, same proportions, same fit, same design language, same visible branding placement, same important visual details.
The recreated product must look steamed and clean, premium, professionally styled, marketplace-ready, maintain realistic textile behavior, realistic folds and stitching.
Do NOT copy defects, wrinkles, distortions, or low-quality details from source product photos.

SMART STYLING LOGIC

The AI must automatically generate matching missing outfit elements.
If uploaded product is: jacket/top/hoodie/shirt/sweater - generate matching pants and shoes; pants/skirt/shorts - generate matching upper clothing and shoes; shoes - generate full matching outfit; full outfit/dress - only add minimal matching accessories or shoes if needed.

COLLAGE STRUCTURE

Generate EXACTLY ONE IMAGE containing EXACTLY 4 frames.
Frame 1: full front view, front-facing product presentation.
Frame 2: full back view, back-facing product presentation.
Frame 3: side profile view, side silhouette clearly visible.
Frame 4: close-up detailed shot, focus on important product details, fabric texture, stitching, footwear, accessories.

POSE RULES

The model must pose like a real professional fashion model. Dynamic posing, asymmetric posing, natural movement, premium fashion body language.
NEVER generate mannequin poses, symmetrical standing, stiff posture, arms hanging straight down, robotic body positioning.

STRICT HAND RULES

NEVER place hands inside pockets. Hands must remain fully visible, natural, elegant, fashion-oriented, dynamically positioned.

CONSISTENCY RULES

All 4 frames must show the SAME person, preserve the SAME identity, the SAME outfit, the SAME product details, the SAME colors, realistic anatomy, realistic fabric behavior.

PHOTOGRAPHY STYLE

Style: premium ecommerce photoshoot, ultra realistic, highly detailed, clean fashion campaign aesthetic.
Lighting: professional studio lighting, soft realistic shadows, balanced exposure, realistic skin rendering, realistic textile rendering.
Background: clean minimal premium studio background, no distracting objects.

IMPORTANT RESTRICTIONS

DO NOT: change face, change identity, change body type, change ethnicity, generate unrealistic anatomy, generate extra fingers, distort product proportions, invent random logos, oversaturate colors, create cartoon/anime style, sexualize the model, over-retouch the face.

FINAL RESULT

Create an ultra realistic premium fashion collage with commercial quality, marketplace-ready presentation, realistic styling, realistic textile physics, consistent identity, consistent product appearance, dynamic professional fashion posing, premium studio photography quality."""
                await db.execute("INSERT INTO app_settings (key, value) VALUES ('collage_prompt', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (collage_text.strip(),))
                print("Migration: Seeded collage_prompt")

        
        async with db.execute("SELECT id, text FROM prompts") as cur:
            rows = await cur.fetchall()
            for pid, ptext in rows:
                if ptext and "(ТУТ УКАЗЫВАЕМ" in ptext:
                    new_text = ptext.replace("(ТУТ УКАЗЫВАЕМ Пол модели)", "{Пол модели}") \
                                     .replace("(ТУТ УКАЗЫВАЕМ ШИРИНУ)", "{Ширина}") \
                                     .replace("(ТУТ УКАЗЫВАЕМ ВЫСОТУ)", "{Высота}") \
                                     .replace("(ТУТ УКАЗЫВАЕМ ДЛИНУ)", "{Длина}") \
                                     .replace("(ТУТ УКАЗЫВАЕМ СЕЗОН)", "{Сезон}") \
                                     .replace("(ТУТ УКАЗЫВАЕМ СТИЛЬ)", "{Стиль}") \
                                     .replace("(ТУТ УКАЗЫВАЕМ Присутствует ли человек на фото)", "{Присутствует ли человек на фото}") \
                                     .replace("(ТУТ УКАЗЫВАЕМ Нагруженность инфографики)", "{Нагруженность}") \
                                     .replace("(ТУТ УКАЗЫВАЕМ Угол камеры)", "{Угол камеры}") \
                                     .replace("(ТУТ УКАЗЫВАЕМ Поза модели)", "{Поза}") \
                                     .replace("(ТУТ УКАЗЫВАЕМ Рост модели)", "{Рост модели}") \
                                     .replace("(ТУТ УКАЗЫВАЕМ Возраст модели)", "{Возраст модели}") \
                                     .replace("(ТУТ УКАЗЫВАЕМ РАЗМЕР МОДЕЛИ)", "{Размер модели}") \
                                     .replace("(ТУТ УКАЗЫВАЕМ ЯЗЫК ИНФОГРАФИКИ)", "{Язык инфографики}") \
                                     .replace("(ТУТ УКАЗЫВАЕМ ПРИМУЩЕСТВО 1)", "{Преимущество 1}") \
                                     .replace("(ТУТ УКАЗЫВАЕМ ПРИМУЩЕСТВО 2)", "{Преимущество 2}") \
                                     .replace("(ТУТ УКАЗЫВАЕМ ПРИМУЩЕСТВО 3)", "{Преимущество 3}") \
                                     .replace("(ТУТ УКАЗЫВАЕМ ДОП ЧТО УГОДНО О ТОВАРЕ)", "{Доп информация}")
                    await db.execute("UPDATE prompts SET text=? WHERE id=?", (new_text, pid))
                    print(f"Migration: Cleaned up prompt ID {pid}")

        await db.commit()
    except Exception as e:
        print(f"Migration error (prompts update): {e}")


def _normalize_placeholder_label(text: str, fallback: str) -> str:
    if not text:
        return fallback
    
    
    clean = re.sub(r"[^0-9A-Za-zА-Яа-яЁё ]+", "", text).strip()
    clean = re.sub(r"\s+", " ", clean).strip()
    
    
    for prefix in ("Выберите ", "Введите ", "Пришлите ", "Загрузите "):
        if clean.lower().startswith(prefix.lower()):
            clean = clean[len(prefix):].strip()
            break
            
    low = clean.lower()
    
    
    mapping = {
        "возраст": "Возраст модели",
        "телосложение": "Телосложение",
        "размер одежды": "Телосложение",
        "рост": "Рост модели",
        "пол": "Пол модели",
        "тип позы": "Поза модели",
        "поза": "Поза модели",
        "угол камеры": "Угол камеры",
        "вид фотографии": "Угол камеры",
        "ракурс": "Ракурс",
        "сезон": "Сезон",
        "праздник": "Праздник",
        "нагруженность": "Нагруженность инфографики",
        "язык": "Язык инфографики",
        "название товара": "Название товара/бренда",
        "преимущество 1": "Топ 1 преимущества товара",
        "преимущество 2": "Топ 2 преимущества товара",
        "преимущество 3": "Топ 3 преимущества товара",
        "преимущества": "Топ 1 преимущества товара",
        "доп текст": "Дополнительная информация о продукте",
        "тип кроя штанов": "Тип кроя штанов",
        "тип рукав": "Тип рукава",
        "длина изделия": "Длина изделия",
        "формат фото": "Формат фото",
        "фото товара": "Фото товара",
        "фото фона": "Фото фона",
        "фотографию модели": "Фото модели",
        "ширину": "Ширина",
        "высоту": "Высота",
        "длину см": "Длина",
        "модель": "Модель"
    }
    
    for key, val in mapping.items():
        if key in low:
            return val

    if len(clean) > 40:
        clean = clean[:40].rstrip()
        
    return clean or fallback


async def _get_prompt_placeholders(db: aiosqlite.Connection) -> list[dict]:
    placeholders: list[dict] = []
    async with db.execute(
        "SELECT id, step_key, question_text FROM steps ORDER BY step_key"
    ) as cur:
        rows = await cur.fetchall()
    for step_id, step_key, question_text in rows:
        label = _normalize_placeholder_label(question_text, step_key)
        placeholders.append({
            "id": step_id,
            "label": label,
            "token": f"{{{label}}}"
        })
    
    seen = set()
    unique = []
    for p in placeholders:
        key = p["label"].lower()
        if key not in seen:
            seen.add(key)
            unique.append(p)
    return unique

async def cleanup_old_history():
    
    while True:
        try:
            async with aiosqlite.connect(DB_PATH) as db:
                
                async with db.execute("SELECT pid, input_paths, result_path FROM generation_history WHERE created_at < datetime('now', '-7 days')") as cur:
                    old_records = await cur.fetchall()
                
                if old_records:
                    print(f"Cleanup: Found {len(old_records)} old generations to delete")
                    for pid, inps_json, res_path in old_records:
                        
                        if res_path and os.path.exists(res_path):
                            try: os.remove(res_path)
                            except: pass
                        if inps_json:
                            try:
                                inps = json.loads(inps_json)
                                for p in inps:
                                    if p and os.path.exists(p):
                                        try: os.remove(p)
                                        except: pass
                            except: pass
                        
                        
                        await db.execute("DELETE FROM generation_history WHERE pid=?", (pid,))
                    await db.commit()
        except Exception as e:
            print(f"Cleanup history error: {e}")
        
        
        await asyncio.sleep(3600)

@asynccontextmanager
async def lifespan(app: FastAPI):
    try:
        os.makedirs(UPLOAD_DIR, exist_ok=True)
        os.makedirs(STATIC_DIR, exist_ok=True)
        os.makedirs(os.path.join(BASE_DIR, "data", "history"), exist_ok=True)
        
        
        try:
            async with aiosqlite.connect(DB_PATH, timeout=30) as db:
                await run_migrations(db)
                
                async with db.execute("PRAGMA table_info(generation_history)") as cur:
                    h_cols = [row[1] for row in await cur.fetchall()]
                if h_cols:
                    if "input_paths" not in h_cols:
                        try:
                            await db.execute("ALTER TABLE generation_history ADD COLUMN input_paths TEXT")
                            await db.commit()
                        except Exception as e: print(f"Migration error (history.input_paths): {e}")
                    if "result_path" not in h_cols:
                        try:
                            await db.execute("ALTER TABLE generation_history ADD COLUMN result_path TEXT")
                            await db.commit()
                        except Exception as e: print(f"Migration error (history.result_path): {e}")
                    if "prompt" not in h_cols:
                        try:
                            await db.execute("ALTER TABLE generation_history ADD COLUMN prompt TEXT")
                            await db.commit()
                        except Exception as e: print(f"Migration error (history.prompt): {e}")
        except Exception as e:
            print(f"ERROR in lifespan migrations: {e}")
            import traceback
            traceback.print_exc()
        
        
        try:
            asyncio.create_task(cleanup_old_history())
        except Exception as e:
            print(f"ERROR creating cleanup task: {e}")
    except Exception as e:
        print(f"CRITICAL ERROR in lifespan: {e}")
        import traceback
        traceback.print_exc()
        
    yield


BASE_URL = os.getenv("BASE_URL", "https://g-box.space").rstrip("/")
BASE_PATH = os.getenv("BASE_PATH", "").rstrip("/")  

app = FastAPI(title="AI-ROOM Admin Panel", lifespan=lifespan, root_path=BASE_PATH or None)


SESSION_SECRET = os.getenv("SESSION_SECRET", "change-me-in-production-" + secrets.token_hex(16))
app.add_middleware(SessionMiddleware, secret_key=SESSION_SECRET, max_age=86400 * 7)


if BASE_PATH:
    from starlette.types import ASGIApp, Receive, Scope, Send
    class PathStripMiddleware:
        def __init__(self, app: ASGIApp) -> None:
            self.app = app
        async def __call__(self, scope: Scope, receive: Receive, send: Send) -> None:
            path = scope.get("path", "")
            if path.startswith(BASE_PATH):
                scope = dict(scope)
                new_path = path[len(BASE_PATH):] or "/"
                scope["path"] = new_path
                if scope.get("root_path"):
                    scope["root_path"] = (scope["root_path"] or "") + BASE_PATH
            await self.app(scope, receive, send)
    app.add_middleware(PathStripMiddleware)


app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
app.mount("/data", StaticFiles(directory=os.path.join(BASE_DIR, "data")), name="data")
app.mount("/uploads", StaticFiles(directory=os.path.join(BASE_DIR, "data", "uploads")), name="uploads")


templates = Jinja2Templates(directory=os.path.join(BASE_DIR, "admin_web", "templates"))
templates.env.filters["from_json"] = json.loads
templates.env.globals["is_site_admin"] = lambda email: email == ADMIN_USER if email else False
templates.env.globals["base_url"] = BASE_URL
templates.env.globals["user"] = None  
from admin_web.site_strings import get_site_string, SITE_STRINGS
templates.env.globals["get_site_string"] = get_site_string
templates.env.globals["SITE_STRINGS"] = SITE_STRINGS
security = HTTPBasic()

try:
    settings = load_settings()
except Exception as e:
    print(f"Error loading settings: {e}")
    class MockSettings:
        bot_token = "MOCK"
        class MockProxy:
            def as_url(self): return None
        proxy = MockProxy()
    settings = MockSettings()

ADMIN_USER = "galimov5500"
ADMIN_PASS = os.getenv("ADMIN_PASS", "")


def get_current_username(credentials: HTTPBasicCredentials = Depends(security)):
    correct_username = secrets.compare_digest(credentials.username, ADMIN_USER)
    correct_password = secrets.compare_digest(credentials.password, ADMIN_PASS)
    if not (correct_username and correct_password):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Incorrect username or password",
            headers={"WWW-Authenticate": "Basic"},
        )
    return credentials.username

async def get_db():
    db = await aiosqlite.connect(DB_PATH, timeout=30)
    db.row_factory = aiosqlite.Row
    try:
        yield db
    finally:
        await db.close()


def _get_site_user_from_session(request: Request) -> dict | None:
    return request.session.get("site_user")


async def get_current_site_user(request: Request, db: aiosqlite.Connection = Depends(get_db)):
    user = _get_site_user_from_session(request)
    if not user:
        return None
    async with db.execute("SELECT id, email, language FROM site_users WHERE id=?", (user["id"],)) as cur:
        row = await cur.fetchone()
    if not row:
        request.session.pop("site_user", None)
        return None
    async with db.execute("SELECT balance FROM users WHERE id=?", (-row[0],)) as cur:
        bal = await cur.fetchone()
    balance = int(bal[0]) if bal and bal[0] is not None else 0
    return {"id": row[0], "email": row[1], "balance": balance, "language": row[2] or "ru"}


async def require_site_user(request: Request, db: aiosqlite.Connection = Depends(get_db)):
    user = await get_current_site_user(request, db)
    if not user:
        raise HTTPException(status_code=302, headers={"Location": "/login?next=" + request.url.path})
    return user


CATEGORIES = ["presets", "female", "male", "child", "boy", "girl", "storefront", "whitebg", "random", "random_other", "own", "own_variant", "infographic_clothing", "infographic_other"]
CORE_CATEGORY_KEYS = set(CATEGORIES)



@app.exception_handler(Exception)
async def global_exception_handler(request: Request, exc: Exception):
    import traceback
    traceback.print_exc()
    return JSONResponse(
        status_code=500,
        content={"error": str(exc), "type": type(exc).__name__}
    )



@app.get("/health")
async def health():
    return {"status": "ok", "base_url": BASE_URL}


@app.get("/download/{rel_path:path}")
async def download_file(rel_path: str):
    import os
    import mimetypes
    from fastapi.responses import FileResponse, JSONResponse

    # Очищаем имя файла от лишнего
    rel = (rel_path or "").replace("\\", "/").lstrip("/")
    if rel.startswith("data/"):
        rel = rel[5:]
    if rel.startswith("history/"):
        rel = rel[8:]

    # Жестко прописываем абсолютные пути для Docker
    path_in_history = f"/app/data/history/{rel}"
    path_in_data = f"/app/data/{rel}"

    # Проверяем наличие файла
    if os.path.exists(path_in_history):
        media_type = mimetypes.guess_type(path_in_history)[0] or "application/octet-stream"
        return FileResponse(path=path_in_history, media_type=media_type, filename=rel)

    elif os.path.exists(path_in_data):
        media_type = mimetypes.guess_type(path_in_data)[0] or "application/octet-stream"
        return FileResponse(path=path_in_data, media_type=media_type, filename=rel)

    # ЕСЛИ ФАЙЛА НЕТ - ВЫВОДИМ ДИАГНОСТИКУ ПРЯМО НА ЭКРАН
    try:
        if os.path.exists("/app/data/history"):
            files_in_history = os.listdir("/app/data/history")[-10:]  # берем 10 последних файлов для примера
        else:
            files_in_history = "Папка /app/data/history не существует в этом контейнере!"
    except Exception as e:
        files_in_history = str(e)

    return JSONResponse(status_code=404, content={
        "error": "Файл физически не найден на диске контейнера admin_web",
        "requested_file": rel,
        "searched_paths": [path_in_history, path_in_data],
        "what_server_sees_in_history_folder": files_in_history
    })


@app.get("/api/site/strings")
async def api_site_strings(lang: str = "ru"):
    
    from admin_web.site_strings import SITE_STRINGS
    strings = SITE_STRINGS.get(lang, SITE_STRINGS["ru"])
    return {"lang": lang, "strings": strings}


@app.get("/api/site/status")
async def api_site_status():
    
    try:
        import sys
        if BASE_DIR not in sys.path:
            sys.path.insert(0, BASE_DIR)
        from bot.db import Database
        bot_db = Database(DB_PATH)
        keys = await bot_db.list_api_keys()
        active = sum(1 for k in keys if k[2])
        return {"api_keys_total": len(keys), "api_keys_active": active, "db_path": DB_PATH}
    except Exception as e:
        return {"error": str(e), "db_path": DB_PATH}



@app.get("/login", response_class=HTMLResponse)
async def login_page(request: Request, next_url: str = Query("/welcome", alias="next")):
    if _get_site_user_from_session(request):
        return RedirectResponse(url=next_url or "/welcome", status_code=302)
    return templates.TemplateResponse("login.html", {"request": request, "next_url": next_url, "user": None})


@app.post("/login", response_class=RedirectResponse)
async def login_submit(
    request: Request,
    email: str = Form(...),
    password: str = Form(...),
    next: str = Form("/welcome"),
    db: aiosqlite.Connection = Depends(get_db),
):
    async with db.execute("SELECT id, email, password_hash, balance, language FROM site_users WHERE LOWER(email)=?", (email.lower(),)) as cur:
        row = await cur.fetchone()
    if not row or not bcrypt.checkpw(password.encode("utf-8"), row[2].encode("utf-8") if isinstance(row[2], str) else row[2]):
        return templates.TemplateResponse("login.html", {"request": request, "error": "Неверный email или пароль", "next_url": next, "user": None})
    request.session["site_user"] = {"id": row[0], "email": row[1], "balance": row[3], "language": row[4] or "ru"}
    return RedirectResponse(url=next or "/welcome", status_code=302)


@app.get("/register", response_class=HTMLResponse)
async def register_page(request: Request):
    if _get_site_user_from_session(request):
        return RedirectResponse(url="/welcome", status_code=302)
    return templates.TemplateResponse("register.html", {"request": request, "user": None})


@app.post("/register", response_class=HTMLResponse)
async def register_submit(
    request: Request,
    email: str = Form(...),
    password: str = Form(...),
    password2: str = Form(...),
    db: aiosqlite.Connection = Depends(get_db),
):
    if password != password2:
        return templates.TemplateResponse("register.html", {"request": request, "error": "Пароли не совпадают", "user": None})
    if len(password) < 6:
        return templates.TemplateResponse("register.html", {"request": request, "error": "Пароль должен быть не менее 6 символов", "user": None})
    try:
        async with db.execute("SELECT id FROM site_users WHERE LOWER(email)=?", (email.lower(),)) as cur:
            if await cur.fetchone():
                return templates.TemplateResponse("register.html", {"request": request, "error": "Пользователь с таким email уже существует", "user": None})
        ph = bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")
        await db.execute("INSERT INTO site_users (email, password_hash, balance, language) VALUES (?, ?, 50, 'ru')", (email, ph))
        await db.commit()
        async with db.execute("SELECT last_insert_rowid()") as cur:
            site_id = (await cur.fetchone())[0]
        await db.execute(
            "INSERT INTO users (id, balance, generation_price, custom_1k_price_rub, language, blocked, accepted_terms, trial_used) "
            "VALUES (?, 50, 30, NULL, 'ru', 0, 0, 0)",
            (-site_id,),
        )
        await db.commit()
        request.session["site_user"] = {"id": site_id, "email": email, "balance": 50, "language": "ru"}
        return RedirectResponse(url="/welcome", status_code=302)
    except Exception as e:
        import traceback
        traceback.print_exc()
        return templates.TemplateResponse("register.html", {
            "request": request,
            "error": "Ошибка регистрации. Попробуйте позже.",
            "user": None
        })


@app.get("/logout", response_class=RedirectResponse)
async def logout(request: Request):
    request.session.pop("site_user", None)
    return RedirectResponse(url="/welcome", status_code=302)


@app.get("/profile", response_class=HTMLResponse)
async def profile_page(request: Request, user=Depends(require_site_user), db: aiosqlite.Connection = Depends(get_db)):
    async with db.execute(
        "SELECT pid, category, result_path, created_at FROM generation_history WHERE user_id=? ORDER BY created_at DESC LIMIT 50",
        (-user["id"],),
    ) as cur:
        history = await cur.fetchall()
    return templates.TemplateResponse(
        "profile.html",
        {"request": request, "user": user, "history": [dict(h) for h in history]},
    )


@app.post("/profile/lang")
async def profile_lang(request: Request, lang: str = Form(...), next_url: str = Form("/profile"), user=Depends(require_site_user), db: aiosqlite.Connection = Depends(get_db)):
    if lang not in ("ru", "en", "vi"):
        lang = "ru"
    await db.execute("UPDATE site_users SET language=? WHERE id=?", (lang, user["id"]))
    await db.execute("UPDATE users SET language=? WHERE id=?", (lang, -user["id"]))
    await db.commit()
    if request.session.get("site_user"):
        request.session["site_user"]["language"] = lang
    return RedirectResponse(url=next_url or "/profile", status_code=302)


@app.post("/site/lang")
async def site_lang_guest(request: Request, lang: str = Form(...), next_url: str = Form("/welcome")):
    
    if lang in ("ru", "en", "vi"):
        request.session["preferred_lang"] = lang
    return RedirectResponse(url=next_url or "/welcome", status_code=302)


@app.get("/welcome", response_class=HTMLResponse)
async def welcome_page(request: Request, user=Depends(get_current_site_user)):
    if not user:
        return RedirectResponse(url="/login?next=/welcome", status_code=302)
    from admin_web.site_strings import SITE_STRINGS
    lang = user.get("language", "ru") or "ru"
    return templates.TemplateResponse("welcome.html", {
        "request": request, "user": user,
        "lang": lang,
        "site_strings": SITE_STRINGS.get(lang, SITE_STRINGS["ru"])
    })


@app.get("/balance", response_class=HTMLResponse)
async def site_balance_page(request: Request, user=Depends(require_site_user), db: aiosqlite.Connection = Depends(get_db)):
    contact = "https://t.me/bnbslow"
    async with db.execute("SELECT value FROM app_settings WHERE key='support_contact'") as cur:
        row = await cur.fetchone()
        if row and row[0]:
            contact = row[0]
    return templates.TemplateResponse("site_balance.html", {"request": request, "user": user, "contact": contact})


@app.get("/help", response_class=HTMLResponse)
async def site_support_page(request: Request, user=Depends(get_current_site_user), db: aiosqlite.Connection = Depends(get_db)):
    contact = "https://t.me/bnbslow"
    async with db.execute("SELECT value FROM app_settings WHERE key='support_contact'") as cur:
        row = await cur.fetchone()
        if row and row[0]:
            contact = row[0]
    return templates.TemplateResponse("site_support.html", {"request": request, "user": user, "contact": contact})


@app.get("/instructions", response_class=HTMLResponse)
async def site_instructions_page(request: Request, user=Depends(get_current_site_user), db: aiosqlite.Connection = Depends(get_db)):
    content = ""
    async with db.execute("SELECT value FROM app_settings WHERE key='howto_text'") as cur:
        row = await cur.fetchone()
        if row and row[0]:
            content = row[0]
    return templates.TemplateResponse("site_instructions.html", {"request": request, "user": user, "content": content})


@app.get("/terms", response_class=HTMLResponse)
async def site_terms_page(request: Request, user=Depends(get_current_site_user), db: aiosqlite.Connection = Depends(get_db)):
    content = ""
    async with db.execute("SELECT value FROM app_settings WHERE key='agreement_text'") as cur:
        row = await cur.fetchone()
        if row and row[0]:
            content = row[0]
    return templates.TemplateResponse("site_terms.html", {"request": request, "user": user, "content": content})


@app.get("/api/site/categories")
async def api_site_categories(db: aiosqlite.Connection = Depends(get_db)):
    async with db.execute(
        "SELECT id, key, name_ru, is_active, order_index FROM categories WHERE is_active=1 ORDER BY order_index, id"
    ) as cur:
        rows = await cur.fetchall()
    return {"categories": [{"id": r[0], "key": r[1], "name_ru": r[2], "is_active": r[3], "order_index": r[4]} for r in rows]}


@app.get("/api/site/category/{key}/enabled")
async def api_site_category_enabled(key: str, db: aiosqlite.Connection = Depends(get_db)):
    async with db.execute("SELECT value FROM app_settings WHERE key=?", (key,)) as cur:
        row = await cur.fetchone()
    enabled = row is None or str(row[0]) != "0"
    return {"enabled": enabled}


@app.get("/api/site/models")
async def api_site_models(category: str, db: aiosqlite.Connection = Depends(get_db)):
    async with db.execute(
        "SELECT m.id, m.name, m.photo_file_id, m.category FROM models m WHERE m.category=? AND m.is_active=1 ORDER BY m.position, m.id",
        (category,),
    ) as cur:
        rows = await cur.fetchall()
    models = []
    for r in rows:
        photo_url = None
        if r[2] and str(r[2]).startswith("data/"):
            photo_url = "/" + r[2].replace("\\", "/")
        elif r[2]:
            photo_url = f"/uploads/{r[2]}" if not str(r[2]).startswith("/") else r[2]
        models.append({"id": r[0], "name": r[1], "photo_url": photo_url, "category": r[3]})
    return {"models": models}


@app.post("/api/site/generate")
async def api_site_generate(
    request: Request,
    user=Depends(require_site_user),
    db: aiosqlite.Connection = Depends(get_db),
):
    try:
        return await _api_site_generate_impl(request, user, db)
    except Exception as e:
        import traceback
        traceback.print_exc()
        return JSONResponse({"error": str(e)}, status_code=500)


async def _api_site_generate_impl(request: Request, user: dict, db):
    form = await request.form()
    category = form.get("category", "")
    model_id = form.get("model_id") or None
    if model_id:
        try:
            model_id = int(model_id)
        except ValueError:
            model_id = None
    aspect = form.get("aspect", "1:1") or "1:1"
    photos = form.getlist("photos")
    if isinstance(photos, str):
        photos = [photos] if photos else []

    if not category:
        return JSONResponse({"error": "Выберите категорию"}, status_code=400)
    min_photos = 2 if category == "own_variant" else 1
    if len(photos) < min_photos:
        return JSONResponse({"error": f"Загрузите минимум {min_photos} фото"}, status_code=400)

    import sys
    if BASE_DIR not in sys.path:
        sys.path.insert(0, BASE_DIR)
    from bot.db import Database
    from bot.gemini import generate_image

    bot_db = Database(DB_PATH)
    user_id = -user["id"]
    price = await bot_db.get_user_generation_price(user_id)

    balance = await bot_db.get_user_balance(user_id)
    if balance < price:
        return JSONResponse({"error": "Недостаточно средств на балансе"}, status_code=402)
    if await bot_db.get_maintenance():
        return JSONResponse({"error": "Идут технические работы"}, status_code=503)
    if await bot_db.get_user_blocked(user_id):
        return JSONResponse({"error": "Доступ заблокирован"}, status_code=403)

    images_bytes = []
    
    needs_model_photo = category in ("female", "male", "child", "boy", "girl", "presets", "own", "storefront") and model_id
    if needs_model_photo:
        async with db.execute("SELECT photo_file_id FROM models WHERE id=?", (model_id,)) as cur:
            row = await cur.fetchone()
        if row and row[0]:
            bg_path = str(row[0]).replace("\\", "/")
            full = os.path.join(BASE_DIR, bg_path) if not os.path.isabs(bg_path) else bg_path
            if not os.path.exists(full):
                full = os.path.join(UPLOAD_DIR, bg_path.split("/")[-1])
            if os.path.exists(full):
                with open(full, "rb") as f:
                    images_bytes.append(f.read())
    for p in photos[:5]:
        if hasattr(p, "read"):
            images_bytes.append(await p.read())
        elif isinstance(p, str) and len(p) > 100:
            import base64
            try:
                images_bytes.append(base64.b64decode(p))
            except Exception:
                pass

    if len(images_bytes) < min_photos:
        return JSONResponse({"error": "Не удалось прочитать фото"}, status_code=400)

    prompt = await _build_web_prompt(category, model_id, aspect, bot_db)
    if not prompt:
        return JSONResponse({"error": "Не удалось сформировать промпт"}, status_code=500)

    api_keys = await bot_db.list_api_keys()
    active_keys = [(k[0], k[1]) for k in api_keys if k[2]]
    if not active_keys:
        return JSONResponse({"error": "Нет активных API ключей"}, status_code=503)

    import random
    import uuid
    random.shuffle(active_keys)
    last_err = None
    for kid, token in active_keys[:5]:
        ok, _ = await bot_db.check_api_key_limits(kid)
        if not ok:
            continue
        try:
            result_path = await generate_image(
                api_key=token,
                prompt=prompt,
                images_bytes=images_bytes,
                aspect_ratio=aspect,
                key_id=kid,
                db_instance=bot_db,
            )
            if result_path:
                pid = f"WEB{str(uuid.uuid4().hex[:10]).upper()}"
                rp = result_path.replace("\\", "/")
                rp_db = rp.split("/")[-1] if "/" in rp else rp  
                await bot_db.subtract_user_balance(user_id, price, reason="generation")
                await bot_db.add_generation_history(
                    pid=pid,
                    user_id=user_id,
                    category=category,
                    params=json.dumps({"model_id": model_id, "aspect": aspect}),
                    input_photos="[]",
                    result_photo_id="",
                    input_paths="[]",
                    result_path=rp_db,
                    prompt=prompt[:2000],
                )
                await bot_db.record_api_usage(kid)
                new_balance = await bot_db.get_user_balance(user_id)
                return JSONResponse({
                    "result_path": rp_db,
                    "result_url": "/data/" + rp_db,
                    "new_balance": new_balance,
                })
        except Exception as e:
            last_err = str(e)
            try:
                from bot.gemini import is_proxy_error, is_fatal_key_error
                await bot_db.record_api_error(
                    kid,
                    token[:10],
                    type(e).__name__,
                    str(e),
                    status_code=getattr(e, "status_code", None),
                    is_proxy_error=is_proxy_error(e),
                )
                if is_fatal_key_error(e):
                    await bot_db.update_api_key(kid, is_active=0)
            except Exception:
                pass
            continue

    return JSONResponse({"error": last_err or "Ошибка генерации"}, status_code=500)


async def _build_web_prompt(category: str, model_id: int | None, aspect: str, db) -> str:
    
    category_prompt_map = {
        "storefront": "storefront_prompt",
        "whitebg": "whitebg_prompt",
        "own_variant": "own_variant_prompt",
        "random": "random_prompt",
        "random_other": "random_other_prompt",
        "infographic_clothing": "infographic_clothing_prompt",
        "infographic_other": "infographic_other_prompt",
        "presets": "presets_prompt",
        "own": "own_prompt",
        "collage": "collage_prompt",
    }
    
    
    
    if category in ("female", "male", "child", "boy", "girl", "presets", "own", "storefront") and model_id:
        base = ""
        
        async with aiosqlite.connect(DB_PATH) as conn:
            async with conn.execute(
                "SELECT p.text FROM prompts p JOIN models m ON m.prompt_id=p.id WHERE m.id=?",
                (model_id,),
            ) as cur:
                row = await cur.fetchone()
        if row:
            base = row[0]
        else:
            base = "Professional commercial photography. High quality, 8k resolution."
        
        
        if category == "presets":
            presets_base = await db.get_app_setting("presets_prompt") or ""
            if presets_base:
                base += "\n\n" + presets_base
        elif category == "own":
            own_base = await db.get_app_setting("own_prompt") or ""
            if own_base:
                base += "\n\n" + own_base
        
        
        
        base += f" Aspect ratio: {aspect}. Produce ONE single image."
        return base
    
    
    prompt_key = category_prompt_map.get(category)
    if prompt_key:
        prompt = await db.get_app_setting(prompt_key) or ""
        if prompt:
            base = prompt
        else:
            base = "Professional commercial photography. High quality, 8k resolution."
    else:
        base = "Professional commercial photography. High quality, 8k resolution."
    
    base += f" Aspect ratio: {aspect}. Produce ONE single image."
    return base


@app.get("/models/toggle/{model_id}")
async def toggle_model_status(model_id: int, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    
    async with db.execute("SELECT is_active FROM models WHERE id=?", (model_id,)) as cur:
        row = await cur.fetchone()
        if not row:
            return RedirectResponse(url="/prompts", status_code=303)
        current = row[0]
    
    new_val = 0 if current == 1 else 1
    await db.execute("UPDATE models SET is_active=? WHERE id=?", (new_val, model_id))
    await db.commit()
    return RedirectResponse(url="/prompts", status_code=303)

@app.get("/models/delete/{model_id}")
async def delete_model(model_id: int, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    async with db.execute("SELECT prompt_id, photo_file_id FROM models WHERE id=?", (model_id,)) as cur:
        row = await cur.fetchone()
        if row:
            prompt_id, photo_path = row
            await db.execute("DELETE FROM models WHERE id=?", (model_id,))
            await db.execute("DELETE FROM prompts WHERE id=?", (prompt_id,))
            if photo_path and photo_path.startswith("data/"):
                full_path = os.path.join(BASE_DIR, photo_path)
                if os.path.exists(full_path):
                    try:
                        os.remove(full_path)
                    except Exception as e:
                        print(f"Error removing file {full_path}: {e}")
            await db.commit()
    return RedirectResponse(url="/prompts", status_code=303)

@app.post("/models/delete_bulk")
async def delete_bulk_models(model_ids: str = Form(...), db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    ids = [int(x) for x in model_ids.split(",") if x.isdigit()]
    for model_id in ids:
        async with db.execute("SELECT prompt_id, photo_file_id FROM models WHERE id=?", (model_id,)) as cur:
            row = await cur.fetchone()
            if row:
                prompt_id, photo_path = row
                await db.execute("DELETE FROM models WHERE id=?", (model_id,))
                await db.execute("DELETE FROM prompts WHERE id=?", (prompt_id,))
                if photo_path and photo_path.startswith("data/"):
                    full_path = os.path.join(BASE_DIR, photo_path)
                    if os.path.exists(full_path):
                        try: os.remove(full_path)
                        except Exception: pass
    await db.commit()
    return RedirectResponse(url="/prompts", status_code=303)


async def get_proxy_url(db: aiosqlite.Connection = None):
    if db:
        async with db.execute("SELECT value FROM app_settings WHERE key='bot_proxy'") as cur:
            row = await cur.fetchone()
            if row and row[0]:
                proxy_val = row[0]
                if "://" not in proxy_val:
                    parts = proxy_val.split(":")
                    if len(parts) == 4:
                        return f"http://{parts[2]}:{parts[3]}@{parts[0]}:{parts[1]}"
                    elif len(parts) == 2:
                        return f"http://{parts[0]}:{parts[1]}"
                return proxy_val
    return settings.proxy.as_url()

async def check_proxy(db: aiosqlite.Connection = None):
    try:
        proxy_url = await get_proxy_url(db)
        if not proxy_url: return "❌ Не настроен"
        test_urls = ["http://api.ipify.org", "https://api.ipify.org", "http://google.com"]
        async with httpx.AsyncClient(proxy=proxy_url, timeout=15, follow_redirects=True, verify=False) as client:
            for url in test_urls:
                try:
                    resp = await client.get(url)
                    if resp.status_code < 500:
                        return f"✅ Работает (код {resp.status_code})"
                except Exception: continue
            return "❌ Прокси не отвечает (Timeout/Auth/IP Block)"
    except Exception as e:
        return f"❌ Ошибка: {str(e)[:50]}"

async def run_broadcast(message_text: str):
    if settings.bot_token == "MOCK": return
    bot = Bot(token=settings.bot_token)
    db = await aiosqlite.connect(DB_PATH)
    async with db.execute("SELECT id FROM users") as cur:
        users = await cur.fetchall()
    for user in users:
        try:
            await bot.send_message(user[0], message_text)
            await asyncio.sleep(0.05)
        except Exception: continue
    await bot.session.close()
    await db.close()

async def send_subscription_notification(user_id: int, plan_name: str, expires_at_str: str, daily_limit: int, lang: str):
    
    if settings.bot_token == "MOCK": return
    
    try:
        
        try:
            expires_str = expires_at_str.replace('Z', '') if 'Z' in expires_at_str else expires_at_str
            if 'T' in expires_str:
                expires_dt = datetime.fromisoformat(expires_str)
            else:
                expires_dt = datetime.fromisoformat(expires_str + "T00:00:00")
            expires_date = expires_dt.strftime("%d.%m.%Y")
            expires_time = expires_dt.strftime("%H:%M")
        except Exception:
            if 'T' in expires_at_str:
                parts = expires_at_str.split('T')
                date_part = parts[0]
                time_part = parts[1][:5] if len(parts[1]) >= 5 else "00:00"
                expires_date = ".".join(reversed(date_part.split("-")))
                expires_time = time_part
            else:
                expires_date = expires_at_str[:10]
                expires_time = "00:00"

        text = get_string("sub_success_congrats", lang, 
                         plan_name=plan_name,
                         expires_date=expires_date,
                         expires_time=expires_time,
                         daily_limit=daily_limit)
        
        bot = Bot(token=settings.bot_token)
        await bot.send_message(user_id, text)
        await bot.session.close()
    except Exception as e:
        print(f"Error sending subscription notification to {user_id}: {e}")



@app.get("/", response_class=HTMLResponse)
async def index(request: Request, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    try:
        
        async with db.execute("SELECT COUNT(*) FROM users") as cur:
            total_users = (await cur.fetchone())[0]
        
        
        async with db.execute("SELECT COUNT(*) FROM users WHERE date(created_at) = date('now')") as cur:
            today_users = (await cur.fetchone())[0]
        
        
        async with db.execute("SELECT COUNT(*) FROM generation_history WHERE date(created_at) = date('now')") as cur:
            today_gens = (await cur.fetchone())[0]

        today_video_gens = 0
        try:
            async with db.execute("SELECT COUNT(*) FROM balance_history WHERE reason='video_generation' AND date(created_at) = date('now')") as cur:
                today_video_gens = (await cur.fetchone())[0]
        except Exception:
            pass

        total_video_gens = 0
        try:
            async with db.execute("SELECT COUNT(*) FROM balance_history WHERE reason='video_generation'") as cur:
                total_video_gens = (await cur.fetchone())[0]
        except Exception:
            pass

        async with db.execute("SELECT SUM(balance) FROM users") as cur:
            row = await cur.fetchone()
            total_balance = row[0] if row and row[0] else 0

        # 1K / 4K breakdown today (exclude video)
        today_1k = 0
        today_4k = 0
        try:
            async with db.execute(
                "SELECT params FROM generation_history WHERE date(created_at) = date('now') AND category NOT LIKE 'video%'"
            ) as cur:
                rows = await cur.fetchall()
            for r in rows:
                p = str(r[0] or "").lower()
                if "4k" in p:
                    today_4k += 1
                else:
                    today_1k += 1
        except Exception:
            pass

        # Total Zveno API calls (all time)
        total_zveno_calls = 0
        try:
            async with db.execute("SELECT COUNT(*) FROM zveno_usage_log WHERE status='ok'") as cur:
                total_zveno_calls = (await cur.fetchone())[0]
        except Exception:
            pass

    except Exception as e:
        print(f"Stats error: {e}")
        total_users = today_users = today_gens = total_balance = 0
        today_1k = today_4k = total_zveno_calls = 0

    maintenance = False
    try:
        async with db.execute("SELECT value FROM app_settings WHERE key='maintenance'") as cur:
            row = await cur.fetchone()
            maintenance = row[0] == '1' if row else False
    except Exception: pass

    disable_4k = False
    try:
        async with db.execute("SELECT value FROM app_settings WHERE key='disable_4k'") as cur:
            row = await cur.fetchone()
            if row is None:
                # Первый заход: зеркалим текущий DISABLE_4K из env, чтобы деплой не менял поведение.
                env_disabled = os.getenv("DISABLE_4K", "0").strip() == "1"
                await db.execute(
                    "INSERT INTO app_settings (key, value) VALUES ('disable_4k', ?) "
                    "ON CONFLICT(key) DO NOTHING",
                    ("1" if env_disabled else "0",),
                )
                await db.commit()
                disable_4k = env_disabled
            else:
                disable_4k = str(row[0]).strip() == "1"
    except Exception:
        pass

    proxy_status = await check_proxy(db)
    
    
    recent_errors = []
    proxy_errors_count = 0
    try:
        async with db.execute(
            "SELECT key_id, api_key_preview, error_type, error_message, status_code, is_proxy_error, created_at FROM api_key_errors ORDER BY created_at DESC LIMIT 10"
        ) as cur:
            recent_errors = await cur.fetchall()
        
        async with db.execute(
            "SELECT COUNT(*) FROM api_key_errors WHERE is_proxy_error=1 AND created_at > datetime('now', '-24 hours')"
        ) as cur:
            row = await cur.fetchone()
            proxy_errors_count = int(row[0]) if row else 0
    except Exception as e:
        print(f"Error fetching API errors: {e}")
    
    return templates.TemplateResponse("dashboard.html", {
        "request": request, 
        "total_users": total_users, "today_users": today_users,
        "today_gens": today_gens, "total_balance": total_balance,
        "today_video_gens": today_video_gens, "total_video_gens": total_video_gens,
        "today_1k": today_1k, "today_4k": today_4k,
        "total_zveno_calls": total_zveno_calls,
        "maintenance": maintenance, "disable_4k": disable_4k, "proxy_status": proxy_status,
        "recent_errors": recent_errors, "proxy_errors_count": proxy_errors_count
    })

@app.get("/users", response_class=HTMLResponse)
async def list_users(request: Request, q: str = "", db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    if q:
        query = """
            SELECT u.id, u.username, u.source, u.blocked, u.balance, u.generation_price, u.custom_1k_price_rub, su.email as site_email
            FROM users u 
            LEFT JOIN site_users su ON su.id = -u.id
            WHERE CAST(u.id AS TEXT) LIKE ? OR u.username LIKE ? OR su.email LIKE ?
            ORDER BY u.created_at DESC LIMIT 100
        """
        async with db.execute(query, (f"%{q}%", f"%{q}%", f"%{q}%")) as cur:
            rows = await cur.fetchall()
        
        if not rows and q.isdigit():
            try:
                uid = int(q)
                await db.execute(
                    "INSERT OR IGNORE INTO users (id, balance, generation_price, custom_1k_price_rub, language, accepted_terms, blocked, trial_used) "
                    "VALUES (?, 50, 30, NULL, 'ru', 0, 0, 0)",
                    (uid,),
                )
                await db.commit()
                async with db.execute(query, (f"%{q}%", f"%{q}%", f"%{q}%")) as cur2:
                    rows = await cur2.fetchall()
            except Exception:
                pass
    else:
        query = """
            SELECT u.id, u.username, u.source, u.blocked, u.balance, u.generation_price, u.custom_1k_price_rub, su.email as site_email
            FROM users u 
            LEFT JOIN site_users su ON su.id = -u.id
            ORDER BY u.created_at DESC LIMIT 50
        """
        async with db.execute(query) as cur:
            rows = await cur.fetchall()
    users = [dict(zip(r.keys(), r)) for r in rows]
    async with db.execute("SELECT value FROM app_settings WHERE key='standard_1k_price_rub'") as cur:
        _sp = await cur.fetchone()
    try:
        standard_price = int(_sp[0]) if _sp and _sp[0] is not None else 30
    except (ValueError, TypeError):
        standard_price = 30
    return templates.TemplateResponse("users.html", {
        "request": request,
        "users": users,
        "q": q,
        "standard_price": standard_price,
    })

@app.post("/users/edit_balance")
async def edit_balance(
    user_id: int = Form(...),
    amount: int = Form(...),
    price: str = Form(""),
    db: aiosqlite.Connection = Depends(get_db),
    user: str = Depends(get_current_username)
):
    from bot.db import DEFAULT_USER_TARIFF_1K_RUB

    async with db.execute("SELECT balance FROM users WHERE id = ?", (user_id,)) as cur:
        row = await cur.fetchone()
        old_balance = row[0] if row else 0
    
    change = amount - old_balance

    raw = (price or "").strip()
    if not raw:
        await db.execute(
            "UPDATE users SET balance = ?, custom_1k_price_rub = NULL, generation_price = ? WHERE id = ?",
            (amount, DEFAULT_USER_TARIFF_1K_RUB, user_id),
        )
    else:
        safe_price = max(1, int(raw))
        await db.execute(
            "UPDATE users SET balance = ?, custom_1k_price_rub = ?, generation_price = ? WHERE id = ?",
            (amount, safe_price, safe_price, user_id),
        )
    
    
    await db.execute(
        "INSERT INTO balance_history (user_id, amount, new_balance, reason, admin_id) VALUES (?, ?, ?, ?, ?)",
        (user_id, change, amount, "admin_edit", user)
    )
    
    await db.commit()
    return RedirectResponse(url=f"/users?q={user_id}", status_code=303)

@app.post("/users/set_standard_price")
async def set_standard_price(
    price: int = Form(...),
    db: aiosqlite.Connection = Depends(get_db),
    user: str = Depends(get_current_username),
):
    """Массовая смена стандартной цены 1K: применяется ко ВСЕМ пользователям БЕЗ своей
    (кастомной) цены. Пользователи с заданным custom_1k_price_rub не затрагиваются."""
    p = max(1, int(price))
    await db.execute(
        "INSERT INTO app_settings (key, value) VALUES ('standard_1k_price_rub', ?) "
        "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (str(p),),
    )
    # Синхроним легаси-колонку (для отображения) только у стандартных пользователей.
    await db.execute(
        "UPDATE users SET generation_price = ? WHERE custom_1k_price_rub IS NULL",
        (p,),
    )
    await db.commit()
    return RedirectResponse(url="/users", status_code=303)

@app.post("/admin/block_user")
async def admin_block_user_route(
    user_id: int = Form(...),
    block: int = Form(...),
    db: aiosqlite.Connection = Depends(get_db),
    user: str = Depends(get_current_username)
):
    await db.execute("UPDATE users SET blocked = ? WHERE id = ?", (block, user_id))
    await db.commit()
    return RedirectResponse(url=f"/users?q={user_id}", status_code=303)

@app.post("/cancel_subscription")
async def cancel_subscription(user_id: int = Form(...), db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    await db.execute("DELETE FROM subscriptions WHERE user_id = ?", (user_id,))
    await db.commit()
    return RedirectResponse(url=f"/users?q={user_id}", status_code=303)

@app.post("/users/add_requests")
async def add_requests(
    user_id: int = Form(...),
    extra: int = Form(...),
    db: aiosqlite.Connection = Depends(get_db),
    user: str = Depends(get_current_username)
):
    if extra <= 0:
        return RedirectResponse(url=f"/users?q={user_id}", status_code=303)
    await db.execute(
        "UPDATE subscriptions SET daily_limit = daily_limit + ? WHERE user_id=? AND datetime(expires_at) > CURRENT_TIMESTAMP",
        (extra, user_id)
    )
    await db.commit()
    return RedirectResponse(url=f"/users?q={user_id}", status_code=303)

@app.get("/mailing", response_class=HTMLResponse)
async def mailing_page(request: Request, user: str = Depends(get_current_username)):
    return templates.TemplateResponse("mailing.html", {"request": request})


@app.get("/balance_history", response_class=HTMLResponse)
async def balance_history_page(
        request: Request,
        db: aiosqlite.Connection = Depends(get_db),
        user: str = Depends(get_current_username)
):
    from datetime import datetime, timedelta

    # 1. История (оставляем как есть)
    history_query = """
        SELECT bh.*, u.username 
        FROM balance_history bh
        JOIN users u ON bh.user_id = u.id
        ORDER BY bh.created_at DESC
        LIMIT 1000
    """
    async with db.execute(history_query) as cur:
        history = await cur.fetchall()

    # 2. График за 30 дней (Сдвиг +3 часа МСК, считаем сумму всех ручных пополнений и корректировок)
    stats_query = """
        SELECT date(datetime(created_at, '+3 hours')) as day, SUM(amount) as total
        FROM balance_history
        WHERE reason IN ('recharge', 'admin_edit')
        AND datetime(created_at, '+3 hours') >= date('now', '+3 hours', '-30 days')
        GROUP BY day
        ORDER BY day
    """
    async with db.execute(stats_query) as cur:
        stats = await cur.fetchall()

    chart_labels = [row[0] for row in stats]
    chart_values = [row[1] for row in stats]

    # 2b. Daily generation counts for tooltip
    chart_1k = []
    chart_4k = []
    chart_video = []
    try:
        gen_daily_query = """
            SELECT date(datetime(created_at, '+3 hours')) as day, category, params
            FROM generation_history
            WHERE datetime(created_at, '+3 hours') >= date('now', '+3 hours', '-30 days')
        """
        async with db.execute(gen_daily_query) as cur:
            gen_rows = await cur.fetchall()
        from collections import defaultdict
        day_1k = defaultdict(int)
        day_4k = defaultdict(int)
        day_vid = defaultdict(int)
        for r in gen_rows:
            d = r[0]
            cat = str(r[1] or "").lower()
            p = str(r[2] or "").lower()
            if cat.startswith("video"):
                day_vid[d] += 1
            elif "4k" in p:
                day_4k[d] += 1
            else:
                day_1k[d] += 1
        for lbl in chart_labels:
            chart_1k.append(day_1k.get(lbl, 0))
            chart_4k.append(day_4k.get(lbl, 0))
            chart_video.append(day_vid.get(lbl, 0))
    except Exception:
        chart_1k = [0] * len(chart_labels)
        chart_4k = [0] * len(chart_labels)
        chart_video = [0] * len(chart_labels)

    # 3. Отчет за текущий месяц (Сдвиг +3 часа МСК. Доход: все admin_edit. Расход: только генерации)
    report_query = """
        SELECT 
            SUM(CASE WHEN reason IN ('recharge', 'admin_edit') THEN amount ELSE 0 END) as total_in,
            SUM(CASE WHEN amount < 0 AND reason = 'generation' THEN ABS(amount) ELSE 0 END) as total_out,
            COUNT(*) as total_changes
        FROM balance_history
        WHERE datetime(created_at, '+3 hours') >= datetime('now', '+3 hours', 'start of month')
    """
    async with db.execute(report_query) as cur:
        report = await cur.fetchone()

    # 4. Отчет за прошлый месяц (Сдвиг +3 часа МСК)
    prev_report_query = """
        SELECT 
            SUM(CASE WHEN reason IN ('recharge', 'admin_edit') THEN amount ELSE 0 END) as total_in,
            SUM(CASE WHEN amount < 0 AND reason = 'generation' THEN ABS(amount) ELSE 0 END) as total_out
        FROM balance_history
        WHERE datetime(created_at, '+3 hours') >= datetime('now', '+3 hours', 'start of month', '-1 month')
          AND datetime(created_at, '+3 hours') < datetime('now', '+3 hours', 'start of month')
    """
    async with db.execute(prev_report_query) as cur:
        prev_report = await cur.fetchone()

    # 5. Выручка по месяцам (Группировка с учетом +3 часов)
    monthly_stats_query = """
        SELECT strftime('%Y-%m', datetime(created_at, '+3 hours')) as month, 
               SUM(CASE WHEN reason IN ('recharge', 'admin_edit') THEN amount ELSE 0 END) as total
        FROM balance_history
        WHERE datetime(created_at, '+3 hours') >= datetime('now', '+3 hours', 'start of month', '-6 months')
        GROUP BY month
        ORDER BY month DESC
    """
    async with db.execute(monthly_stats_query) as cur:
        monthly_stats_raw = await cur.fetchall()

    # --- Применяем ручные корректировки из базы ---
    async with db.execute("SELECT value FROM app_settings WHERE key='revenue_manual_correction'") as cur:
        row = await cur.fetchone()
        try:
            revenue_manual_correction = int(str(row[0]).strip()) if row and row[0] is not None else 0
        except Exception:
            revenue_manual_correction = 0

    async with db.execute("SELECT value FROM app_settings WHERE key='revenue_manual_correction_note'") as cur:
        row = await cur.fetchone()
        revenue_manual_correction_note = str(row[0]) if row and row[0] is not None else ""

    # Корректируем главный виджет
    report_total_in = int(report[0] or 0) + revenue_manual_correction
    report = (
        report_total_in,
        int(report[1] or 0),
        int(report[2] or 0),
    )

    # Корректируем список "Выручка по месяцам" (чтобы цифры с главным виджетом совпадали)
    current_month_str = (datetime.utcnow() + timedelta(hours=3)).strftime('%Y-%m')
    monthly_stats = []
    for m_row in monthly_stats_raw:
        m_name = m_row[0]
        m_total = m_row[1] or 0
        if m_name == current_month_str:
            m_total += revenue_manual_correction
        monthly_stats.append((m_name, m_total))

    # 6. Generation breakdown this month (1K / 4K / video)
    month_1k = 0
    month_4k = 0
    month_video = 0
    try:
        async with db.execute(
            "SELECT category, params FROM generation_history WHERE datetime(created_at, '+3 hours') >= datetime('now', '+3 hours', 'start of month')"
        ) as cur:
            gen_rows = await cur.fetchall()
        for r in gen_rows:
            cat = str(r[0] or "").lower()
            p = str(r[1] or "").lower()
            if cat.startswith("video"):
                month_video += 1
            elif "4k" in p:
                month_4k += 1
            else:
                month_1k += 1
    except Exception:
        pass

    return templates.TemplateResponse("balance_history.html", {
        "request": request,
        "history": history,
        "chart_labels": chart_labels,
        "chart_values": chart_values,
        "chart_1k": chart_1k,
        "chart_4k": chart_4k,
        "chart_video": chart_video,
        "report": report,
        "prev_report": prev_report,
        "monthly_stats": monthly_stats,
        "revenue_manual_correction": revenue_manual_correction,
        "revenue_manual_correction_note": revenue_manual_correction_note,
        "month_1k": month_1k,
        "month_4k": month_4k,
        "month_video": month_video,
    })


@app.post("/balance_history/revenue_correction")
async def set_revenue_correction(
    amount: str = Form("0"),
    note: str = Form(""),
    db: aiosqlite.Connection = Depends(get_db),
    user: str = Depends(get_current_username),
):
    raw = (amount or "0").strip()
    try:
        correction = int(raw)
    except Exception:
        correction = 0
    correction = max(-10_000_000, min(correction, 10_000_000))
    safe_note = (note or "").strip()[:500]

    await db.execute(
        "INSERT INTO app_settings (key, value) VALUES ('revenue_manual_correction', ?) "
        "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (str(correction),),
    )
    await db.execute(
        "INSERT INTO app_settings (key, value) VALUES ('revenue_manual_correction_note', ?) "
        "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (safe_note,),
    )
    await db.commit()
    return RedirectResponse(url="/balance_history", status_code=303)


# ─── Metrics ────────────────────────────────────────────────────────────────

def _last_12_months():
    """Return list of {'key': 'YYYY-MM', 'label': 'Мес ГГ'} for last 12 months."""
    MONTH_SHORT_RU = {
        1: "Янв", 2: "Фев", 3: "Мар", 4: "Апр", 5: "Май", 6: "Июн",
        7: "Июл", 8: "Авг", 9: "Сен", 10: "Окт", 11: "Ноя", 12: "Дек",
    }
    now = datetime.utcnow() + timedelta(hours=3)
    months = []
    for i in range(11, -1, -1):
        dt = now - timedelta(days=30 * i)
        dt = dt.replace(day=1)
        key = dt.strftime("%Y-%m")
        label = f"{MONTH_SHORT_RU[dt.month]} {dt.strftime('%y')}"
        months.append({"key": key, "label": label})
    seen = set()
    deduped = []
    for m in months:
        if m["key"] not in seen:
            seen.add(m["key"])
            deduped.append(m)
    return deduped


async def _compute_metrics(db: aiosqlite.Connection):
    months = _last_12_months()
    month_keys = [m["key"] for m in months]

    # 1. New paying customers per month (first recharge/admin_edit in that month)
    q_new_paying = """
        SELECT strftime('%Y-%m', datetime(min_date, '+3 hours')) AS m, COUNT(*) AS cnt
        FROM (
            SELECT user_id, MIN(created_at) AS min_date
            FROM balance_history
            WHERE reason IN ('recharge', 'admin_edit') AND amount > 0
            GROUP BY user_id
        )
        GROUP BY m
    """
    new_paying_map = {}
    async with db.execute(q_new_paying) as cur:
        for row in await cur.fetchall():
            new_paying_map[row[0]] = row[1]

    # 2. Active paying per month
    q_active = """
        SELECT strftime('%Y-%m', datetime(created_at, '+3 hours')) AS m,
               COUNT(DISTINCT user_id) AS cnt
        FROM balance_history
        WHERE reason IN ('recharge', 'admin_edit') AND amount > 0
        GROUP BY m
    """
    active_map = {}
    async with db.execute(q_active) as cur:
        for row in await cur.fetchall():
            active_map[row[0]] = row[1]

    # 2b. Active paying sets for churn calculation
    q_active_sets = """
        SELECT strftime('%Y-%m', datetime(created_at, '+3 hours')) AS m,
               user_id
        FROM balance_history
        WHERE reason IN ('recharge', 'admin_edit') AND amount > 0
        GROUP BY m, user_id
    """
    active_sets = {}
    async with db.execute(q_active_sets) as cur:
        for row in await cur.fetchall():
            active_sets.setdefault(row[0], set()).add(row[1])

    # 3. Revenue per month
    q_revenue = """
        SELECT strftime('%Y-%m', datetime(created_at, '+3 hours')) AS m,
               SUM(amount) AS total,
               COUNT(*) AS cnt
        FROM balance_history
        WHERE reason IN ('recharge', 'admin_edit') AND amount > 0
        GROUP BY m
    """
    revenue_map = {}
    recharge_cnt_map = {}
    async with db.execute(q_revenue) as cur:
        for row in await cur.fetchall():
            revenue_map[row[0]] = row[1] or 0
            recharge_cnt_map[row[0]] = row[2] or 0

    # 4. API cost per month (from zveno_usage_log)
    q_api = """
        SELECT strftime('%Y-%m', datetime(created_at, '+3 hours')) AS m,
               SUM(total_rub) AS cost
        FROM zveno_usage_log
        GROUP BY m
    """
    api_cost_map = {}
    async with db.execute(q_api) as cur:
        for row in await cur.fetchall():
            api_cost_map[row[0]] = row[1] or 0

    # 5. Generations per month (photo)
    q_gens = """
        SELECT strftime('%Y-%m', datetime(created_at, '+3 hours')) AS m,
               COUNT(*) AS cnt
        FROM balance_history
        WHERE reason = 'generation'
        GROUP BY m
    """
    gens_map = {}
    async with db.execute(q_gens) as cur:
        for row in await cur.fetchall():
            gens_map[row[0]] = row[1]

    # 5b. Video generations per month
    q_video_gens = """
        SELECT strftime('%Y-%m', datetime(created_at, '+3 hours')) AS m,
               COUNT(*) AS cnt
        FROM balance_history
        WHERE reason = 'video_generation'
        GROUP BY m
    """
    video_gens_map = {}
    async with db.execute(q_video_gens) as cur:
        for row in await cur.fetchall():
            video_gens_map[row[0]] = row[1]

    # Build monthly dict
    monthly = {}
    for i, mk in enumerate(month_keys):
        prev_mk = month_keys[i - 1] if i > 0 else None
        prev_set = active_sets.get(prev_mk, set()) if prev_mk else set()
        cur_set = active_sets.get(mk, set())
        churned = len(prev_set - cur_set)

        rev = revenue_map.get(mk, 0)
        cnt = recharge_cnt_map.get(mk, 0)
        monthly[mk] = {
            "new_paying": new_paying_map.get(mk, 0),
            "churned": churned,
            "active_paying": active_map.get(mk, 0),
            "revenue": rev,
            "api_cost": api_cost_map.get(mk, 0),
            "generations": gens_map.get(mk, 0),
            "video_generations": video_gens_map.get(mk, 0),
            "avg_recharge": (rev / cnt) if cnt else 0,
        }

    # Totals
    async with db.execute("SELECT COUNT(*) FROM users WHERE id > 0") as cur:
        total_users = (await cur.fetchone())[0]
    async with db.execute(
        "SELECT COUNT(DISTINCT user_id) FROM balance_history WHERE reason IN ('recharge','admin_edit') AND amount > 0"
    ) as cur:
        total_paying = (await cur.fetchone())[0]
    async with db.execute(
        "SELECT COALESCE(SUM(amount),0) FROM balance_history WHERE reason IN ('recharge','admin_edit') AND amount > 0"
    ) as cur:
        total_revenue = (await cur.fetchone())[0]
    async with db.execute(
        "SELECT COUNT(*) FROM balance_history WHERE reason = 'generation'"
    ) as cur:
        total_generations = (await cur.fetchone())[0]
    async with db.execute(
        "SELECT COUNT(*) FROM balance_history WHERE reason = 'video_generation'"
    ) as cur:
        total_video_generations = (await cur.fetchone())[0]

    # Paying this month (MSK +3h)
    async with db.execute(
        "SELECT COUNT(DISTINCT user_id) FROM balance_history "
        "WHERE reason IN ('recharge','admin_edit') AND amount > 0 "
        "AND strftime('%Y-%m', datetime(created_at, '+3 hours')) = strftime('%Y-%m', datetime('now', '+3 hours'))"
    ) as cur:
        paying_this_month = (await cur.fetchone())[0]

    totals = {
        "total_users": total_users,
        "total_paying": total_paying,
        "paying_this_month": paying_this_month,
        "total_revenue": total_revenue,
        "total_generations": total_generations,
        "total_video_generations": total_video_generations,
    }

    # Per-user metrics
    q_users = """
        SELECT
            u.id,
            u.username,
            u.first_name,
            u.created_at,
            u.source,
            u.balance,
            COALESCE(r.total_in, 0) AS total_recharged,
            COALESCE(r.cnt_in, 0) AS recharge_count,
            COALESCE(g.cnt_gen, 0) AS generation_count,
            COALESCE(g.total_out, 0) AS total_spent,
            COALESCE(r.last_in, g.last_gen) AS last_activity
        FROM users u
        LEFT JOIN (
            SELECT user_id,
                   SUM(amount) AS total_in,
                   COUNT(*) AS cnt_in,
                   MAX(created_at) AS last_in
            FROM balance_history
            WHERE reason IN ('recharge', 'admin_edit') AND amount > 0
            GROUP BY user_id
        ) r ON r.user_id = u.id
        LEFT JOIN (
            SELECT user_id,
                   COUNT(*) AS cnt_gen,
                   SUM(ABS(amount)) AS total_out,
                   MAX(created_at) AS last_gen
            FROM balance_history
            WHERE reason = 'generation'
            GROUP BY user_id
        ) g ON g.user_id = u.id
        WHERE u.id > 0
        ORDER BY total_recharged DESC
    """
    user_metrics = []
    async with db.execute(q_users) as cur:
        for row in await cur.fetchall():
            total_in = row[6] or 0
            cnt_in = row[7] or 0
            user_metrics.append({
                "user_id": row[0],
                "username": row[1],
                "first_name": row[2],
                "created_at": row[3] or "",
                "source": row[4] or "",
                "balance": row[5] or 0,
                "total_recharged": total_in,
                "recharge_count": cnt_in,
                "avg_recharge": (total_in / cnt_in) if cnt_in else 0,
                "generation_count": row[8] or 0,
                "total_spent": row[9] or 0,
                "last_activity": row[10] or "",
            })

    # Top-10 by revenue
    top_users = [u for u in user_metrics if u["total_recharged"] > 0][:10]
    for u in top_users:
        u["share_pct"] = (u["total_recharged"] / total_revenue * 100) if total_revenue else 0

    # At-risk: paid exactly once, last payment 14+ days ago
    q_at_risk = """
        SELECT u.id, u.username, u.first_name, u.source,
               bh.total_in, bh.last_pay
        FROM users u
        JOIN (
            SELECT user_id,
                   COUNT(*) AS cnt,
                   SUM(amount) AS total_in,
                   MAX(created_at) AS last_pay
            FROM balance_history
            WHERE reason IN ('recharge', 'admin_edit') AND amount > 0
            GROUP BY user_id
            HAVING cnt = 1
               AND datetime(MAX(created_at), '+3 hours') <= datetime('now', '+3 hours', '-14 days')
        ) bh ON bh.user_id = u.id
        ORDER BY bh.last_pay DESC
    """
    at_risk_users = []
    now_msk = datetime.utcnow() + timedelta(hours=3)
    async with db.execute(q_at_risk) as cur:
        for row in await cur.fetchall():
            last_pay_str = row[5] or ""
            try:
                last_pay_dt = datetime.strptime(last_pay_str[:10], "%Y-%m-%d")
                days_ago = (now_msk.date() - last_pay_dt.date()).days
            except Exception:
                days_ago = 0
            at_risk_users.append({
                "user_id": row[0],
                "username": row[1],
                "first_name": row[2],
                "source": row[3] or "",
                "total_in": row[4] or 0,
                "last_pay": last_pay_str,
                "days_ago": days_ago,
            })

    # Chart data
    chart_labels = [m["label"] for m in months]
    chart_revenue = [monthly[m["key"]]["revenue"] for m in months]
    chart_api_cost = [round(monthly[m["key"]]["api_cost"], 2) for m in months]
    chart_profit = [round(monthly[m["key"]]["revenue"] - monthly[m["key"]]["api_cost"], 2) for m in months]

    return {
        "months": months,
        "monthly": monthly,
        "totals": type("T", (), totals),
        "user_metrics": user_metrics,
        "top_users": top_users,
        "at_risk_users": at_risk_users,
        "chart_labels": chart_labels,
        "chart_revenue": chart_revenue,
        "chart_api_cost": chart_api_cost,
        "chart_profit": chart_profit,
    }


@app.get("/metrics", response_class=HTMLResponse)
async def metrics_page(
    request: Request,
    db: aiosqlite.Connection = Depends(get_db),
    user: str = Depends(get_current_username),
):
    data = await _compute_metrics(db)
    return templates.TemplateResponse("metrics.html", {"request": request, **data})


@app.get("/metrics/export")
async def metrics_export(
    db: aiosqlite.Connection = Depends(get_db),
    user: str = Depends(get_current_username),
):
    import io
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return Response("openpyxl not installed on server", status_code=500)

    data = await _compute_metrics(db)
    months = data["months"]
    monthly = data["monthly"]

    wb = Workbook()

    # Sheet 1: Monthly metrics
    ws = wb.active
    ws.title = "Метрики бизнеса"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=1 + len(months))
    ws["A1"] = "Динамика бизнеса по месяцам"
    ws["A1"].font = Font(bold=True, size=14)

    headers = ["Метрика"] + [m["label"] for m in months]
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=ci, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    rows_data = [
        ("Новых платящих клиентов (первое пополнение)", [monthly[m["key"]]["new_paying"] for m in months]),
        ("Клиенты перестали платить в этом месяце", [monthly[m["key"]]["churned"] for m in months]),
        ("Всего активных платящих на конец месяца", [monthly[m["key"]]["active_paying"] for m in months]),
        ("Выручка за месяц (сумма пополнений), ₽", [monthly[m["key"]]["revenue"] for m in months]),
        ("Расход API за месяц, ₽", [round(monthly[m["key"]]["api_cost"], 2) for m in months]),
        ("Чистая прибыль за месяц, ₽", [round(monthly[m["key"]]["revenue"] - monthly[m["key"]]["api_cost"], 2) for m in months]),
        ("Генераций за месяц", [monthly[m["key"]]["generations"] for m in months]),
        ("Средний чек пополнения, ₽", [round(monthly[m["key"]]["avg_recharge"]) if monthly[m["key"]]["avg_recharge"] else "" for m in months]),
    ]

    for ri, (label, values) in enumerate(rows_data, 4):
        ws.cell(row=ri, column=1, value=label).font = Font(bold=True)
        ws.cell(row=ri, column=1).border = thin_border
        for ci, v in enumerate(values, 2):
            cell = ws.cell(row=ri, column=ci, value=v)
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

    ws.column_dimensions["A"].width = 45
    for ci in range(2, 2 + len(months)):
        ws.column_dimensions[get_column_letter(ci)].width = 12

    # Sheet 2: Per-user metrics
    ws2 = wb.create_sheet("Пользователи")
    user_headers = [
        "ID", "Username", "Имя", "Регистрация", "Источник", "Баланс ₽",
        "Всего пополнений ₽", "Кол-во пополнений", "Средний чек ₽",
        "Генераций", "Расход на генерации ₽", "Последняя активность",
    ]
    for ci, h in enumerate(user_headers, 1):
        cell = ws2.cell(row=1, column=ci, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for ri, u in enumerate(data["user_metrics"], 2):
        vals = [
            u["user_id"], u["username"] or "", u["first_name"] or "",
            u["created_at"][:10] if u["created_at"] else "",
            u["source"], u["balance"],
            u["total_recharged"], u["recharge_count"],
            round(u["avg_recharge"]) if u["avg_recharge"] else "",
            u["generation_count"], u["total_spent"],
            u["last_activity"][:10] if u["last_activity"] else "",
        ]
        for ci, v in enumerate(vals, 1):
            cell = ws2.cell(row=ri, column=ci, value=v)
            cell.border = thin_border

    for ci in range(1, len(user_headers) + 1):
        ws2.column_dimensions[get_column_letter(ci)].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    now_str = (datetime.utcnow() + timedelta(hours=3)).strftime("%Y-%m-%d")
    filename = f"metrics_{now_str}.xlsx"

    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.post("/mailing/send")
async def send_mailing(background_tasks: BackgroundTasks, text: str = Form(...), user: str = Depends(get_current_username)):
    background_tasks.add_task(run_broadcast, text)
    return RedirectResponse(url="/mailing?status=sent", status_code=303)

@app.get("/prompts", response_class=HTMLResponse)
async def list_prompts(request: Request, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    try:
        async with db.execute("SELECT m.*, p.text as prompt_text, p.title as prompt_title FROM models m JOIN prompts p ON m.prompt_id = p.id") as cur:
            models = await cur.fetchall()
    except Exception: models = []
    
    categorized_models = {}
    for m in models:
        cat = m['category']
        if cat not in categorized_models:
            categorized_models[cat] = []
        categorized_models[cat].append(m)
    prompt_placeholders = await _get_prompt_placeholders(db)
    
    cat_prompt_keys = [
        "whitebg_prompt",
        "collage_prompt",
        "random_prompt",
        "random_other_prompt",
        "storefront_prompt",
        "infographic_clothing_prompt",
        "infographic_other_prompt",
        "own_prompt",
        "own_variant_prompt",
        "presets_prompt",
        "photostudio_wb_prompt",
        "video_prompt",
        "video_clothing_prompt",
    ]
    cat_prompts = {}
    async with db.execute(
        f"SELECT key, value FROM app_settings WHERE key IN ({','.join(['?']*len(cat_prompt_keys))})",
        cat_prompt_keys
    ) as cur:
        rows = await cur.fetchall()
        for r in rows:
            cat_prompts[r[0]] = r[1]
    for k in cat_prompt_keys:
        cat_prompts.setdefault(k, "")

    return templates.TemplateResponse("prompts.html", {
        "request": request, 
        "categorized_models": categorized_models, 
        "categories": CATEGORIES,
        "prompt_placeholders": prompt_placeholders,
        "cat_prompts": cat_prompts
    })

@app.get("/proxy", response_class=HTMLResponse)
async def list_proxies_page(request: Request, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    async with db.execute("SELECT * FROM proxies ORDER BY created_at DESC") as cur:
        proxies = await cur.fetchall()
    try:
        async with db.execute("SELECT * FROM api_keys ORDER BY is_active DESC, id") as cur:
            gemini_keys = await cur.fetchall()
    except Exception:
        gemini_keys = []

    
    proxy_to_key: dict[int, int] = {}
    try:
        async with db.execute("SELECT proxy_id, key_id FROM api_key_proxies ORDER BY id DESC") as cur:
            rows = await cur.fetchall()
        for r in rows:
            pid = int(r[0])
            kid = int(r[1])
            proxy_to_key.setdefault(pid, kid)
    except Exception:
        proxy_to_key = {}

    return templates.TemplateResponse("proxy.html", {
        "request": request,
        "proxies": proxies,
        "gemini_keys": gemini_keys,
        "proxy_to_key": proxy_to_key,
    })


@app.post("/admin/proxy/bind_key")
async def bind_proxy_to_key(
    proxy_id: int = Form(...),
    key_id: int = Form(0),
    db: aiosqlite.Connection = Depends(get_db),
    user: str = Depends(get_current_username),
):
    pid = int(proxy_id)
    kid = int(key_id or 0)
    if pid <= 0:
        return RedirectResponse(url="/proxy", status_code=303)

    
    if kid <= 0:
        await db.execute("DELETE FROM api_key_proxies WHERE proxy_id=?", (pid,))
        await db.commit()
        return RedirectResponse(url="/proxy", status_code=303)

    
    async with db.execute("SELECT COUNT(*) FROM api_key_proxies WHERE key_id=?", (kid,)) as cur:
        cnt = int((await cur.fetchone())[0] or 0)
    async with db.execute("SELECT 1 FROM api_key_proxies WHERE key_id=? AND proxy_id=? LIMIT 1", (kid, pid)) as cur:
        already = await cur.fetchone()
    if cnt >= 5 and not already:
        return RedirectResponse(url="/proxy?error=limit", status_code=303)

    
    await db.execute("DELETE FROM api_key_proxies WHERE proxy_id=?", (pid,))
    await db.execute("INSERT OR IGNORE INTO api_key_proxies (key_id, proxy_id) VALUES (?, ?)", (kid, pid))
    await db.commit()
    return RedirectResponse(url="/proxy", status_code=303)

@app.post("/admin/proxy/add")
async def add_proxy_route(urls: str = Form(...), db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    proxy_list = [u.strip() for u in urls.split("\n") if u.strip()]
    for url in proxy_list:
        
        parts = url.split(':')
        if len(parts) == 4 and "://" not in url:
            converted_url = f"http://{parts[2]}:{parts[3]}@{parts[0]}:{parts[1]}"
        elif not url.startswith("http") and "@" in url:
            converted_url = f"http://{url}"
        elif not url.startswith("http"):
            converted_url = f"http://{url}"
        else:
            converted_url = url
            
        await db.execute("INSERT INTO proxies (url) VALUES (?)", (converted_url,))
    await db.commit()
    return RedirectResponse(url="/proxy", status_code=303)

@app.post("/admin/proxy/delete")
async def delete_proxy_route(id: int = Form(...), db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    
    await db.execute("DELETE FROM api_key_proxies WHERE proxy_id = ?", (id,))
    await db.execute("DELETE FROM proxies WHERE id = ?", (id,))
    await db.commit()
    return RedirectResponse(url="/proxy", status_code=303)

@app.post("/admin/proxy/set_name")
async def set_proxy_name_route(
    id: int = Form(...),
    name: str = Form(""),
    db: aiosqlite.Connection = Depends(get_db),
    user: str = Depends(get_current_username),
):
    nm = (name or "").strip()
    if len(nm) > 64:
        nm = nm[:64]
    await db.execute(
        "UPDATE proxies SET name=?, updated_at=CURRENT_TIMESTAMP WHERE id=?",
        (nm, int(id)),
    )
    await db.commit()
    return RedirectResponse(url="/proxy", status_code=303)

@app.post("/admin/proxy/toggle")
async def toggle_proxy_route(id: int = Form(...), active: int = Form(...), db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    await db.execute("UPDATE proxies SET is_active = ? WHERE id = ?", (active, id))
    await db.commit()
    return RedirectResponse(url="/proxy", status_code=303)

@app.post("/admin/proxy/check")
async def check_proxy_route(id: int = Form(...), db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    async with db.execute("SELECT url FROM proxies WHERE id = ?", (id,)) as cur:
        row = await cur.fetchone()
        if not row:
            return RedirectResponse(url="/proxy", status_code=303)
        proxy_url = row[0]

    import httpx
    import time
    
    status = "failed"
    error_msg = None
    
    
    protocols = ["http://", "socks5://"]
    
    
    base_url = proxy_url
    if "://" in proxy_url:
        base_url = proxy_url.split("://", 1)[1]

    for proto in protocols:
        test_url = f"{proto}{base_url}"
        try:
            
            async with httpx.AsyncClient(proxy=test_url, timeout=15.0, verify=False) as client:
                
                resp = await client.get("https://www.google.com", follow_redirects=True)
                if resp.status_code < 400:
                    status = "working"
                    proxy_url = test_url 
                    break
                else:
                    error_msg = f"HTTP {resp.status_code} ({proto})"
        except Exception as e:
            error_msg = f"{str(e)} ({proto})"
            continue

    await db.execute(
        "UPDATE proxies SET url = ?, status = ?, error_message = ?, last_check = CURRENT_TIMESTAMP WHERE id = ?",
        (proxy_url, status, error_msg, id)
    )
    await db.commit()
    return RedirectResponse(url="/proxy", status_code=303)

@app.post("/prompts/category_prompts")
async def update_category_prompts(
    request: Request,
    whitebg_prompt: str = Form(""),
    collage_prompt: str = Form(""),
    random_prompt: str = Form(""),
    random_other_prompt: str = Form(""),
    storefront_prompt: str = Form(""),
    infographic_clothing_prompt: str = Form(""),
    infographic_other_prompt: str = Form(""),
    own_prompt: str = Form(""),
    own_variant_prompt: str = Form(""),
    presets_prompt: str = Form(""),
    db: aiosqlite.Connection = Depends(get_db),
    user: str = Depends(get_current_username)
):
    payload = {
        "whitebg_prompt": whitebg_prompt,
        "collage_prompt": collage_prompt,
        "random_prompt": random_prompt,
        "random_other_prompt": random_other_prompt,
        "storefront_prompt": storefront_prompt,
        "infographic_clothing_prompt": infographic_clothing_prompt,
        "infographic_other_prompt": infographic_other_prompt,
        "own_prompt": own_prompt,
        "own_variant_prompt": own_variant_prompt,
        "presets_prompt": presets_prompt,

        "own_prompt1": own_prompt,
        "own_prompt2": own_prompt,
        "own_prompt3": own_prompt,
    }
    for key, value in payload.items():
        await db.execute(
            "INSERT INTO app_settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (key, (value or "").strip())
        )
    await db.commit()
    return RedirectResponse(url="/prompts", status_code=303)


@app.post("/prompts/category_prompt/update")
async def update_single_category_prompt(
    key: str = Form(...),
    value: str = Form(""),
    db: aiosqlite.Connection = Depends(get_db),
    user: str = Depends(get_current_username),
):
    allowed_keys = {
        "whitebg_prompt",
        "collage_prompt",
        "random_prompt",
        "random_other_prompt",
        "storefront_prompt",
        "infographic_clothing_prompt",
        "infographic_other_prompt",
        "own_prompt",
        "own_variant_prompt",
        "presets_prompt",
        "photostudio_wb_prompt",
        "video_prompt",
        "video_clothing_prompt",
    }
    safe_key = (key or "").strip()
    if safe_key not in allowed_keys:
        return RedirectResponse(url="/prompts", status_code=303)

    safe_value = (value or "").strip()
    await db.execute(
        "INSERT INTO app_settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (safe_key, safe_value),
    )
    
    if safe_key == "own_prompt":
        for own_k in ("own_prompt1", "own_prompt2", "own_prompt3"):
            await db.execute(
                "INSERT INTO app_settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                (own_k, safe_value),
            )
    await db.commit()
    return RedirectResponse(url="/prompts", status_code=303)

@app.post("/prompts/edit_model")
async def edit_model_prompt(
    model_id: int = Form(...), 
    name: str = Form(...), 
    prompt_text: str = Form(...), 
    photo: UploadFile = File(None),
    db: aiosqlite.Connection = Depends(get_db), 
    user: str = Depends(get_current_username)
):
    await db.execute("UPDATE models SET name=? WHERE id=?", (name, model_id))
    if photo and photo.filename:
        file_path = f"{UPLOAD_DIR}/model_{model_id}.jpg"
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(photo.file, buffer)
        
        rel_path = os.path.relpath(file_path, BASE_DIR).replace("\\", "/")
        await db.execute("UPDATE models SET photo_file_id=? WHERE id=?", (rel_path, model_id))
    async with db.execute("SELECT prompt_id FROM models WHERE id=?", (model_id,)) as cur:
        row = await cur.fetchone()
        if row:
            await db.execute("UPDATE prompts SET text=? WHERE id=?", (prompt_text, row[0]))
    await db.commit()
    return RedirectResponse(url="/prompts", status_code=303)

@app.post("/models/add_full")
async def add_full_model(
    category: str = Form(...),
    name: str = Form(...),
    prompt_text: str = Form(...),
    cloth: str = Form("all"),
    photo: UploadFile = File(None),
    db: aiosqlite.Connection = Depends(get_db),
    user: str = Depends(get_current_username)
):
    await db.execute("INSERT INTO prompts (title, text) VALUES (?, ?)", (f"Prompt for {name}", prompt_text))
    async with db.execute("SELECT last_insert_rowid()") as cur:
        prompt_id = (await cur.fetchone())[0]
    await db.execute("INSERT INTO models (category, cloth, name, prompt_id) VALUES (?, ?, ?, ?)", 
                     (category, cloth, name, prompt_id))
    async with db.execute("SELECT last_insert_rowid()") as cur:
        model_id = (await cur.fetchone())[0]
    
    if photo and photo.filename:
        file_path = f"{UPLOAD_DIR}/model_{model_id}.jpg"
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(photo.file, buffer)
        
        rel_path = os.path.relpath(file_path, BASE_DIR).replace("\\", "/")
        await db.execute("UPDATE models SET photo_file_id=? WHERE id=?", (rel_path, model_id))
    
    await db.commit()
    return RedirectResponse(url="/prompts", status_code=303)

@app.get("/models/delete_photo/{model_id}")
async def delete_model_photo(model_id: int, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    async with db.execute("SELECT photo_file_id FROM models WHERE id=?", (model_id,)) as cur:
        row = await cur.fetchone()
        if row and row[0]:
            photo_path = row[0]
            if photo_path.startswith("data/"):
                full_path = os.path.join(BASE_DIR, photo_path)
                if os.path.exists(full_path):
                    try:
                        os.remove(full_path)
                    except Exception as e:
                        print(f"Error removing file {full_path}: {e}")
            await db.execute("UPDATE models SET photo_file_id=NULL WHERE id=?", (model_id,))
            await db.commit()
    return RedirectResponse(url="/prompts", status_code=303)

@app.get("/settings", response_class=HTMLResponse)
async def settings_page(request: Request, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    try:
        async with db.execute(
            "SELECT key, value FROM app_settings "
            "WHERE key IN ('agreement_text', 'howto_text', 'required_channel_id', 'required_channel_url', "
            "'required_vk_group_id', 'required_vk_group_url', 'support_contact', 'support_contact_vk')"
        ) as cur:
            rows = await cur.fetchall()
            settings_dict = {r[0]: r[1] for r in rows}
    except Exception: settings_dict = {}
    return templates.TemplateResponse("settings.html", {
        "request": request, 
        "agreement": settings_dict.get('agreement_text', ""), 
        "howto": settings_dict.get('howto_text', ""),
        "channel_id": settings_dict.get('required_channel_id', ""),
        "channel_url": settings_dict.get('required_channel_url', "https://t.me/bnbslow"),
        "vk_group_id": settings_dict.get('required_vk_group_id', ""),
        "vk_group_url": settings_dict.get('required_vk_group_url', ""),
        "support_contact": settings_dict.get('support_contact', "https://t.me/bnbslow"),
        "support_contact_vk": settings_dict.get('support_contact_vk', ""),
    })

@app.post("/settings/update")
async def update_app_settings(
    agreement: str = Form(""), 
    howto: str = Form(""), 
    channel_id: str = Form(None),
    channel_url: str = Form(None),
    vk_group_id: str = Form(None),
    vk_group_url: str = Form(None),
    support_contact: str = Form(None),
    support_contact_vk: str = Form(None),
    db: aiosqlite.Connection = Depends(get_db), 
    user: str = Depends(get_current_username)
):
    
    if channel_id:
        channel_id = channel_id.strip()
        if not channel_id.startswith("-100") and channel_id.isdigit():
            
            channel_id = f"-100{channel_id}"
        elif not channel_id.startswith("-") and not channel_id.startswith("@"):
            
            channel_id = f"@{channel_id}"

    await db.execute("INSERT INTO app_settings (key, value) VALUES ('agreement_text', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (agreement,))
    await db.execute("INSERT INTO app_settings (key, value) VALUES ('howto_text', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (howto,))
    await db.execute("INSERT INTO app_settings (key, value) VALUES ('required_channel_id', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (channel_id or "",))
    await db.execute("INSERT INTO app_settings (key, value) VALUES ('required_channel_url', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (channel_url or "https://t.me/bnbslow",))
    await db.execute("INSERT INTO app_settings (key, value) VALUES ('required_vk_group_id', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", ((vk_group_id or "").strip(),))
    await db.execute("INSERT INTO app_settings (key, value) VALUES ('required_vk_group_url', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", ((vk_group_url or "").strip(),))
    await db.execute("INSERT INTO app_settings (key, value) VALUES ('support_contact', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (support_contact or "https://t.me/bnbslow",))
    await db.execute("INSERT INTO app_settings (key, value) VALUES ('support_contact_vk', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", ((support_contact_vk or "").strip(),))
    await db.commit()
    return RedirectResponse(url="/settings", status_code=303)

@app.post("/settings/test_channel")
async def test_channel_settings(
    db: aiosqlite.Connection = Depends(get_db),
    user: str = Depends(get_current_username)
):
    channel_id = ""
    async with db.execute("SELECT value FROM app_settings WHERE key='required_channel_id'") as cur:
        row = await cur.fetchone()
        if row: channel_id = row[0]
    
    if not channel_id:
        return JSONResponse({"status": "error", "message": "ID канала не установлен"})

    bot = Bot(token=settings.bot_token)
    try:
        chat = await bot.get_chat(channel_id)
        
        member = await bot.get_chat_member(channel_id, (await bot.get_me()).id)
        can_check = member.status in ("administrator", "creator")
        
        return JSONResponse({
            "status": "success", 
            "message": f"Канал найден: {chat.title}. Бот имеет права администратора: {'Да' if can_check else 'Нет'}"
        })
    except Exception as e:
        return JSONResponse({"status": "error", "message": f"Ошибка: {str(e)}"})
    finally:
        await bot.session.close()

@app.get("/api_keys", response_class=HTMLResponse)
async def list_keys(request: Request, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    try:
        async with db.execute("SELECT * FROM api_keys ORDER BY is_active DESC, id") as cur:
            gemini_keys = await cur.fetchall()
    except Exception: gemini_keys = []

    key_proxy_lists: dict[int, list[dict]] = {}
    try:
        async with db.execute(
            "SELECT akp.key_id, p.id, p.name, p.url, p.is_active, p.status, p.error_message, p.last_check "
            "FROM api_key_proxies akp "
            "JOIN proxies p ON p.id = akp.proxy_id "
            "ORDER BY akp.key_id, p.id"
        ) as cur:
            rows = await cur.fetchall()
    except Exception:
        rows = []

    for r in rows:
        kid = int(r[0])
        key_proxy_lists.setdefault(kid, []).append({
            "id": int(r[1]),
            "name": r[2],
            "url": r[3],
            "is_active": int(r[4]) == 1,
            "status": r[5],
            "error_message": r[6],
            "last_check": r[7],
        })

    try:
        async with db.execute("PRAGMA table_info(zveno_usage_log)") as cur:
            zveno_cols = await cur.fetchall()
        zveno_col_names = {str(r[1]) for r in (zveno_cols or [])}
        if "user_id" not in zveno_col_names:
            await db.execute("ALTER TABLE zveno_usage_log ADD COLUMN user_id INTEGER")
            await db.commit()
    except Exception:
        pass

    try:
        async with db.execute(
            """
            SELECT
                COALESCE(SUM(prompt_tokens), 0),
                COALESCE(SUM(completion_tokens), 0),
                COALESCE(SUM(total_tokens), 0),
                COALESCE(SUM(input_rub), 0.0),
                COALESCE(SUM(output_rub), 0.0),
                COALESCE(SUM(total_rub), 0.0),
                COUNT(*)
            FROM zveno_usage_log
            WHERE status='ok'
            """
        ) as cur:
            usage_row = await cur.fetchone()
    except Exception:
        usage_row = (0, 0, 0, 0.0, 0.0, 0.0, 0)

    try:
        async with db.execute(
            """
            SELECT
                created_at,
                COALESCE(user_id, 0),
                COALESCE(key_preview, ''),
                COALESCE(model, ''),
                COALESCE(status, ''),
                COALESCE(error_text, '')
            FROM zveno_usage_log
            WHERE status != 'ok'
            ORDER BY id DESC
            LIMIT 80
            """
        ) as cur:
            zveno_error_rows = await cur.fetchall()
    except Exception:
        zveno_error_rows = []

    try:
        async with db.execute(
            "SELECT key, value FROM app_settings WHERE key IN ('zveno_total_balance_rub', 'zveno_note')"
        ) as cur:
            settings_rows = await cur.fetchall()
        settings_map = {str(r[0]): str(r[1]) for r in settings_rows}
    except Exception:
        settings_map = {}

    total_balance_rub_raw = settings_map.get("zveno_total_balance_rub", "0")
    try:
        total_balance_rub = float(total_balance_rub_raw)
    except Exception:
        total_balance_rub = 0.0
    spent_rub = float(usage_row[5] or 0.0)
    remaining_rub = max(0.0, total_balance_rub - spent_rub)
    zveno_note = settings_map.get("zveno_note", "")

    flat_rate = float(os.environ.get("ZVENO_FLAT_RATE_PER_CALL_RUB", "6.3"))
    return templates.TemplateResponse("api_keys.html", {
        "request": request,
        "gemini_keys": gemini_keys,
        "key_proxy_lists": key_proxy_lists,
        "zveno_usage": {
            "prompt_tokens": int(usage_row[0] or 0),
            "completion_tokens": int(usage_row[1] or 0),
            "total_tokens": int(usage_row[2] or 0),
            "input_rub": float(usage_row[3] or 0.0),
            "output_rub": float(usage_row[4] or 0.0),
            "spent_rub": spent_rub,
            "request_count": int(usage_row[6] or 0),
            "total_balance_rub": total_balance_rub,
            "remaining_rub": remaining_rub,
            "note": zveno_note,
            "flat_rate": flat_rate,
        },
        "zveno_error_rows": zveno_error_rows,
    })


@app.post("/api_keys/zveno_budget")
async def update_zveno_budget(
    total_balance_rub: str = Form("0"),
    note: str = Form(""),
    db: aiosqlite.Connection = Depends(get_db),
    user: str = Depends(get_current_username),
):
    try:
        total_value = max(0.0, float((total_balance_rub or "0").replace(",", ".")))
    except Exception:
        total_value = 0.0
    await db.execute(
        "INSERT INTO app_settings (key, value) VALUES ('zveno_total_balance_rub', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (f"{total_value:.2f}",),
    )
    await db.execute(
        "INSERT INTO app_settings (key, value) VALUES ('zveno_note', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        ((note or "").strip()[:300],),
    )
    await db.commit()
    return RedirectResponse(url="/api_keys", status_code=303)


@app.post("/api_keys/{key_id}/proxies")
async def api_key_set_proxies(
    key_id: int,
    proxy_ids: list[int] = Form([]),
    db: aiosqlite.Connection = Depends(get_db),
    user: str = Depends(get_current_username),
):
    uniq: list[int] = []
    seen = set()
    for pid in proxy_ids or []:
        v = int(pid)
        if v <= 0 or v in seen:
            continue
        seen.add(v)
        uniq.append(v)
        if len(uniq) >= 5:
            break
    await db.execute("DELETE FROM api_key_proxies WHERE key_id=?", (int(key_id),))
    for pid in uniq:
        await db.execute(
            "INSERT OR IGNORE INTO api_key_proxies (key_id, proxy_id) VALUES (?, ?)",
            (int(key_id), int(pid)),
        )
    await db.commit()
    return RedirectResponse(url="/api_keys", status_code=303)

@app.post("/api_keys/add")
async def add_key(token: str = Form(...), db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    
    token = "".join(token.split())
    
    
    cyr_to_lat = {
        'А': 'A', 'В': 'B', 'Е': 'E', 'К': 'K', 'М': 'M', 'Н': 'H', 
        'О': 'O', 'Р': 'P', 'С': 'C', 'Т': 'T', 'У': 'y', 'Х': 'X',
        'а': 'a', 'е': 'e', 'о': 'o', 'р': 'p', 'с': 'c', 'у': 'y', 'х': 'x'
    }
    new_token = ""
    for char in token:
        new_token += cyr_to_lat.get(char, char)
    token = new_token

    import re
    token = re.sub(r'[^A-Za-z0-9_\-]', '', token)
    
    await db.execute("INSERT INTO api_keys (token) VALUES (?)", (token,))
    await db.commit()
    return RedirectResponse(url="/api_keys", status_code=303)

@app.get("/api_keys/toggle/{key_id}")
async def toggle_key(key_id: int, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    async with db.execute("SELECT is_active FROM api_keys WHERE id=?", (key_id,)) as cur:
        row = await cur.fetchone()
    if not row:
        return RedirectResponse(url="/api_keys", status_code=303)
    new_val = 0 if int(row[0]) == 1 else 1
    await db.execute("UPDATE api_keys SET is_active=? WHERE id=?", (new_val, key_id))
    await db.commit()
    return RedirectResponse(url="/api_keys", status_code=303)

@app.get("/api_keys/delete/{key_id}")
async def delete_key(key_id: int, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    await db.execute("DELETE FROM api_keys WHERE id=?", (key_id,))
    await db.commit()
    return RedirectResponse(url="/api_keys", status_code=303)

@app.get("/tg_img/{file_id}")
async def get_telegram_image(file_id: str, user: str = Depends(get_current_username)):
    
    if not file_id or file_id == "None" or file_id == "FILE_ID_MOCK":
        return Response(status_code=404)
    if settings.bot_token == "MOCK":
        return Response(status_code=500, content="Bot token not configured")
    
    bot = Bot(token=settings.bot_token)
    try:
        
        file = await bot.get_file(file_id)
        file_url = f"https://api.telegram.org/file/bot{settings.bot_token}/{file.file_path}"
        
        
        async with httpx.AsyncClient(follow_redirects=True) as client:
            
            proxy_url = settings.proxy.as_url() if hasattr(settings, "proxy") else None
            
            
            
            resp = await client.get(file_url, timeout=30)
            
            if resp.status_code == 200:
                return Response(content=resp.content, media_type="image/jpeg")
            
            
            if proxy_url:
                async with httpx.AsyncClient(proxy=proxy_url, follow_redirects=True) as p_client:
                    resp = await p_client.get(file_url, timeout=30)
                    if resp.status_code == 200:
                        return Response(content=resp.content, media_type="image/jpeg")
            
            print(f"Telegram image {file_id} error: {resp.status_code}")
            return Response(status_code=resp.status_code)
            
    except Exception as e:
        print(f"Critical error proxying Telegram image {file_id}: {e}")
        return Response(status_code=500, content=str(e))
    finally:
        await bot.session.close()

@app.get("/history", response_class=HTMLResponse)
async def list_history(request: Request, q: str = "", db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    try:
        if q:
            # Добавили OR category LIKE ? и третий параметр f"%{q}%"
            query = "SELECT * FROM generation_history WHERE pid LIKE ? OR user_id LIKE ? OR category LIKE ? ORDER BY created_at DESC LIMIT 50"
            async with db.execute(query, (f"%{q}%", f"%{q}%", f"%{q}%")) as cur:
                history = await cur.fetchall()
        else:
            async with db.execute("SELECT * FROM generation_history ORDER BY created_at DESC LIMIT 50") as cur:
                history = await cur.fetchall()
    except Exception: history = []
    return templates.TemplateResponse("history.html", {"request": request, "history": history, "q": q})

@app.get("/payments", response_class=HTMLResponse)
async def list_payments(request: Request, q: str = "", db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    try:
        if q:
            query = """
                SELECT p.*, sp.name_ru as plan_name 
                FROM payments p 
                LEFT JOIN subscription_plans sp ON p.plan_id = sp.id 
                WHERE p.user_id LIKE ? 
                ORDER BY p.created_at DESC LIMIT 100
            """
            async with db.execute(query, (f"%{q}%",)) as cur:
                payments = await cur.fetchall()
        else:
            query = """
                SELECT p.*, sp.name_ru as plan_name 
                FROM payments p 
                LEFT JOIN subscription_plans sp ON p.plan_id = sp.id 
                ORDER BY p.created_at DESC LIMIT 100
            """
            async with db.execute(query) as cur:
                payments = await cur.fetchall()
    except Exception as e:
        print(f"Payments error: {e}")
        payments = []
        
    return templates.TemplateResponse("payments.html", {"request": request, "payments": payments, "q": q})

@app.get("/prices", response_class=HTMLResponse)
async def list_prices(request: Request, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    status_data = []
    for cat in CATEGORIES:
        status_key = f"{cat}" 
        
        async with db.execute("SELECT value FROM app_settings WHERE key=?", (status_key,)) as cur:
            row = await cur.fetchone()
            
            status_val = row[0] if row else "1"
            
        status_data.append({
            "status_key": status_key,
            "cat": cat, 
            "is_enabled": status_val == "1"
        })
            
    async with db.execute("SELECT * FROM subscription_plans") as cur:
        plans = await cur.fetchall()
        
    return templates.TemplateResponse("prices.html", {"request": request, "category_status": status_data, "plans": plans})

@app.post("/plans/edit")
async def edit_plan(
    plan_id: int = Form(...),
    name_ru: str = Form(...),
    name_en: str = Form(...),
    name_vi: str = Form(...),
    price: int = Form(...),
    duration: int = Form(...),
    limit: int = Form(...),
    desc_ru: str = Form(None),
    desc_en: str = Form(None),
    desc_vi: str = Form(None),
    db: aiosqlite.Connection = Depends(get_db),
    user: str = Depends(get_current_username)
):
    try:
        await db.execute(
            """UPDATE subscription_plans 
               SET name_ru=?, name_en=?, name_vi=?, price=?, duration_days=?, daily_limit=?, 
                   description_ru=?, description_en=?, description_vi=? 
               WHERE id=?""",
            (name_ru, name_en, name_vi, price, duration, limit, 
             desc_ru, desc_en, desc_vi, plan_id)
        )
        await db.commit()
        return RedirectResponse(url="/prices", status_code=303)
    except Exception as e:
        print(f"Error in edit_plan: {e}")
        return JSONResponse(status_code=500, content={"detail": str(e)})

@app.post("/categories/toggle")
async def toggle_category(key: str = Form(...), db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    async with db.execute("SELECT value FROM app_settings WHERE key=?", (key,)) as cur:
        row = await cur.fetchone()
        current = row[0] if row else "1"
    
    new_val = "0" if current == "1" else "1"
    await db.execute("INSERT INTO app_settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, new_val))
    await db.commit()
    return RedirectResponse(url="/prices", status_code=303)

@app.post("/edit_subscription")
async def edit_subscription(
    user_id: int = Form(...), 
    plan_id: str = Form(None),
    days: int = Form(...), 
    limit: int = Form(...), 
    api_key: str = Form(None),
    db: aiosqlite.Connection = Depends(get_db), 
    user: str = Depends(get_current_username)
):
    try:
        expires_at = datetime.utcnow() + timedelta(days=days)
        expires_str = expires_at.strftime("%Y-%m-%d %H:%M:%S")
        plan_type = "custom"
        
        plan_id_val = None
        if plan_id and plan_id.isdigit():
            plan_id_val = int(plan_id)
            async with db.execute("SELECT name_ru FROM subscription_plans WHERE id=?", (plan_id_val,)) as cur:
                row = await cur.fetchone()
                if row:
                    plan_type = row[0]
        elif plan_id == "custom":
            plan_type = "custom"
        
        
        safe_api_key = None
        if api_key and api_key.strip():
            
            token = "".join(api_key.split())
            
            cyr_to_lat = {
                'А': 'A', 'В': 'B', 'Е': 'E', 'К': 'K', 'М': 'M', 'Н': 'H', 
                'О': 'O', 'Р': 'P', 'С': 'C', 'Т': 'T', 'У': 'y', 'Х': 'X',
                'а': 'a', 'е': 'e', 'о': 'o', 'р': 'p', 'с': 'c', 'у': 'y', 'х': 'x'
            }
            new_token = ""
            for char in token:
                new_token += cyr_to_lat.get(char, char)
            
            import re
            safe_api_key = re.sub(r'[^A-Za-z0-9_\-]', '', new_token)
        
        
        
        await db.execute("DELETE FROM subscriptions WHERE user_id = ?", (user_id,))
        
        await db.execute(
            "INSERT INTO subscriptions (user_id, plan_id, plan_type, expires_at, daily_limit, daily_usage, last_usage_reset, individual_api_key) VALUES (?, ?, ?, ?, ?, 0, CURRENT_TIMESTAMP, ?)",
            (user_id, plan_id_val, plan_type, expires_str, limit, safe_api_key)
        )
        
        
        await db.execute("UPDATE users SET trial_used=1 WHERE id=?", (user_id,))
        
        
        amount = 0
        if plan_id_val:
            async with db.execute("SELECT price FROM subscription_plans WHERE id=?", (plan_id_val,)) as cur:
                row = await cur.fetchone()
                if row:
                    amount = row[0]
        
        await db.execute(
            "INSERT INTO payments (user_id, plan_id, amount, status) VALUES (?, ?, ?, 'admin_granted')",
            (user_id, plan_id_val, amount)
        )
            
        await db.commit()

        
        try:
            async with db.execute("SELECT language FROM users WHERE id=?", (user_id,)) as cur:
                row = await cur.fetchone()
                lang = row[0] if row and row[0] else "ru"
            
            await send_subscription_notification(user_id, plan_type, expires_str, limit, lang)
        except Exception as e:
            print(f"Error in post-grant notification: {e}")

        return RedirectResponse(url=f"/users?q={user_id}", status_code=303)
    except Exception as e:
        print(f"Error in edit_subscription: {e}")
        return JSONResponse(status_code=500, content={"detail": str(e)})



@app.get("/proxy_single", response_class=HTMLResponse)
async def proxy_page(request: Request, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    current_proxy = ""
    try:
        async with db.execute("SELECT value FROM app_settings WHERE key='bot_proxy'") as cur:
            row = await cur.fetchone()
            current_proxy = row[0] if row else ""
    except Exception: pass
    status_text = await check_proxy(db)
    return templates.TemplateResponse("proxy.html", {"request": request, "current_proxy": current_proxy, "status": status_text})

@app.post("/proxy_single/edit")
async def edit_proxy(proxy_url: str = Form(...), db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    await db.execute("INSERT INTO app_settings (key, value) VALUES ('bot_proxy', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (proxy_url,))
    await db.commit()
    return RedirectResponse(url="/proxy_single", status_code=303)


@app.get("/api/mtproxy/public")
async def mtproxy_public_link(db: aiosqlite.Connection = Depends(get_db)):
    
    import base64
    
    
    async with db.execute("SELECT value FROM app_settings WHERE key='mtproxy_running'") as cur:
        row = await cur.fetchone()
        running = row and row[0] == "1"
    
    if not running:
        return JSONResponse({"available": False})
    
    
    async with db.execute("SELECT value FROM app_settings WHERE key='mtproxy_secret'") as cur:
        row = await cur.fetchone()
        secret = row[0] if row else None
    
    async with db.execute("SELECT value FROM app_settings WHERE key='mtproxy_port'") as cur:
        row = await cur.fetchone()
        port = int(row[0]) if row and row[0] else 8888
    
    if not secret:
        return JSONResponse({"available": False})
    
    
    server_ip = os.getenv("MTPROXY_SERVER_IP", "130.49.148.147")
    
    try:
        secret_hex = secret.replace("-", "")
        if secret_hex.startswith("dd") and len(secret_hex) == 34:
            secret_hex = secret_hex[2:]
        if len(secret_hex) == 32:
            link = f"tg://proxy?server={server_ip}&port={port}&secret=dd{secret_hex}"
            return JSONResponse({"available": True, "link": link})
        else:
            return JSONResponse({"available": False})
    except Exception:
        return JSONResponse({"available": False})


@app.get("/api/mtproxy/status")
async def mtproxy_status(db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    
    import subprocess
    import base64
    
    
    async with db.execute("SELECT value FROM app_settings WHERE key='mtproxy_running'") as cur:
        row = await cur.fetchone()
        db_running = row and row[0] == "1"
    
    
    try:
        result = subprocess.run(
            ["docker", "inspect", "-f", "{{.State.Running}}", "mtproxy"],
            capture_output=True,
            text=True,
            timeout=5
        )
        docker_running = result.returncode == 0 and result.stdout.strip().lower() == "true"
        
        if docker_running != db_running:
            await db.execute(
                "INSERT INTO app_settings (key, value) VALUES ('mtproxy_running', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                ("1" if docker_running else "0",)
            )
            await db.commit()
        running = docker_running
    except Exception as e:
        print(f"Error checking Docker status: {e}")
        running = db_running
    
    
    async with db.execute("SELECT value FROM app_settings WHERE key='mtproxy_secret'") as cur:
        row = await cur.fetchone()
        secret = row[0] if row else None
    
    async with db.execute("SELECT value FROM app_settings WHERE key='mtproxy_port'") as cur:
        row = await cur.fetchone()
        port = int(row[0]) if row and row[0] else 8888
    
    link = None
    if secret and running:
        
        server_ip = os.getenv("MTPROXY_SERVER_IP", "130.49.148.147")
        
        secret_hex = secret.replace("-", "")
        if secret_hex.startswith("dd") and len(secret_hex) == 34:
            secret_hex = secret_hex[2:]
        if len(secret_hex) == 32:
            link = f"tg://proxy?server={server_ip}&port={port}&secret=dd{secret_hex}"
        else:
            
            try:
                secret_bytes = bytes.fromhex(secret_hex) if len(secret_hex) == 32 else base64.b64decode(secret)
                secret_hex_clean = secret_bytes.hex()
                link = f"tg://proxy?server={server_ip}&port={port}&secret=dd{secret_hex_clean}"
            except:
                pass
    
    return JSONResponse({
        "running": running,
        "link": link,
        "port": port
    })


@app.post("/api/mtproxy/generate")
async def mtproxy_generate(db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    
    import secrets
    import subprocess
    import os
    
    
    secret_bytes = secrets.token_bytes(16)
    secret_hex = secret_bytes.hex()
    
    
    await db.execute(
        "INSERT INTO app_settings (key, value) VALUES ('mtproxy_secret', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (secret_hex,)
    )
    
    
    async with db.execute("SELECT value FROM app_settings WHERE key='mtproxy_port'") as cur:
        row = await cur.fetchone()
        if not row or not row[0]:
            await db.execute(
                "INSERT INTO app_settings (key, value) VALUES ('mtproxy_port', '8888') ON CONFLICT(key) DO UPDATE SET value=excluded.value"
            )
    
    await db.commit()
    
    
    server_ip = os.getenv("MTPROXY_SERVER_IP", "130.49.148.147")
    async with db.execute("SELECT value FROM app_settings WHERE key='mtproxy_port'") as cur:
        row = await cur.fetchone()
        port = int(row[0]) if row and row[0] else 8888
    
    
    link = f"tg://proxy?server={server_ip}&port={port}&secret=dd{secret_hex}"
    
    
    try:
        result = subprocess.run(
            ["docker", "inspect", "-f", "{{.State.Running}}", "mtproxy"],
            capture_output=True,
            text=True,
            timeout=5
        )
        is_running = result.returncode == 0 and result.stdout.strip().lower() == "true"
        
        if not is_running:
            
            secret_for_env = secret_hex
            env = os.environ.copy()
            env["MTPROXY_SECRET"] = secret_for_env
            os.makedirs(os.path.join(BASE_DIR, "data", "mtproxy"), exist_ok=True)
            
            
            subprocess.run(["docker", "rm", "-f", "mtproxy"], timeout=10, check=False)
            
            
            cmd = [
                "docker", "run", "-d",
                "--name", "mtproxy",
                "--restart", "unless-stopped",
                "-p", f"{port}:443",
                "-e", f"SECRET={secret_for_env}",
                "-v", f"{BASE_DIR}/data/mtproxy:/data",
                "telegrammessenger/proxy:latest"
            ]
            run_result = subprocess.run(cmd, capture_output=True, text=True, timeout=30, check=False, env=env)
            if run_result.returncode != 0:
                return JSONResponse({"error": (run_result.stderr or run_result.stdout or "Не удалось запустить MTProxy").strip()}, status_code=500)

            verify = subprocess.run(
                ["docker", "inspect", "-f", "{{.State.Running}}", "mtproxy"],
                capture_output=True,
                text=True,
                timeout=5
            )
            if not (verify.returncode == 0 and verify.stdout.strip().lower() == "true"):
                return JSONResponse({"error": "MTProxy контейнер создан, но не запущен"}, status_code=500)

        await db.execute(
            "INSERT INTO app_settings (key, value) VALUES ('mtproxy_running', '1') ON CONFLICT(key) DO UPDATE SET value=excluded.value"
        )
        await db.commit()
    except Exception as e:
        return JSONResponse({"error": f"Ошибка запуска MTProxy: {e}"}, status_code=500)
    
    return JSONResponse({"link": link, "secret": secret_hex, "running": True})


@app.post("/api/mtproxy/toggle")
async def mtproxy_toggle(db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    
    import subprocess
    import os
    
    
    try:
        result = subprocess.run(
            ["docker", "inspect", "-f", "{{.State.Running}}", "mtproxy"],
            capture_output=True,
            text=True,
            timeout=5
        )
        running = result.returncode == 0 and result.stdout.strip().lower() == "true"
    except Exception as e:
        print(f"Error checking MTProxy status in toggle: {e}")
        running = False
    
    
    async with db.execute("SELECT value FROM app_settings WHERE key='mtproxy_secret'") as cur:
        row = await cur.fetchone()
        secret = row[0] if row else None
    
    if not secret and not running:
        return JSONResponse({"error": "Сначала сгенерируйте секрет"}, status_code=400)
    
    try:
        if running:
            
            subprocess.run(["docker", "rm", "-f", "mtproxy"], timeout=10, check=False)
            
            await db.execute(
                "INSERT INTO app_settings (key, value) VALUES ('mtproxy_running', '0') ON CONFLICT(key) DO UPDATE SET value=excluded.value"
            )
        else:
            
            async with db.execute("SELECT value FROM app_settings WHERE key='mtproxy_port'") as cur:
                row = await cur.fetchone()
                port = int(row[0]) if row and row[0] else 8888
            
            
            secret_hex = secret.replace("-", "")
            if secret_hex.startswith("dd") and len(secret_hex) == 34:
                secret_hex = secret_hex[2:]
            if len(secret_hex) != 32:
                return JSONResponse({"error": "Некорректный формат секрета MTProxy"}, status_code=400)
            
            
            env = os.environ.copy()
            env["MTPROXY_SECRET"] = secret_hex
            os.makedirs(os.path.join(BASE_DIR, "data", "mtproxy"), exist_ok=True)
            
            
            subprocess.run(["docker", "rm", "-f", "mtproxy"], timeout=10, check=False)
            
            
            cmd = [
                "docker", "run", "-d",
                "--name", "mtproxy",
                "--restart", "unless-stopped",
                "-p", f"{port}:443",
                "-e", f"SECRET={secret_hex}",
                "-v", f"{BASE_DIR}/data/mtproxy:/data",
                "telegrammessenger/proxy:latest"
            ]
            run_result = subprocess.run(cmd, capture_output=True, text=True, timeout=30, check=False, env=env)
            if run_result.returncode != 0:
                return JSONResponse({"error": (run_result.stderr or run_result.stdout or "Не удалось запустить MTProxy").strip()}, status_code=500)
            verify = subprocess.run(
                ["docker", "inspect", "-f", "{{.State.Running}}", "mtproxy"],
                capture_output=True,
                text=True,
                timeout=5
            )
            if not (verify.returncode == 0 and verify.stdout.strip().lower() == "true"):
                return JSONResponse({"error": "MTProxy контейнер создан, но не запущен"}, status_code=500)
            
            await db.execute(
                "INSERT INTO app_settings (key, value) VALUES ('mtproxy_running', '1') ON CONFLICT(key) DO UPDATE SET value=excluded.value"
            )
        await db.commit()
        
        new_running = not running
        return JSONResponse({"success": True, "running": new_running})
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)

@app.post("/maintenance/toggle")
async def toggle_maintenance(db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    current = False
    try:
        async with db.execute("SELECT value FROM app_settings WHERE key='maintenance'") as cur:
            row = await cur.fetchone()
            current = row[0] == '1' if row else False
    except Exception: pass
    
    new_val = '0' if current else '1'
    now_str = datetime.now().isoformat()

    if new_val == '1':
        await db.execute("INSERT INTO app_settings (key, value) VALUES ('maintenance_start', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (now_str,))
    else:
        try:
            async with db.execute("SELECT value FROM app_settings WHERE key='maintenance_start'") as cur:
                row = await cur.fetchone()
                if row:
                    start_time = datetime.fromisoformat(row[0])
                    duration_seconds = int((datetime.now() - start_time).total_seconds())
                    if duration_seconds > 0:
                        await db.execute(
                            "UPDATE subscriptions SET expires_at = datetime(expires_at, '+' || ? || ' seconds') WHERE expires_at > CURRENT_TIMESTAMP",
                            (str(duration_seconds),)
                        )
        except Exception:
            pass

    await db.execute("INSERT INTO app_settings (key, value) VALUES ('maintenance', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (new_val,))
    await db.commit()
    return RedirectResponse(url="/", status_code=303)


@app.post("/settings/disable_4k/toggle")
async def toggle_disable_4k(db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    current = False
    try:
        async with db.execute("SELECT value FROM app_settings WHERE key='disable_4k'") as cur:
            row = await cur.fetchone()
            if row is None:
                current = os.getenv("DISABLE_4K", "0").strip() == "1"
            else:
                current = str(row[0]).strip() == "1"
    except Exception:
        pass

    new_val = "0" if current else "1"
    await db.execute(
        "INSERT INTO app_settings (key, value) VALUES ('disable_4k', ?) "
        "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (new_val,),
    )
    await db.commit()
    return RedirectResponse(url="/", status_code=303)


@app.get("/constructor", response_class=HTMLResponse)
async def constructor_page(request: Request, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    async with db.execute("SELECT id, key, name_ru, is_active, order_index FROM categories ORDER BY order_index, id") as cur:
        categories = await cur.fetchall()
    return templates.TemplateResponse("constructor.html", {"request": request, "categories": categories})

@app.post("/constructor/category/add")
async def admin_add_category(key: str = Form(...), name: str = Form(...), order: int = Form(0), db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    await db.execute(
        "INSERT OR IGNORE INTO categories (key, name_ru, is_active, order_index) VALUES (?, ?, 1, ?)",
        (key, name, order)
    )
    await db.commit()
    return RedirectResponse("/constructor", status_code=303)

@app.post("/constructor/category/delete/{cat_id}")
async def admin_delete_category(cat_id: int, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    
    
    async with db.execute("SELECT key FROM categories WHERE id=?", (cat_id,)) as cur:
        row = await cur.fetchone()
    if row and row["key"] in CORE_CATEGORY_KEYS:
        await db.execute("UPDATE categories SET is_active=0 WHERE id=?", (cat_id,))
        await db.commit()
        return RedirectResponse("/constructor", status_code=303)

    
    await db.execute("DELETE FROM step_options WHERE step_id IN (SELECT id FROM steps WHERE category_id=?)", (cat_id,))
    await db.execute("DELETE FROM steps WHERE category_id=?", (cat_id,))
    await db.execute("DELETE FROM categories WHERE id=?", (cat_id,))
    await db.commit()
    return RedirectResponse("/constructor", status_code=303)

@app.post("/constructor/category/toggle/{cat_id}")
async def admin_toggle_category(cat_id: int, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    async with db.execute("SELECT is_active FROM categories WHERE id=?", (cat_id,)) as cur:
        row = await cur.fetchone()
    if not row:
        return RedirectResponse("/constructor", status_code=303)
    new_val = 0 if int(row["is_active"]) == 1 else 1
    await db.execute("UPDATE categories SET is_active=? WHERE id=?", (new_val, cat_id))
    await db.commit()
    return RedirectResponse("/constructor", status_code=303)

@app.get("/constructor/category/{cat_id}", response_class=HTMLResponse)
async def category_steps_page(request: Request, cat_id: int, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    async with db.execute("SELECT id, key, name_ru, is_active, order_index FROM categories WHERE id=?", (cat_id,)) as cur:
        category = await cur.fetchone()
    if not category:
        return RedirectResponse("/constructor")
    
    async with db.execute(
        "SELECT id, step_key, question_text, question_text_en, question_text_vi, input_type, is_optional, order_index FROM steps WHERE category_id=? ORDER BY order_index, id",
        (cat_id,)
    ) as cur:
        steps = await cur.fetchall()
    
    steps_with_options = []
    for step in steps:
        async with db.execute(
            "SELECT id, option_text, option_text_en, option_text_vi, option_value, order_index, custom_prompt FROM step_options WHERE step_id=? ORDER BY order_index, id",
            (step['id'],)
        ) as cur_opt:
            options_raw = await cur_opt.fetchall()
            options = [
                {
                    "id": r[0],
                    "text": r[1],
                    "text_en": r[2] or "",
                    "text_vi": r[3] or "",
                    "value": r[4],
                    "order": r[5],
                    "prompt": r[6] or ""
                } for r in options_raw
            ]
        
        steps_with_options.append({
            "id": step['id'],
            "key": step['step_key'],
            "question": step['question_text'],
            "question_en": step['question_text_en'] or "",
            "question_vi": step['question_text_vi'] or "",
            "type": step['input_type'],
            "optional": step['is_optional'],
            "order": step['order_index'],
            "options": options
        })

    
    async with db.execute("SELECT id, step_key, question_text, question_text_en, question_text_vi, input_type FROM library_steps ORDER BY id") as cur:
        lib_steps_raw = await cur.fetchall()

    lib_steps = []
    for step in lib_steps_raw:
        step_id, step_key, question_text, question_text_en, question_text_vi, input_type = step
        async with db.execute(
            "SELECT option_text, option_text_en, option_text_vi, option_value, custom_prompt FROM library_step_options WHERE step_id=? ORDER BY order_index, id",
            (step_id,)
        ) as cur:
            opts = await cur.fetchall()
        default_buttons = json.dumps(
            [{"text": o[0], "text_en": o[1] or "", "text_vi": o[2] or "", "value": o[3], "prompt": o[4] or ""} for o in opts],
            ensure_ascii=False
        )
        lib_steps.append({
            "id": step_id,
            "step_key": step_key,
            "question_text": question_text,
            "question_text_en": question_text_en or "",
            "question_text_vi": question_text_vi or "",
            "input_type": input_type,
            "default_buttons": default_buttons
        })

    async with db.execute("SELECT id, name FROM button_categories ORDER BY id") as cur:
        button_cats = await cur.fetchall()

    lib_buttons = {}
    for bcat in button_cats:
        async with db.execute(
            "SELECT id, option_text, option_text_en, option_text_vi, option_value, custom_prompt FROM library_options WHERE category_id=? ORDER BY id", 
            (bcat['id'],)
        ) as cur:
            lib_buttons[bcat['name']] = await cur.fetchall()

    prompt_placeholders = await _get_prompt_placeholders(db)
    
    return templates.TemplateResponse("category_edit.html", {
        "request": request, 
        "category": category, 
        "steps": steps_with_options,
        "lib_steps": lib_steps,
        "lib_buttons": lib_buttons,
        "prompt_placeholders": prompt_placeholders
    })


@app.get("/constructor/category/{cat_id}/languages", response_class=HTMLResponse)
async def category_languages_page(request: Request, cat_id: int, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    async with db.execute("SELECT id, key, name_ru, is_active, order_index FROM categories WHERE id=?", (cat_id,)) as cur:
        category = await cur.fetchone()
    if not category:
        return RedirectResponse("/constructor")

    async with db.execute(
        "SELECT id, step_key, question_text, question_text_en, question_text_vi "
        "FROM steps WHERE category_id=? ORDER BY order_index, id",
        (cat_id,)
    ) as cur:
        steps = await cur.fetchall()

    options_by_step = {}
    async with db.execute(
        "SELECT id, step_id, option_text, option_text_en, option_text_vi, option_value "
        "FROM step_options WHERE step_id IN (SELECT id FROM steps WHERE category_id=?) "
        "ORDER BY step_id, order_index, id",
        (cat_id,)
    ) as cur:
        rows = await cur.fetchall()
        for row in rows:
            options_by_step.setdefault(row["step_id"], []).append(row)

    return templates.TemplateResponse("category_languages.html", {
        "request": request,
        "category": category,
        "steps": steps,
        "options_by_step": options_by_step
    })


@app.post("/constructor/category/{cat_id}/save_translations")
async def save_category_translations(request: Request, cat_id: int, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    form = await request.form()
    for key, val in form.items():
        value = (val or "").strip()
        if key.startswith("step_"):
            parts = key.split("_")
            if len(parts) == 3:
                step_id = int(parts[1])
                lang = parts[2]
                if lang in ("en", "vi"):
                    col = f"question_text_{lang}"
                    await db.execute(f"UPDATE steps SET {col}=? WHERE id=?", (value, step_id))
        elif key.startswith("opt_"):
            parts = key.split("_")
            if len(parts) == 3:
                opt_id = int(parts[1])
                lang = parts[2]
                if lang in ("en", "vi"):
                    col = f"option_text_{lang}"
                    await db.execute(f"UPDATE step_options SET {col}=? WHERE id=?", (value, opt_id))
    await db.commit()
    return RedirectResponse(f"/constructor/category/{cat_id}/languages", status_code=303)

@app.post("/constructor/category/update/{cat_id}")
async def admin_update_category(cat_id: int, name: str = Form(...), key: str = Form(...), is_active: int = Form(1), order: int = Form(0), db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    await db.execute(
        "UPDATE categories SET name_ru=?, key=?, is_active=?, order_index=? WHERE id=?",
        (name, key, is_active, order, cat_id)
    )
    await db.commit()
    return RedirectResponse(f"/constructor/category/{cat_id}", status_code=303)

@app.post("/constructor/step/add/{cat_id}")
async def admin_add_step(cat_id: int, step_key: str = Form(...), question: str = Form(...), input_type: str = Form(...), is_optional: int = Form(0), order: int = Form(0), db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    if order == 0:
        async with db.execute("SELECT MAX(order_index) FROM steps WHERE category_id=?", (cat_id,)) as cur:
            row = await cur.fetchone()
            order = (row[0] or 0) + 1
            
    await db.execute(
        "INSERT INTO steps (category_id, step_key, question_text, input_type, is_optional, order_index) VALUES (?, ?, ?, ?, ?, ?)",
        (cat_id, step_key, question, input_type, is_optional, order)
    )
    await db.commit()
    return RedirectResponse(f"/constructor/category/{cat_id}", status_code=303)

@app.post("/constructor/step/delete/{cat_id}/{step_id}")
async def admin_delete_step(cat_id: int, step_id: int, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    await db.execute("DELETE FROM step_options WHERE step_id=?", (step_id,))
    await db.execute("DELETE FROM steps WHERE id=?", (step_id,))
    await db.commit()
    return RedirectResponse(f"/constructor/category/{cat_id}", status_code=303)

@app.post("/constructor/option/add/{cat_id}/{step_id}")
async def admin_add_option(cat_id: int, step_id: int, text: str = Form(...), value: str = Form(...), order: int = Form(0), db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    await db.execute(
        "INSERT INTO step_options (step_id, option_text, option_value, order_index) VALUES (?, ?, ?, ?)",
        (step_id, text, value, order)
    )
    await db.commit()
    return RedirectResponse(f"/constructor/category/{cat_id}", status_code=303)

@app.post("/constructor/option/delete/{cat_id}/{opt_id}")
async def admin_delete_option(cat_id: int, opt_id: int, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    await db.execute("DELETE FROM step_options WHERE id=?", (opt_id,))
    await db.commit()
    return RedirectResponse(f"/constructor/category/{cat_id}", status_code=303)

@app.post("/constructor/steps/reorder/{cat_id}")
async def admin_reorder_steps(cat_id: int, request: Request, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    data = await request.json()
    step_ids = data.get("step_ids", [])
    
    
    for index, step_id in enumerate(step_ids, 1):
        await db.execute("UPDATE steps SET order_index=? WHERE id=?", (index, step_id))
    
    await db.commit()
    return JSONResponse({"status": "ok"})

@app.post("/constructor/category/{cat_id}/save_all")
async def admin_save_all_steps(cat_id: int, request: Request, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    try:
        data = await request.json()
        print(f"Starting save_all for category {cat_id}. Data: {len(data.get('steps', []))} steps.")
        steps_data = data.get("steps", [])
        
        
        async with db.execute("SELECT id FROM steps WHERE category_id=?", (cat_id,)) as cur:
            existing_step_ids = [row[0] for row in await cur.fetchall()]
        
        received_step_ids = []
        
        for s_data in steps_data:
            step_id = s_data.get("id")
            step_key = s_data.get("key")
            question = s_data.get("question")
            question_en = s_data.get("question_en") or ""
            question_vi = s_data.get("question_vi") or ""
            input_type = s_data.get("type")
            is_optional = s_data.get("optional")
            order_index = int(s_data.get("order") or 0)
            buttons = s_data.get("buttons", [])
            
            
            if step_id == 'null' or step_id is None:
                step_id = None
            else:
                try:
                    step_id = int(step_id)
                except (ValueError, TypeError):
                    step_id = None
            
            if step_id and step_id in existing_step_ids:
                
                await db.execute(
                    "UPDATE steps SET question_text=?, question_text_en=?, question_text_vi=?, input_type=?, is_optional=?, order_index=?, step_key=? WHERE id=?",
                    (question, question_en, question_vi, input_type, is_optional, order_index, step_key, step_id)
                )
                received_step_ids.append(step_id)
            else:
                
                await db.execute(
                    "INSERT INTO steps (category_id, step_key, question_text, question_text_en, question_text_vi, input_type, is_optional, order_index) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                    (cat_id, step_key, question, question_en, question_vi, input_type, is_optional, order_index)
                )
                async with db.execute("SELECT last_insert_rowid()") as cur_last_step:
                    step_id = (await cur_last_step.fetchone())[0]
                received_step_ids.append(step_id)
                
            
            async with db.execute("SELECT id FROM step_options WHERE step_id=?", (int(step_id),)) as cur_opts:
                existing_opt_ids = [row[0] for row in await cur_opts.fetchall()]
                
            received_opt_ids = []
            for b_data in buttons:
                opt_id = b_data.get("id")
                opt_text = b_data.get("text")
                opt_text_en = b_data.get("text_en")
                opt_text_vi = b_data.get("text_vi")
                opt_value = b_data.get("value")
                opt_prompt = b_data.get("prompt")
                opt_order = int(b_data.get("order") or 0)
                
                
                if opt_id == 'null' or opt_id is None:
                    opt_id = None
                else:
                    try:
                        opt_id = int(opt_id)
                    except (ValueError, TypeError):
                        opt_id = None
                
                if opt_id and opt_id in existing_opt_ids:
                    await db.execute(
                        "UPDATE step_options SET option_text=?, option_text_en=?, option_text_vi=?, option_value=?, order_index=?, custom_prompt=? WHERE id=?",
                        (opt_text, opt_text_en, opt_text_vi, opt_value, opt_order, opt_prompt, opt_id)
                    )
                    received_opt_ids.append(opt_id)
                else:
                    await db.execute(
                        "INSERT INTO step_options (step_id, option_text, option_text_en, option_text_vi, option_value, order_index, custom_prompt) VALUES (?, ?, ?, ?, ?, ?, ?)",
                        (int(step_id), opt_text, opt_text_en, opt_text_vi, opt_value, opt_order, opt_prompt)
                    )
                    async with db.execute("SELECT last_insert_rowid()") as cur_last_opt:
                        received_opt_ids.append((await cur_last_opt.fetchone())[0])
                        
            
            for old_opt_id in existing_opt_ids:
                if old_opt_id not in received_opt_ids:
                    await db.execute("DELETE FROM step_options WHERE id=?", (old_opt_id,))
                    
        
        for old_step_id in existing_step_ids:
            if old_step_id not in received_step_ids:
                await db.execute("DELETE FROM step_options WHERE step_id=?", (old_step_id,))
                await db.execute("DELETE FROM steps WHERE id=?", (old_step_id,))
                
        await db.commit()
        return JSONResponse({"status": "ok"})
    except Exception as e:
        print(f"Error in save_all_steps: {e}")
        return JSONResponse({"status": "error", "message": str(e)}, status_code=500)

@app.post("/constructor/library/step/add")
async def admin_add_library_step(step_key: str = Form(...), question: str = Form(...), input_type: str = Form(...), db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    await db.execute(
        "INSERT INTO library_steps (step_key, question_text, input_type) VALUES (?, ?, ?)",
        (step_key, question, input_type)
    )
    await db.commit()
    return RedirectResponse(request.headers.get("referer", "/constructor"), status_code=303)

@app.post("/constructor/library/step/update/{step_id}")
async def admin_update_library_step(step_id: int, request: Request, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    data = await request.json()
    step_data = data.get("step", {})
    options = data.get("options", [])
    await db.execute(
        "UPDATE library_steps SET step_key=?, question_text=?, input_type=? WHERE id=?",
        (step_data.get("key"), step_data.get("question"), step_data.get("type"), step_id)
    )

    
    async with db.execute("SELECT id FROM library_step_options WHERE step_id=?", (step_id,)) as cur:
        existing_ids = [row[0] for row in await cur.fetchall()]
    received_ids = []
    for b_data in options:
        opt_id = b_data.get("id")
        opt_text = b_data.get("text")
        opt_text_en = b_data.get("text_en") or ""
        opt_text_vi = b_data.get("text_vi") or ""
        opt_value = b_data.get("value")
        opt_prompt = b_data.get("prompt")
        opt_order = b_data.get("order")

        if opt_id == 'null' or opt_id is None:
            opt_id = None
        else:
            try:
                opt_id = int(opt_id)
            except (ValueError, TypeError):
                opt_id = None

        if opt_id and opt_id in existing_ids:
            await db.execute(
                "UPDATE library_step_options SET option_text=?, option_text_en=?, option_text_vi=?, option_value=?, order_index=?, custom_prompt=? WHERE id=?",
                (opt_text, opt_text_en, opt_text_vi, opt_value, opt_order, opt_prompt, opt_id)
            )
            received_ids.append(opt_id)
        else:
            await db.execute(
                "INSERT INTO library_step_options (step_id, option_text, option_text_en, option_text_vi, option_value, order_index, custom_prompt) VALUES (?, ?, ?, ?, ?, ?, ?)",
                (step_id, opt_text, opt_text_en, opt_text_vi, opt_value, opt_order, opt_prompt)
            )
            async with db.execute("SELECT last_insert_rowid()") as cur_last_opt:
                received_ids.append((await cur_last_opt.fetchone())[0])

    for old_id in existing_ids:
        if old_id not in received_ids:
            await db.execute("DELETE FROM library_step_options WHERE id=?", (old_id,))

    await db.commit()
    return JSONResponse({"status": "ok"})

@app.post("/constructor/library/step/delete/{step_id}")
async def admin_delete_library_step(step_id: int, request: Request, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    await db.execute("DELETE FROM library_step_options WHERE step_id=?", (step_id,))
    await db.execute("DELETE FROM library_steps WHERE id=?", (step_id,))
    await db.commit()
    return RedirectResponse(request.headers.get("referer", "/constructor"), status_code=303)

@app.post("/constructor/library/button/add")
async def admin_add_library_button(request: Request, text: str = Form(...), text_en: str = Form(""), text_vi: str = Form(""), value: str = Form(...), category: str = Form("Системные"), db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    
    async with db.execute("SELECT id FROM button_categories WHERE name=?", (category,)) as cur:
        row = await cur.fetchone()
        if row:
            cat_id = row[0]
        else:
            await db.execute("INSERT INTO button_categories (name) VALUES (?)", (category,))
            async with db.execute("SELECT last_insert_rowid()") as cur_new:
                cat_id = (await cur_new.fetchone())[0]
                
    await db.execute(
        "INSERT INTO library_options (category_id, option_text, option_text_en, option_text_vi, option_value) VALUES (?, ?, ?, ?, ?)",
        (cat_id, text, text_en, text_vi, value)
    )
    await db.commit()
    return RedirectResponse(request.headers.get("referer", "/constructor"), status_code=303)

@app.post("/constructor/library/button/delete/{btn_id}")
async def admin_delete_library_button(request: Request, btn_id: int, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    await db.execute("DELETE FROM library_options WHERE id=?", (btn_id,))
    await db.commit()
    return RedirectResponse(request.headers.get("referer", "/constructor"), status_code=303)

@app.post("/constructor/library/category/add")
async def admin_add_library_category(request: Request, name: str = Form(...), db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    await db.execute("INSERT OR IGNORE INTO button_categories (name) VALUES (?)", (name,))
    await db.commit()
    return RedirectResponse(request.headers.get("referer", "/constructor"), status_code=303)



@app.get("/support", response_class=HTMLResponse)
async def get_support(request: Request, user_id: int = None, db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    
    async with db.execute("""
        SELECT u.id, u.username, u.first_name, COALESCE(u.source, 'telegram') as source,
               (SELECT COUNT(*) FROM support_messages WHERE user_id = u.id AND is_admin = 0 AND is_read = 0) as unread_count,
               (SELECT MAX(created_at) FROM support_messages WHERE user_id = u.id) as last_msg_at
        FROM users u
        WHERE EXISTS (SELECT 1 FROM support_messages WHERE user_id = u.id)
        ORDER BY last_msg_at DESC
    """) as cur:
        support_users = await cur.fetchall()

    messages = []
    current_user = None
    if user_id:
        
        await db.execute("UPDATE support_messages SET is_read = 1 WHERE user_id = ? AND is_admin = 0", (user_id,))
        await db.commit()
        
        async with db.execute("SELECT message_text, is_admin, created_at, file_id, file_type FROM support_messages WHERE user_id = ? ORDER BY created_at ASC", (user_id,)) as cur:
            messages = await cur.fetchall()
        
        async with db.execute("SELECT id, username, first_name, COALESCE(source, 'telegram') FROM users WHERE id = ?", (user_id,)) as cur:
            current_user = await cur.fetchone()

    return templates.TemplateResponse("support.html", {
        "request": request, 
        "support_users": support_users, 
        "messages": messages, 
        "current_user": current_user
    })

@app.get("/support/file/{file_id}")
async def proxy_telegram_file(file_id: str, user: str = Depends(get_current_username)):
    
    try:
        settings = load_settings()
        async with httpx.AsyncClient() as client:
            
            get_file_url = f"https://api.telegram.org/bot{settings.bot_token}/getFile?file_id={file_id}"
            resp = await client.get(get_file_url)
            file_data = resp.json()
            if not file_data.get("ok"):
                return Response(status_code=404)
            
            file_path = file_data["result"]["file_path"]
            
            download_url = f"https://api.telegram.org/file/bot{settings.bot_token}/{file_path}"
            file_resp = await client.get(download_url)
            
            return Response(content=file_resp.content, media_type=file_resp.headers.get("content-type"))
    except Exception as e:
        print(f"Error proxying file {file_id}: {e}")
        return Response(status_code=500)

@app.post("/support/send")
async def send_support_reply(
    user_id: int = Form(...),
    message_text: str = Form(None),
    file: UploadFile = File(None),
    db: aiosqlite.Connection = Depends(get_db),
    user: str = Depends(get_current_username)
):
    try:
        settings = load_settings()
        file_id = None
        file_type = 'text'

        async with db.execute("SELECT COALESCE(source, 'telegram') FROM users WHERE id = ?", (user_id,)) as cur:
            row = await cur.fetchone()
        user_source = row[0] if row else "telegram"

        async with httpx.AsyncClient() as client:
            if user_source == "vk":
                if file and file.filename:
                    return JSONResponse(status_code=400, content={"detail": "VK support replies support text only"})
                if not settings.vk_bot_token:
                    raise RuntimeError("VK_BOT_TOKEN is required for VK support replies")
                if message_text:
                    resp = await client.post(
                        "https://api.vk.com/method/messages.send",
                        data={
                            "access_token": settings.vk_bot_token,
                            "v": "5.199",
                            "peer_id": user_id,
                            "random_id": int(datetime.utcnow().timestamp() * 1000) & 0x7FFFFFFF,
                            "message": f"Ответ поддержки:\n\n{message_text}",
                        },
                    )
                    result = resp.json()
                    if "error" in result:
                        raise RuntimeError(result["error"].get("error_msg", "VK send failed"))
            elif file and file.filename:
                
                file_content = await file.read()
                files = {"photo" if "image" in file.content_type else "video": (file.filename, file_content, file.content_type)}
                
                method = "sendPhoto" if "image" in file.content_type else "sendVideo"
                url = f"https://api.telegram.org/bot{settings.bot_token}/{method}"
                
                payload = {"chat_id": user_id}
                if message_text:
                    payload["caption"] = f"👨‍💻 <b>Ответ поддержки:</b>\n\n{message_text}"
                    payload["parse_mode"] = "HTML"
                else:
                    payload["caption"] = "👨‍💻 <b>Ответ поддержки</b>"
                    payload["parse_mode"] = "HTML"

                resp = await client.post(url, data=payload, files=files)
                result = resp.json()
                
                if result.get("ok"):
                    msg = result["result"]
                    if "photo" in msg:
                        file_id = msg["photo"][-1]["file_id"]
                        file_type = 'photo'
                    elif "video" in msg:
                        file_id = msg["video"]["file_id"]
                        file_type = 'video'
            
            elif message_text:
                url = f"https://api.telegram.org/bot{settings.bot_token}/sendMessage"
                payload = {
                    "chat_id": user_id,
                    "text": f"👨‍💻 <b>Ответ поддержки:</b>\n\n{message_text}",
                    "parse_mode": "HTML"
                }
                resp = await client.post(url, json=payload)
                result = resp.json()

            
            await db.execute(
                "INSERT INTO support_messages (user_id, message_text, file_id, file_type, is_admin, is_read) VALUES (?, ?, ?, ?, 1, 1)",
                (user_id, message_text, file_id, file_type)
            )
            await db.commit()

        return RedirectResponse(url=f"/support?user_id={user_id}", status_code=303)
    except Exception as e:
        print(f"Error sending support reply: {e}")
        return JSONResponse(status_code=500, content={"detail": str(e)})

@app.post("/support/mark_read")
async def mark_read(user_id: int = Form(...), db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    await db.execute("UPDATE support_messages SET is_read = 1 WHERE user_id = ? AND is_admin = 0", (user_id,))
    await db.commit()
    return {"status": "ok"}


@app.get("/dl/{filename:path}")
async def force_download_file(filename: str):
    import os
    from fastapi.responses import FileResponse, JSONResponse

    # Очищаем имя файла от лишнего (оставляем только result_XXX.jpg)
    clean_filename = os.path.basename(filename)

    # Жесткий поиск по всем возможным папкам внутри контейнера
    search_paths = [
        f"/app/data/history/{clean_filename}",
        f"/app/data/{clean_filename}",
        f"/app/bot/data/history/{clean_filename}",
        f"/app/admin_web/data/history/{clean_filename}"
    ]

    found_path = None
    for path in search_paths:
        if os.path.exists(path):
            found_path = path
            break

    # Если файла физически нет на диске сервера:
    if not found_path:
        return JSONResponse(status_code=404, content={
            "error": "Файл физически отсутствует на сервере",
            "requested_file": clean_filename,
            "searched_paths": search_paths
        })

    # Секретный соус: заголовок attachment заставляет телефон/ПК СКАЧАТЬ файл
    return FileResponse(
        path=found_path,
        filename=clean_filename,
        media_type="application/octet-stream",
        headers={"Content-Disposition": f'attachment; filename="{clean_filename}"'}
    )
@app.get("/view/{filename:path}")
async def view_file_inline(filename: str):
    """
    Новый эндпоинт: если нужно, чтобы картинка просто открылась на странице
    в браузере (например, для истории), используйте ссылку /view/имя_файла.jpg
    """
    import os
    from fastapi.responses import FileResponse

    clean_filename = os.path.basename(filename)
    possible_paths = [
        f"/app/data/history/{clean_filename}",
        f"/app/data/{clean_filename}",
        f"/app/bot/data/history/{clean_filename}",
        f"/app/admin_web/data/history/{clean_filename}"
    ]

    found_path = None
    for p in possible_paths:
        if os.path.exists(p):
            found_path = p
            break

    if not found_path:
        return {"detail": "Не найдено", "searched_paths": possible_paths}

    # Заголовок inline заставляет браузер ОТКРЫТЬ картинку, а не скачивать её
    return FileResponse(
        path=found_path,
        filename=clean_filename,
        media_type="image/jpeg",
        headers={"Content-Disposition": f'inline; filename="{clean_filename}"'}
    )


@app.get("/admin/backup_db")
async def backup_database(user: str = Depends(get_current_username)):
    import os
    from fastapi.responses import FileResponse, JSONResponse
    from datetime import datetime

    db_path = "/app/data/bot.db"

    # Проверяем, существует ли файл
    if not os.path.exists(db_path):
        return JSONResponse(status_code=404, content={"error": "База данных физически отсутствует"})

    # Добавляем текущую дату и время в название скачиваемого файла (чтобы не путаться в бэкапах)
    date_str = datetime.now().strftime("%Y-%m-%d_%H-%M")
    filename = f"bot_backup_{date_str}.db"

    # Заголовок attachment заставляет браузер именно СКАЧАТЬ файл
    return FileResponse(
        path=db_path,
        filename=filename,
        media_type="application/octet-stream",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@app.get("/api/accounts/live")
async def live_accounts_status(db: aiosqlite.Connection = Depends(get_db), user: str = Depends(get_current_username)):
    import os
    import json
    import time

    status_file = os.path.join(BASE_DIR, "data", "keys_status.json")
    live_statuses = {}
    if os.path.exists(status_file):
        try:
            with open(status_file, "r", encoding="utf-8") as f:
                live_statuses = json.load(f)
        except:
            pass

    db_stats = {}
    try:
        async with db.execute("SELECT project_id, generation_count FROM gcp_usage") as cur:
            rows = await cur.fetchall()
            for r in rows:
                db_stats[r[0]] = int(r[1])
    except Exception:
        pass

    now = time.time()
    result = {}
    for proj_id, bot_data in live_statuses.items():
        gen_count = db_stats.get(proj_id, 0)
        spent_dollars = gen_count * 0.039
        left_dollars = max(0, 300.0 - spent_dollars)
        max_gens = 7692
        load_percent = min(100, int((gen_count / max_gens) * 100))

        # Вычисляем живой статус прямо сейчас
        time_left = 0
        is_active = False

        if bot_data.get("is_unsupported"):
            state = "Отключен (Ошибка)"
        elif bot_data.get("transient_error_until", 0) > now:
            time_left = int(bot_data["transient_error_until"] - now)
            state = "Сетевой сбой (500)"
        elif bot_data.get("rate_limit_until", 0) > now:
            time_left = int(bot_data["rate_limit_until"] - now)
            state = "Охлаждение (429)"
        elif bot_data.get("cooldown_until", 0) > now:
            time_left = int(bot_data["cooldown_until"] - now)
            state = "Внутренняя пауза"
        else:
            state = "Активен"
            is_active = True

        # Живой RPM (если минута с момента первого запроса прошла, сбрасываем в 0 визуально)
        if bot_data.get("window_started_at", 0) + 60 < now:
            rpm_current = 0
        else:
            rpm_current = bot_data.get("requests_in_window", 0)

        result[proj_id] = {
            "status": state,
            "time_left": time_left,
            "rpm_current": rpm_current,
            "rpm_max": bot_data.get("rpm_max", 8),
            "is_active": is_active,
            "last_error": bot_data.get("last_error", ""),
            "gen_count": gen_count,
            "spent_dollars": round(spent_dollars, 2),
            "left_dollars": round(left_dollars, 2),
            "load_percent": load_percent
        }

    return JSONResponse(result)
# ==========================================
# МОНИТОРИНГ АККАУНТОВ VERTEX AI (TRIAL $300)
# ==========================================

@app.get("/accounts", response_class=HTMLResponse)
async def view_accounts(request: Request, db: aiosqlite.Connection = Depends(get_db),
                        user: str = Depends(get_current_username)):
    import os
    import json

    keys_dir = os.path.join(BASE_DIR, "data", "keys")
    status_file = os.path.join(BASE_DIR, "data", "keys_status.json")
    os.makedirs(keys_dir, exist_ok=True)

    # Читаем живые статусы от бота
    live_statuses = {}
    if os.path.exists(status_file):
        try:
            with open(status_file, "r", encoding="utf-8") as f:
                live_statuses = json.load(f)
        except:
            pass

    accounts = []
    if os.path.exists(keys_dir):
        for filename in os.listdir(keys_dir):
            if filename.endswith(".json"):
                filepath = os.path.join(keys_dir, filename)
                try:
                    with open(filepath, "r", encoding="utf-8") as f:
                        data = json.load(f)
                        project_id = data.get("project_id") or data.get("quota_project_id") or "Неизвестный проект"
                        acc_type = data.get("type", "Неизвестный")
                        client_email = data.get("client_email", "Авторизованный пользователь")

                        # Достаем статистику из базы (сколько фото сгенерировано)
                        async with db.execute("SELECT generation_count FROM gcp_usage WHERE project_id=?",
                                              (project_id,)) as cur:
                            row = await cur.fetchone()
                            gen_count = int(row[0]) if row else 0

                        # Новая математика: ~$0.039 за генерацию (1K). Триал $300.
                        # $300 / $0.039 = ~7692 фото на один аккаунт.
                        max_gens = 7692
                        spent_dollars = gen_count * 0.039
                        left_dollars = max(0, 300.0 - spent_dollars)
                        load_percent = min(100, int((gen_count / max_gens) * 100))

                        # Получаем статус из памяти бота (или ставим дефолтный)
                        bot_memory = live_statuses.get(project_id, {})
                        live_status = bot_memory.get("status", "В пуле (Ожидает)")
                        if load_percent >= 100:
                            live_status = "ЛИМИТ $300 ИСЧЕРПАН"

                        accounts.append({
                            "filename": filename,
                            "project_id": project_id,
                            "type": acc_type,
                            "client_email": client_email,
                            "status": live_status,
                            "last_error": bot_memory.get("last_error", ""),
                            "gen_count": gen_count,
                            "spent_dollars": round(spent_dollars, 2),
                            "left_dollars": round(left_dollars, 2),
                            "load_percent": load_percent
                        })
                except Exception as e:
                    accounts.append(
                        {"filename": filename, "project_id": "Ошибка", "status": "Ошибка файла", "load_percent": 0})

    return templates.TemplateResponse("accounts.html", {"request": request, "accounts": accounts})
